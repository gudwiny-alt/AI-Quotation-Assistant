from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import datetime
from itertools import count
from io import BytesIO
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
from typing import Iterator
import warnings

from openpyxl.drawing.image import Image as OpenpyxlImage  # type: ignore[import-untyped]
from openpyxl.utils.units import points_to_pixels  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.datavalidation import (  # type: ignore[import-untyped]
    DataValidation,
)
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from quote_app.core.formulas import (
    formula_cells,
    is_new,
    needs_six_month_adjustment,
)
from quote_app.core.months import required_price_months
from quote_app.core.normalization import normalize_code
from quote_app.core.precheck import month_header
from quote_app.domain.models import QuoteMonth, QuoteRow
from quote_app.excel.template_builder import (
    STYLE_SKELETON_ROW,
    TEMPLATE_SHEET_NAME,
    load_clean_template,
)


HEADER_ROW = 1
FIRST_DATA_ROW = 2
LAST_TEMPLATE_COLUMN = 42
MANUAL_COLUMNS = frozenset(("K", "L", "M", "N", "P", "Q"))
PERCENTAGE_COLUMNS = ("X", "Y", "Z")
N_VALIDATION_FORMULA = '"货源充足,货源紧缺,新品上市,尾货期"'
_MONTH_TOKEN = re.compile(r"\d{4}年\d{1,2}月")
_EVIDENCE_ANCHOR = re.compile(r"^(?:AL|AM|AN)(?:[2-9]|[1-9]\d+)$")
_TEMPLATE_COLUMNS = frozenset(
    get_column_letter(column) for column in range(1, LAST_TEMPLATE_COLUMN + 1)
)


@dataclass(frozen=True, slots=True)
class QuoteEvidenceImage:
    anchor: str
    payload: bytes

    def __post_init__(self) -> None:
        if (
            not isinstance(self.anchor, str)
            or _EVIDENCE_ANCHOR.fullmatch(self.anchor) is None
        ):
            raise ValueError("evidence image anchor must be AL/AM/AN data cell")
        if not isinstance(self.payload, bytes) or not self.payload:
            raise ValueError("evidence image payload must be non-empty bytes")


@dataclass(frozen=True, slots=True)
class QuoteWriteRequest:
    quote_month: QuoteMonth
    rows: list[QuoteRow]
    template_path: str | Path
    output_dir: Path
    evidence_images: tuple[QuoteEvidenceImage, ...] = ()
    destination_path: Path | None = None


def write_quote_workbook(request: QuoteWriteRequest) -> Path:
    """Generate a new quotation workbook while leaving its template unchanged."""
    output_dir = _ensure_output_directory(Path(request.output_dir))
    temporary_path: Path | None = None
    published_path: Path | None = None
    try:
        workbook = load_clean_template(Path(request.template_path))
        try:
            if TEMPLATE_SHEET_NAME not in workbook.sheetnames:
                raise ValueError(f"报价模板缺少工作表：{TEMPLATE_SHEET_NAME}")
            sheet = workbook[TEMPLATE_SHEET_NAME]
            for extra_sheet in tuple(workbook.worksheets):
                if extra_sheet is not sheet:
                    workbook.remove(extra_sheet)

            _update_month_headers(sheet, request.quote_month)
            _write_rows(sheet, request.rows)
            _insert_evidence_images(sheet, request.evidence_images)
            temporary_path = _create_temporary_path(output_dir)
            try:
                workbook.save(temporary_path)
            except OSError:
                raise ValueError(f"报价表无法写入输出目录：{output_dir}") from None
        finally:
            workbook.close()

        published_path = (
            _publish_to_destination(
                temporary_path,
                output_dir,
                request.destination_path,
            )
            if request.destination_path is not None
            else _publish_without_overwrite(
                temporary_path,
                output_dir,
                request.quote_month,
            )
        )
        return published_path
    finally:
        if temporary_path is not None:
            try:
                _remove_temporary_file(temporary_path)
            except ValueError:
                if published_path is None:
                    raise
                warnings.warn(
                    (
                        f"正式报价表已生成：{published_path}；"
                        f"临时文件未能清理：{temporary_path}"
                    ),
                    RuntimeWarning,
                    stacklevel=2,
                )


def _ensure_output_directory(output_dir: Path) -> Path:
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except OSError:
        raise ValueError(f"输出目录无法创建或使用：{output_dir}") from None
    if not output_dir.is_dir():
        raise ValueError(f"输出目录无法创建或使用：{output_dir}")
    return output_dir


def _create_temporary_path(output_dir: Path) -> Path:
    try:
        with NamedTemporaryFile(
            mode="wb",
            prefix=".quote-",
            suffix=".tmp.xlsx",
            dir=output_dir,
            delete=False,
        ) as temporary_file:
            return Path(temporary_file.name)
    except OSError:
        raise ValueError(f"报价表无法写入输出目录：{output_dir}") from None


def _candidate_output_paths(
    output_dir: Path,
    quote_month: QuoteMonth,
) -> Iterator[Path]:
    stem = f"{quote_month.year}年{quote_month.month:02d}月终端供货价报价表"
    yield output_dir / f"{stem}.xlsx"
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    yield output_dir / f"{stem}-{timestamp}.xlsx"
    for suffix in count(2):
        yield output_dir / f"{stem}-{timestamp}-{suffix}.xlsx"


def _publish_without_overwrite(
    temporary_path: Path,
    output_dir: Path,
    quote_month: QuoteMonth,
) -> Path:
    for candidate in _candidate_output_paths(output_dir, quote_month):
        try:
            os.link(temporary_path, candidate)
        except FileExistsError:
            continue
        except OSError:
            raise ValueError(
                f"报价表无法安全发布到输出目录：{output_dir}"
            ) from None
        return candidate
    raise AssertionError("unreachable")


def _publish_to_destination(
    temporary_path: Path,
    output_dir: Path,
    destination_path: Path,
) -> Path:
    destination = Path(destination_path)
    if destination.resolve().parent != output_dir.resolve():
        raise ValueError("报价表指定输出路径必须位于输出目录内")
    try:
        os.replace(temporary_path, destination)
    except OSError:
        raise ValueError(f"报价表无法安全发布到输出目录：{output_dir}") from None
    return destination


def _remove_temporary_file(temporary_path: Path) -> None:
    try:
        temporary_path.unlink(missing_ok=True)
    except OSError:
        raise ValueError(f"报价表临时文件无法清理：{temporary_path}") from None


def _update_month_headers(sheet: Worksheet, quote_month: QuoteMonth) -> None:
    history_month, previous_month = required_price_months(quote_month)
    for column, value in (
        ("I", history_month),
        ("J", previous_month),
        ("K", quote_month),
    ):
        cell = sheet[f"{column}{HEADER_ROW}"]
        if not isinstance(cell.value, str) or _MONTH_TOKEN.search(cell.value) is None:
            raise ValueError(f"报价模板月份表头无效：{cell.coordinate}")
        replacement = month_header(value).split("结算报价", maxsplit=1)[0]
        cell.value = _MONTH_TOKEN.sub(replacement, cell.value, count=1)


def _write_rows(sheet: Worksheet, rows: list[QuoteRow]) -> None:
    sheet.data_validations.dataValidation.clear()
    if not rows:
        if sheet.max_row >= FIRST_DATA_ROW:
            sheet.delete_rows(FIRST_DATA_ROW, sheet.max_row - FIRST_DATA_ROW + 1)
        for row_number in tuple(sheet.row_dimensions):
            if row_number >= FIRST_DATA_ROW:
                del sheet.row_dimensions[row_number]
        return

    last_data_row = FIRST_DATA_ROW + len(rows) - 1
    if sheet.max_row > last_data_row:
        sheet.delete_rows(last_data_row + 1, sheet.max_row - last_data_row)

    for target_row in range(FIRST_DATA_ROW + 1, last_data_row + 1):
        _copy_row_style(sheet, STYLE_SKELETON_ROW, target_row)

    for row_number, quote_row in enumerate(rows, start=FIRST_DATA_ROW):
        for column_number in range(1, LAST_TEMPLATE_COLUMN + 1):
            sheet.cell(row_number, column_number).value = None

        for column, value in quote_row.cells.items():
            if column in _TEMPLATE_COLUMNS and column not in MANUAL_COLUMNS:
                sheet[f"{column}{row_number}"] = value

        sheet[f"C{row_number}"] = normalize_code(quote_row.material_code)
        sheet[f"C{row_number}"].number_format = "@"
        sheet[f"T{row_number}"] = "无"
        sheet[f"U{row_number}"] = "无"
        sheet[f"V{row_number}"] = is_new(
            quote_row.cells.get("J"),
            quote_row.cells.get("F"),
        )
        sheet[f"W{row_number}"] = needs_six_month_adjustment(
            quote_row.cells.get("I")
        )
        for column, formula in formula_cells(row_number).items():
            sheet[f"{column}{row_number}"] = formula
        for column in PERCENTAGE_COLUMNS:
            sheet[f"{column}{row_number}"].number_format = "0.00%"

    validation = DataValidation(
        type="list",
        formula1=N_VALIDATION_FORMULA,
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"N{FIRST_DATA_ROW}:N{last_data_row}")


def _copy_row_style(
    sheet: Worksheet,
    source_row: int,
    target_row: int,
) -> None:
    for column in range(1, LAST_TEMPLATE_COLUMN + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        target._style = copy(source._style)
        if source.has_style:
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)
            target.number_format = source.number_format

    source_dimension = sheet.row_dimensions[source_row]
    target_dimension = sheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed
    target_dimension.thickTop = source_dimension.thickTop
    target_dimension.thickBot = source_dimension.thickBot
    target_dimension._style = copy(source_dimension._style)


def _insert_evidence_images(
    sheet: Worksheet,
    evidence_images: tuple[QuoteEvidenceImage, ...],
) -> None:
    anchors = [evidence.anchor for evidence in evidence_images]
    if len(set(anchors)) != len(anchors):
        raise ValueError("evidence image anchors must be unique")

    for evidence in evidence_images:
        image = OpenpyxlImage(BytesIO(evidence.payload))
        column = re.match(r"[A-Z]+", evidence.anchor)
        row = re.search(r"\d+$", evidence.anchor)
        if column is None or row is None:
            raise AssertionError("validated evidence anchor must be parseable")
        column_letter = column.group()
        row_number = int(row.group())
        width = sheet.column_dimensions[column_letter].width
        width_pixels = _column_width_pixels(float(width or 13))
        height_points = (
            sheet.row_dimensions[row_number].height
            or sheet.sheet_format.defaultRowHeight
            or 15
        )
        height_pixels = points_to_pixels(float(height_points))
        available_width = max(1, width_pixels - 4)
        available_height = max(1, height_pixels - 4)
        if image.width <= 0 or image.height <= 0:
            raise ValueError("evidence image dimensions must be positive")
        scale = min(
            available_width / image.width,
            available_height / image.height,
            1,
        )
        image.width *= scale
        image.height *= scale
        image.anchor = evidence.anchor
        sheet.add_image(image)


def _column_width_pixels(width: float) -> int:
    if width < 1:
        return int(width * 12 + 0.5)
    return int(width * 7 + 5)
