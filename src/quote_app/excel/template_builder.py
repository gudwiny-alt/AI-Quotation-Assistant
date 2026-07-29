from __future__ import annotations

import os
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.workbook.workbook import Workbook  # type: ignore[import-untyped]


TEMPLATE_SHEET_NAME = "5G手机"
HEADER_ROW = 1
STYLE_SKELETON_ROW = 2


def _paths_refer_to_same_file(source: Path, destination: Path) -> bool:
    if source.resolve(strict=False) == destination.resolve(strict=False):
        return True
    try:
        return source.exists() and destination.exists() and os.path.samefile(
            source, destination
        )
    except OSError:
        return False


def build_template(source: Path, destination: Path) -> None:
    """Build a clean quotation template without modifying the source workbook."""
    if _paths_refer_to_same_file(source, destination):
        raise ValueError("源文件和目标文件不能是同一个文件")

    workbook = load_workbook(source, data_only=False)
    try:
        template_sheet = workbook[TEMPLATE_SHEET_NAME]

        for sheet in tuple(workbook.worksheets):
            if sheet is not template_sheet:
                workbook.remove(sheet)

        for cell in template_sheet[STYLE_SKELETON_ROW]:
            cell.value = None
            cell.comment = None
            cell._hyperlink = None

        if template_sheet.max_row > STYLE_SKELETON_ROW:
            template_sheet.delete_rows(
                STYLE_SKELETON_ROW + 1,
                template_sheet.max_row - STYLE_SKELETON_ROW,
            )

        for row_index in tuple(template_sheet.row_dimensions):
            if row_index > STYLE_SKELETON_ROW:
                del template_sheet.row_dimensions[row_index]

        for merged_range in tuple(template_sheet.merged_cells.ranges):
            if merged_range.max_row > STYLE_SKELETON_ROW:
                template_sheet.unmerge_cells(str(merged_range))

        template_sheet._images.clear()

        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
    finally:
        workbook.close()


def load_clean_template(path: Path) -> Workbook:
    """Load a clean quotation template as an editable workbook."""
    return load_workbook(path, data_only=False, keep_links=False)
