from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from itertools import count
import os
from pathlib import Path
import platform as host_platform
import re
from tempfile import NamedTemporaryFile
from typing import Iterator
import warnings

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from quote_app.core.normalization import normalize_text
from quote_app.domain.models import InputPaths, QuoteMonth, QuoteRow
from quote_app.domain.statuses import RowStatus
from quote_app.evidence.models import (
    MacCapturePolicy,
    accepts_capture_validation,
    validate_mac_capture_policy,
)
from quote_app.tasks.models import (
    BusinessOutcome,
    TaskState,
    WebsiteChannel,
    WebsiteObservationCheckpoint,
    WebsiteResult,
    WebsiteTask,
)


SUPPORTED_BRANDS = frozenset(("HONOR", "华为", "维沃", "欧珀", "小米", "苹果", "ZTE中兴"))
CHANNEL_PENDING = "待人工补充"
CHANNEL_WAITING = "待运行"
CHANNEL_MANUAL_VERIFICATION = "等待人工验证"
CHANNEL_UNSUPPORTED = "不支持"
_CHANNEL_ORDER = (
    WebsiteChannel.JD,
    WebsiteChannel.TMALL,
    WebsiteChannel.OFFICIAL,
)
_LEGAL_NO_LABELS = {
    BusinessOutcome.NO_MODEL: "无（无该机型）",
    BusinessOutcome.CAPACITY_UNAVAILABLE: "无（容量不可用）",
    BusinessOutcome.COLOR_UNAVAILABLE: "无（颜色不可用）",
    BusinessOutcome.SOLD_OUT: "无（已售罄）",
}
STATUS_LABELS = {
    RowStatus.COMPLETED: "完成",
    RowStatus.PARTIAL: "部分完成",
    RowStatus.FAILED: "失败",
    RowStatus.UNSUPPORTED: "不支持",
}
STATUS_FILLS = {
    RowStatus.COMPLETED: PatternFill("solid", fgColor="70AD47"),
    RowStatus.PARTIAL: PatternFill("solid", fgColor="F4B183"),
    RowStatus.FAILED: PatternFill("solid", fgColor="F8696B"),
    RowStatus.UNSUPPORTED: PatternFill("solid", fgColor="F4B183"),
}
DETAIL_HEADERS = (
    "产品经理",
    "品牌",
    "集团一级库物料编码",
    "产品名称",
    "配置",
    "基础表关联状态",
    "营销表关联状态",
    "BOP关联状态",
    "京东渠道状态",
    "天猫渠道状态",
    "官网渠道状态",
    "最终状态",
    "失败步骤",
    "原因",
    "建议操作",
)
_HARD_ASSOCIATION_ISSUES = frozenset(("MARKETING_NOT_FOUND", "MARKETING_CONFLICT"))


@dataclass(frozen=True, slots=True)
class RunSummary:
    total_rows: int
    completed_rows: int
    partial_rows: int
    failed_rows: int
    unsupported_rows: int
    manual_supplement_rows: int


@dataclass(frozen=True, slots=True)
class ReportWriteRequest:
    quote_month: QuoteMonth
    rows: list[QuoteRow]
    output_dir: Path
    input_paths: InputPaths | None = None
    quote_path: Path | None = None
    run_at: datetime | None = None
    website_tasks: tuple[WebsiteTask, ...] = ()
    website_results: tuple[WebsiteResult, ...] = ()
    website_observations: tuple[WebsiteObservationCheckpoint, ...] = ()
    waiting_task_ids: frozenset[str] = frozenset()
    website_run: bool = False
    capture_acceptance_policy: MacCapturePolicy = MacCapturePolicy.STRICT

    def __post_init__(self) -> None:
        validate_mac_capture_policy(
            self.capture_acceptance_policy,
            platform_name=host_platform.system(),
        )


@dataclass(frozen=True, slots=True)
class RowAssessment:
    row: QuoteRow
    status: RowStatus
    manual_supplement: bool
    manager: str
    brand: str
    base_state: str
    marketing_state: str
    bop_state: str
    channel_states: tuple[str, str, str]
    failed_step: str
    reason: str
    recommendation: str


@dataclass(frozen=True, slots=True)
class _WebsiteReportContext:
    active: bool
    by_row: dict[
        int,
        dict[
            WebsiteChannel,
            tuple[WebsiteResult | None, WebsiteObservationCheckpoint | None],
        ],
    ]
    waiting_channels: dict[int, frozenset[WebsiteChannel]]
    capture_acceptance_policy: MacCapturePolicy


def assess_rows(
    rows: list[QuoteRow],
    website_tasks: tuple[WebsiteTask, ...] = (),
    website_results: tuple[WebsiteResult, ...] = (),
    website_observations: tuple[WebsiteObservationCheckpoint, ...] = (),
    waiting_task_ids: frozenset[str] = frozenset(),
    website_run: bool = False,
    *,
    capture_acceptance_policy: MacCapturePolicy = MacCapturePolicy.STRICT,
) -> list[RowAssessment]:
    """Classify core-only rows using one explicit source of report truth."""
    context = _website_report_context(
        website_tasks,
        website_results,
        website_observations,
        waiting_task_ids,
        website_run,
        capture_acceptance_policy=capture_acceptance_policy,
    )
    return [
        _assess_row(row, output_row_number, context)
        for output_row_number, row in enumerate(rows, start=2)
    ]


def summarize_assessments(assessments: list[RowAssessment]) -> RunSummary:
    status_counts = Counter(assessment.status for assessment in assessments)
    return RunSummary(
        total_rows=len(assessments),
        completed_rows=status_counts[RowStatus.COMPLETED],
        partial_rows=status_counts[RowStatus.PARTIAL],
        failed_rows=status_counts[RowStatus.FAILED],
        unsupported_rows=status_counts[RowStatus.UNSUPPORTED],
        manual_supplement_rows=sum(
            assessment.manual_supplement for assessment in assessments
        ),
    )


def summarize_rows(rows: list[QuoteRow]) -> RunSummary:
    return summarize_assessments(assess_rows(rows))


def write_execution_report(request: ReportWriteRequest) -> Path:
    """Write a collision-safe report without modifying any source workbook."""
    output_dir = _ensure_output_directory(Path(request.output_dir))
    assessments = assess_rows(
        request.rows,
        request.website_tasks,
        request.website_results,
        request.website_observations,
        request.waiting_task_ids,
        request.website_run,
        capture_acceptance_policy=request.capture_acceptance_policy,
    )
    summary = summarize_assessments(assessments)
    workbook = Workbook()
    temporary_path: Path | None = None
    published_path: Path | None = None
    try:
        overview = workbook.active
        overview.title = "运行总览"
        detail = workbook.create_sheet("处理明细")
        _write_overview(
            overview,
            request,
            assessments,
            summary,
            request.run_at or datetime.now(),
        )
        _write_detail(detail, assessments)
        temporary_path = _create_temporary_path(output_dir)
        try:
            workbook.save(temporary_path)
        except OSError:
            raise ValueError(f"执行报告无法写入输出目录：{output_dir}") from None
        workbook.close()
        published_path = _publish_without_overwrite(
            temporary_path,
            output_dir,
            request.quote_month,
        )
        return published_path
    finally:
        workbook.close()
        if temporary_path is not None:
            try:
                _remove_temporary_file(temporary_path)
            except ValueError:
                if published_path is None:
                    raise
                warnings.warn(
                    (
                        f"正式执行报告已生成：{published_path}；"
                        f"临时文件未能清理：{temporary_path}"
                    ),
                    RuntimeWarning,
                    stacklevel=2,
                )


def _assess_row(
    row: QuoteRow,
    output_row_number: int,
    website_context: _WebsiteReportContext,
) -> RowAssessment:
    issue_codes = {issue.code for issue in row.issues}
    brand = normalize_text(row.cells.get("B")) or "未识别"
    manager = normalize_text(row.cells.get("AG")) or "未分配"
    issue_reason = "；".join(f"{issue.code}：{issue.message}" for issue in row.issues)
    channel_states = _channel_states(
        row,
        output_row_number,
        brand,
        website_context,
    )

    if issue_codes & _HARD_ASSOCIATION_ISSUES:
        status = RowStatus.FAILED
        manual_supplement = False
        failed_step = "营销商品信息关联"
        reason = issue_reason
        recommendation = "核对营销商品信息查询表中的物料编码及重复记录后重新运行"
    elif "WEB_FIELDS_MISSING" in issue_codes:
        status = RowStatus.PARTIAL
        manual_supplement = True
        failed_step = "营销商品信息完整性"
        reason = issue_reason
        recommendation = "补齐营销商品信息中的品牌、型号、内存、存储和颜色字段后重新运行"
    elif brand not in SUPPORTED_BRANDS:
        status = RowStatus.UNSUPPORTED
        manual_supplement = True
        failed_step = "网站渠道自动处理"
        reason = (
            f"品牌“{brand}”不在首版支持的7个品牌范围内"
            + (f"；{issue_reason}" if issue_reason else "")
        )
        recommendation = "人工补充 AI:AN 的价格、链接和截图，并反馈新增品牌规则"
    elif website_context.active:
        completed_channels = sum(
            _channel_is_completed(state) for state in channel_states
        )
        technical_channels = sum(
            _channel_is_technical(state) for state in channel_states
        )
        channel_reason = "；".join(
            f"{channel}：{state}"
            for channel, state in zip(
                ("京东", "天猫", "官网"),
                channel_states,
                strict=True,
            )
        )
        if CHANNEL_MANUAL_VERIFICATION in channel_states:
            status = RowStatus.PARTIAL
            manual_supplement = True
            failed_step = "网站人工验证"
            reason = "；".join(
                part for part in (issue_reason, channel_reason) if part
            )
            recommendation = "完成人工验证后继续当前任务"
        elif completed_channels == 3 and not issue_codes:
            status = RowStatus.COMPLETED
            manual_supplement = False
            failed_step = ""
            reason = "核心关联及三个网站渠道均已完成"
            recommendation = "无需操作"
        elif technical_channels == 3 and completed_channels == 0:
            status = RowStatus.FAILED
            manual_supplement = False
            failed_step = "网站渠道自动处理"
            reason = "；".join(
                part for part in (issue_reason, channel_reason) if part
            )
            recommendation = "检查三个渠道的稳定错误代码后重试失败任务"
        else:
            status = RowStatus.PARTIAL
            manual_supplement = True
            failed_step = "网站渠道自动处理"
            reason = "；".join(
                part for part in (issue_reason, channel_reason) if part
            )
            recommendation = "重试失败或待运行渠道；仍未完成时人工补充 AI:AN"
    elif all(state == "已完成" for state in channel_states) and not issue_codes:
        status = RowStatus.COMPLETED
        manual_supplement = False
        failed_step = ""
        reason = "核心关联及三个网站渠道均已完成"
        recommendation = "无需操作"
    else:
        status = RowStatus.PARTIAL
        manual_supplement = True
        failed_step = "网站渠道自动处理"
        reason = issue_reason or "核心表格关联已完成，网站价格及截图尚未处理"
        recommendation = "人工补充 AI:AN，或在后续浏览器自动化版本中继续处理"

    return RowAssessment(
        row=row,
        status=status,
        manual_supplement=manual_supplement,
        manager=manager,
        brand=brand,
        base_state="已关联",
        marketing_state=_marketing_state(issue_codes),
        bop_state=_bop_state(row, issue_codes),
        channel_states=channel_states,
        failed_step=failed_step,
        reason=reason,
        recommendation=recommendation,
    )


def _marketing_state(issue_codes: set[str]) -> str:
    if "MARKETING_NOT_FOUND" in issue_codes:
        return "未找到"
    if "MARKETING_CONFLICT" in issue_codes:
        return "记录冲突"
    if "WEB_FIELDS_MISSING" in issue_codes:
        return "已关联（网站字段不完整）"
    return "已关联"


def _bop_state(row: QuoteRow, issue_codes: set[str]) -> str:
    if "BOP_CONFLICT" in issue_codes:
        return "记录冲突"
    if row.cells.get("F") == "申请配置" and row.cells.get("G") == "申请配置":
        return "无需匹配（申请配置）"
    return "已关联"


def _channel_states(
    row: QuoteRow,
    output_row_number: int,
    brand: str,
    website_context: _WebsiteReportContext,
) -> tuple[str, str, str]:
    if website_context.active:
        if brand not in SUPPORTED_BRANDS:
            return (CHANNEL_UNSUPPORTED,) * 3
        channel_observations = website_context.by_row.get(
            output_row_number,
            {},
        )
        states = tuple(
            _website_channel_state(
                *(channel_observations.get(channel, (None, None))),
                waiting_for_manual_verification=(
                    channel
                    in website_context.waiting_channels.get(
                        output_row_number,
                        frozenset(),
                    )
                ),
                capture_acceptance_policy=website_context.capture_acceptance_policy,
            )
            for channel in _CHANNEL_ORDER
        )
        return states[0], states[1], states[2]
    return (
        _channel_state(row, "AI", "AL"),
        _channel_state(row, "AJ", "AM"),
        _channel_state(row, "AK", "AN"),
    )


def _channel_state(row: QuoteRow, price_column: str, evidence_column: str) -> str:
    if normalize_text(row.cells.get(price_column)) and normalize_text(
        row.cells.get(evidence_column)
    ):
        return "已完成"
    return CHANNEL_PENDING


def _website_report_context(
    tasks: tuple[WebsiteTask, ...],
    results: tuple[WebsiteResult, ...],
    observations: tuple[WebsiteObservationCheckpoint, ...],
    waiting_task_ids: frozenset[str],
    website_run: bool,
    *,
    capture_acceptance_policy: MacCapturePolicy,
) -> _WebsiteReportContext:
    validate_mac_capture_policy(
        capture_acceptance_policy,
        platform_name=host_platform.system(),
    )
    if not website_run and not tasks and not results and not observations:
        return _WebsiteReportContext(
            active=False,
            by_row={},
            waiting_channels={},
            capture_acceptance_policy=capture_acceptance_policy,
        )
    results_by_task: dict[str, WebsiteResult] = {}
    for result in results:
        if result.task_id in results_by_task:
            raise ValueError("website results contain a duplicate task_id")
        results_by_task[result.task_id] = result
    observations_by_task: dict[str, WebsiteObservationCheckpoint] = {}
    for observation in observations:
        if observation.task_id in observations_by_task:
            raise ValueError("website observations contain a duplicate task_id")
        observations_by_task[observation.task_id] = observation

    by_row: dict[
        int,
        dict[
            WebsiteChannel,
            tuple[WebsiteResult | None, WebsiteObservationCheckpoint | None],
        ],
    ] = {}
    waiting_channels: dict[int, set[WebsiteChannel]] = {}
    task_ids: set[str] = set()
    for task in tasks:
        if task.task_id in task_ids:
            raise ValueError("website tasks contain a duplicate task_id")
        task_ids.add(task.task_id)
        row_channels = by_row.setdefault(task.output_row_number, {})
        if task.channel in row_channels:
            raise ValueError("website tasks contain a duplicate row/channel")
        row_channels[task.channel] = (
            results_by_task.get(task.task_id),
            observations_by_task.get(task.task_id),
        )
        if task.task_id in waiting_task_ids:
            waiting_channels.setdefault(task.output_row_number, set()).add(task.channel)

    unknown_results = set(results_by_task) - task_ids
    if unknown_results:
        raise ValueError("website result does not belong to a supplied task")
    unknown_observations = set(observations_by_task) - task_ids
    if unknown_observations:
        raise ValueError("website observation does not belong to a supplied task")
    unknown_waiting = set(waiting_task_ids) - task_ids
    if unknown_waiting:
        raise ValueError("waiting website task does not belong to a supplied task")
    return _WebsiteReportContext(
        active=True,
        by_row=by_row,
        waiting_channels={
            row_number: frozenset(channels)
            for row_number, channels in waiting_channels.items()
        },
        capture_acceptance_policy=capture_acceptance_policy,
    )


def _website_channel_state(
    result: WebsiteResult | None,
    observation: WebsiteObservationCheckpoint | None = None,
    *,
    waiting_for_manual_verification: bool = False,
    capture_acceptance_policy: MacCapturePolicy,
) -> str:
    if waiting_for_manual_verification:
        return CHANNEL_MANUAL_VERIFICATION
    if result is not None and result.state is TaskState.SUCCEEDED:
        evidence = result.evidence
        if evidence is None or not accepts_capture_validation(
            evidence.validation_code,
            policy=capture_acceptance_policy,
            platform_name=host_platform.system(),
        ):
            raise ValueError("capture validation is not accepted for this report")
        if result.outcome is BusinessOutcome.PRICE_FOUND:
            if evidence.validation_code == "CAPTURE_OK_MAC_VISUAL_REVIEW":
                return f"价格成功（{result.price}）；截图成功（Mac 视觉复核）"
            return f"价格成功（{result.price}）；截图成功"
        if result.outcome in _LEGAL_NO_LABELS:
            return _LEGAL_NO_LABELS[result.outcome]
        raise ValueError("website result has an unsupported completed outcome")
    if observation is not None:
        code = (
            normalize_text(result.error_code)
            if result is not None and result.state is TaskState.TECHNICAL_FAILURE
            else "CAPTURE_PENDING"
        )
        if observation.outcome is BusinessOutcome.PRICE_FOUND:
            return f"价格成功（{observation.price}）；截图待补（{code}）"
        if observation.outcome in _LEGAL_NO_LABELS:
            return f"{_LEGAL_NO_LABELS[observation.outcome]}；截图待补（{code}）"
        raise ValueError("website observation has an unsupported outcome")
    if result is None:
        return CHANNEL_WAITING
    if result.state is TaskState.TECHNICAL_FAILURE:
        code = normalize_text(result.error_code)
        message = _concise_diagnostic(result.error_message)
        return f"技术失败（{code}：{message}）"


def _concise_diagnostic(message: str | None) -> str:
    normalized = normalize_text(message)
    if "<" in normalized or ">" in normalized:
        return "页面诊断已省略"
    return re.sub(r"\s+", " ", normalized)[:160]


def _channel_is_completed(state: str) -> bool:
    return (
        state == "已完成"
        or state.startswith("价格成功（")
        or state.startswith("无（")
    )


def _channel_is_technical(state: str) -> bool:
    return state == "失败" or state.startswith("技术失败（")


def _write_overview(
    sheet: Worksheet,
    request: ReportWriteRequest,
    assessments: list[RowAssessment],
    summary: RunSummary,
    run_at: datetime,
) -> None:
    sheet.append(["资金物流平台铺货报价执行报告"])
    sheet.append(["报价月份", f"{request.quote_month.year}年{request.quote_month.month:02d}月"])
    sheet.append(["运行时间", run_at.strftime("%Y-%m-%d %H:%M:%S")])
    for label, path in _metadata_rows(request):
        sheet.append([label, path.name if path is not None else "未提供"])
    sheet.append([])
    sheet.append(["汇总指标", "数量"])
    for label, value in (
        ("报价总行数", summary.total_rows),
        ("处理完成", summary.completed_rows),
        ("部分完成", summary.partial_rows),
        ("处理失败", summary.failed_rows),
        ("不支持品牌", summary.unsupported_rows),
        ("待人工补充", summary.manual_supplement_rows),
    ):
        sheet.append([label, value])

    sheet.append([])
    sheet.append(["品牌汇总"])
    sheet.append(["品牌", "总行数", "完成", "部分完成", "失败", "不支持", "待人工补充"])
    _append_breakdown(sheet, assessments, "brand")

    sheet.append([])
    sheet.append(["产品经理汇总"])
    sheet.append(
        ["产品经理", "总行数", "完成", "部分完成", "失败", "不支持", "待人工补充"]
    )
    _append_breakdown(sheet, assessments, "manager")

    sheet.append([])
    sheet.append(["渠道状态汇总"])
    sheet.append(["渠道", "已完成", "待人工补充", "失败"])
    for channel_index, channel in enumerate(("京东", "天猫", "官网")):
        states = Counter(
            _channel_summary_bucket(
                assessment.channel_states[channel_index]
            )
            for assessment in assessments
        )
        sheet.append(
            [
                channel,
                states["已完成"],
                states[CHANNEL_PENDING],
                states["失败"],
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 72
    for column in ("C", "D", "E", "F", "G"):
        sheet.column_dimensions[column].width = 14
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="4472C4")
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row, 1).value in (
            "汇总指标",
            "品牌",
            "产品经理",
            "渠道",
        ):
            for cell in sheet[row]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="5B9BD5")


def _metadata_rows(request: ReportWriteRequest) -> tuple[tuple[str, Path | None], ...]:
    inputs = request.input_paths
    return (
        ("基础表", inputs.base if inputs else None),
        ("营销商品信息查询表", inputs.marketing if inputs else None),
        ("BOP资源信息表", inputs.bop if inputs else None),
        ("报价表输出", request.quote_path),
    )


def _channel_summary_bucket(state: str) -> str:
    if _channel_is_completed(state):
        return "已完成"
    if _channel_is_technical(state):
        return "失败"
    return CHANNEL_PENDING


def _append_breakdown(
    sheet: Worksheet,
    assessments: list[RowAssessment],
    field: str,
) -> None:
    groups: defaultdict[str, list[RowAssessment]] = defaultdict(list)
    for assessment in assessments:
        groups[getattr(assessment, field)].append(assessment)
    for name in sorted(groups):
        group = groups[name]
        counts = Counter(item.status for item in group)
        sheet.append(
            [
                name,
                len(group),
                counts[RowStatus.COMPLETED],
                counts[RowStatus.PARTIAL],
                counts[RowStatus.FAILED],
                counts[RowStatus.UNSUPPORTED],
                sum(item.manual_supplement for item in group),
            ]
        )


def _write_detail(sheet: Worksheet, assessments: list[RowAssessment]) -> None:
    sheet.append(list(DETAIL_HEADERS))
    for assessment in assessments:
        row = assessment.row
        sheet.append(
            [
                assessment.manager,
                assessment.brand,
                row.material_code,
                row.cells.get("E"),
                row.cells.get("G"),
                assessment.base_state,
                assessment.marketing_state,
                assessment.bop_state,
                *assessment.channel_states,
                STATUS_LABELS[assessment.status],
                assessment.failed_step,
                assessment.reason,
                assessment.recommendation,
            ]
        )
        status_cell = sheet.cell(sheet.max_row, 12)
        status_cell.fill = STATUS_FILLS[assessment.status]
        sheet.cell(sheet.max_row, 3).number_format = "@"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:O{max(1, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    widths = (16, 14, 22, 24, 18, 20, 24, 24, 16, 16, 16, 14, 22, 52, 52)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _ensure_output_directory(output_dir: Path) -> Path:
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except OSError:
        raise ValueError(f"输出目录无法创建或使用：{output_dir}") from None
    if not output_dir.is_dir():
        raise ValueError(f"输出目录无法创建或使用：{output_dir}")
    return output_dir


def _create_temporary_path(output_dir: Path) -> Path:
    try:
        with NamedTemporaryFile(
            mode="wb",
            prefix=".report-",
            suffix=".tmp.xlsx",
            dir=output_dir,
            delete=False,
        ) as temporary_file:
            return Path(temporary_file.name)
    except OSError:
        raise ValueError(f"执行报告无法写入输出目录：{output_dir}") from None


def _candidate_output_paths(
    output_dir: Path,
    quote_month: QuoteMonth,
) -> Iterator[Path]:
    stem = f"{quote_month.year}年{quote_month.month:02d}月报价执行报告"
    yield output_dir / f"{stem}.xlsx"
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    yield output_dir / f"{stem}-{timestamp}.xlsx"
    for suffix in count(2):
        yield output_dir / f"{stem}-{timestamp}-{suffix}.xlsx"


def _publish_without_overwrite(
    temporary_path: Path,
    output_dir: Path,
    quote_month: QuoteMonth,
) -> Path:
    for candidate in _candidate_output_paths(output_dir, quote_month):
        try:
            os.link(temporary_path, candidate)
        except FileExistsError:
            continue
        except OSError:
            raise ValueError(f"执行报告无法安全发布到输出目录：{output_dir}") from None
        return candidate
    raise AssertionError("unreachable")


def _remove_temporary_file(temporary_path: Path) -> None:
    try:
        temporary_path.unlink(missing_ok=True)
    except OSError:
        raise ValueError(f"执行报告临时文件无法清理：{temporary_path}") from None
