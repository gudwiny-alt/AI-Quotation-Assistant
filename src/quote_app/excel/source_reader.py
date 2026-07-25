from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]


class WorkbookReadError(Exception):
    """A source workbook could not be read as a table."""


@dataclass(frozen=True, slots=True)
class SheetTable:
    path: Path
    sheet_name: str
    headers: dict[str, int]
    header_values: tuple[str, ...]
    column_count: int
    rows: tuple[tuple[object, ...], ...]

    def records_by_column_letter(self) -> tuple[dict[str, object], ...]:
        return tuple(
            {
                get_column_letter(index): value
                for index, value in enumerate(row, start=1)
            }
            for row in self.rows
        )


def read_first_table(path: Path) -> SheetTable:
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        if not workbook.sheetnames:
            raise ValueError("workbook has no worksheets")
        sheet = workbook[workbook.sheetnames[0]]
        values = tuple(tuple(row) for row in sheet.iter_rows(values_only=True))
        if not values or not any(value is not None for value in values[0]):
            raise ValueError("first worksheet has no header row")
        header_values = tuple("" if value is None else str(value) for value in values[0])
        headers = {
            value: index
            for index, value in enumerate(header_values)
            if value
        }
        return SheetTable(
            path=path,
            sheet_name=sheet.title,
            headers=headers,
            header_values=header_values,
            column_count=len(header_values),
            rows=values[1:],
        )
    except WorkbookReadError:
        raise
    except Exception as error:
        raise WorkbookReadError(f"could not read workbook {path.name}") from error
    finally:
        if workbook is not None:
            workbook.close()
