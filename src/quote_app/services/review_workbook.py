"""Conservative OOXML writeback: only six whitelisted cells may change.

Unlike a load/save round-trip this preserves every other ZIP member verbatim,
including drawings, cached formulas, relationships and application extensions.
"""

from __future__ import annotations
import hashlib
import os
import re
import tempfile
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from decimal import Decimal, InvalidOperation
from datetime import datetime
from openpyxl import load_workbook

COLUMNS = ("K", "L", "M", "P", "Q", "AO")
HEADERS = {
    "C": "集团一级库物料编码",
    "L": "终端公司采购价",
    "M": "终端公司全国采购系统均价",
    "P": "渠道买断价",
    "Q": "分销零售价",
    "AO": "备注",
}
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

_DIGEST_CACHE = {}


def digest(path: Path | None, *, fresh=False) -> str:
    if path is None:
        return "missing"
    try:
        stat = path.stat()
        signature = (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns)
        key = str(path.resolve())
        if not fresh and key in _DIGEST_CACHE and _DIGEST_CACHE[key][0] == signature:
            return _DIGEST_CACHE[key][1]
        value = hashlib.sha256(path.read_bytes()).hexdigest()
        _DIGEST_CACHE[key] = (signature, value)
        return value
    except OSError:
        return "missing"


def money(value) -> Decimal | None:
    try:
        raw = str(value).strip()
        if len(raw) > 1000:
            return None
        d = Decimal(raw)
        return d if d.is_finite() and d > 0 and -300 <= d.adjusted() <= 300 else None
    except (InvalidOperation, ValueError, TypeError):
        return None


def inspect(path, month):
    if path is None or not path.is_file():
        raise ValueError("尚无已生成的报价工作簿")
    try:
        w = load_workbook(path, read_only=True, data_only=False)
        try:
            if "5G手机" not in w.sheetnames:
                raise ValueError("报价模板页签不匹配：缺少5G手机")
            s = w["5G手机"]
            for col, title in HEADERS.items():
                if s[f"{col}1"].value != title:
                    raise ValueError(f"报价模板列不匹配：{col}应为{title}")
            if s["K1"].value != f"{month.year}年{month.month}月结算报价（元/台）":
                raise ValueError("工作簿报价月份与当前任务不一致")
            return {
                number: {c.column_letter: c.value for c in row if c.value is not None}
                for number, row in enumerate(s.iter_rows(min_row=2), start=2)
            }
        finally:
            w.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"报价工作簿无法核验：{exc}") from exc


def atomic_json(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".review-", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as out:
            out.write(text)
            out.flush()
            os.fsync(out.fileno())
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def make_backup(path, expected):
    if digest(path, fresh=True) != expected:
        raise ValueError("报价工作簿版本在保存期间变化，已取消写回")
    folder = path.parent / "历史备份"
    folder.mkdir(exist_ok=True)
    backup = folder / (path.name + "." + datetime.now().strftime("%Y%m%d-%H%M%S-%f") + ".bak")
    with backup.open("xb") as out:
        out.write(path.read_bytes())
        out.flush()
        os.fsync(out.fileno())
    if digest(path, fresh=True) != expected:
        raise ValueError("报价工作簿版本在备份期间变化，已取消写回")
    return backup


def write_cells(path: Path, expected: str, month, product, original_identity):
    if digest(path, fresh=True) != expected:
        raise ValueError("报价工作簿版本已变化，请重新打开本批次后核验")
    rows = inspect(path, month)
    row = rows.get(product.output_row, {})
    current = tuple(str(row.get(c, "") or "") for c in ("C", "D", "E", "F", "G"))
    if current != original_identity or str(row.get("C", "")) != product.material_code:
        raise ValueError("商品行定位不一致，已取消写回以避免覆盖其他商品")
    numeric = {}
    for col in COLUMNS[:-1]:
        raw = product.values.get(col, "").strip()
        if raw and money(raw) is None:
            raise ValueError(f"{col}请输入大于0的有效金额，或留空保存草稿")
        numeric[col] = str(money(raw)) if raw else None
    unchanged = all(
        (money(row.get(col)) == money(numeric[col]) if numeric[col] else row.get(col) in (None, ""))
        for col in COLUMNS[:-1]
    ) and str(row.get("AO") or "") == product.values.get("AO", "")
    if unchanged:
        if digest(path, fresh=True) != expected:
            raise ValueError("报价工作簿版本已变化，请重新打开本批次")
        return None
    with ZipFile(path) as z:
        # Resolve worksheet by workbook relationships; never assume sheet1.
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        sheet = next(
            e for e in wb.findall(f"{{{NS}}}sheets/{{{NS}}}sheet") if e.attrib["name"] == "5G手机"
        )
        rid = sheet.attrib[
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        ]
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        target = next(e.attrib["Target"] for e in rels if e.attrib["Id"] == rid)
        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        xml = z.read(part).decode("utf-8")
        row_pattern = rf'(<row\b[^>]*\br="{product.output_row}"[^>]*>)(.*?)(</row>)'
        match = re.search(row_pattern, xml, re.S)
        if match is None:
            raise ValueError("商品行结构无法安全写回")
        body = match.group(2)
        for col in COLUMNS:
            address = f"{col}{product.output_row}"
            pattern = rf'<c\b[^>]*\br="{address}"(?=[\s/>])[^>]*?(?:/>|>.*?</c>)'
            old = re.search(pattern, body, re.S)
            style = re.search(r'\bs="[^"]*"', old.group(0)) if old else None
            cell = ET.Element("c", {"r": address})
            if style:
                cell.set("s", style.group(0)[3:-1])
            value = product.values.get(col, "") if col == "AO" else numeric[col]
            if value is not None and value != "":
                if col == "AO":
                    cell.set("t", "inlineStr")
                    inline = ET.SubElement(cell, "is")
                    text = ET.SubElement(inline, "t")
                    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    text.text = value
                else:
                    ET.SubElement(cell, "v").text = value
            replacement = ET.tostring(cell, encoding="unicode")
            if old:
                body = body[: old.start()] + replacement + body[old.end() :]
            else:
                # Cell order is significant to some spreadsheet readers.
                def number(name):
                    n = 0
                    for ch in name:
                        n = n * 26 + ord(ch) - 64
                    return n

                following = next(
                    (
                        m
                        for m in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', body)
                        if number(m.group(1)) > number(col)
                    ),
                    None,
                )
                offset = following.start() if following else len(body)
                body = body[:offset] + replacement + body[offset:]
        changed = xml[: match.start(2)] + body + xml[match.end(2) :]
        ET.fromstring(changed)
        fd, name = tempfile.mkstemp(suffix=".xlsx", prefix=".quote-review-", dir=path.parent)
        os.close(fd)
        try:
            with ZipFile(name, "w") as out:
                out.comment = z.comment
                for item in z.infolist():
                    out.writestr(
                        item, changed.encode() if item.filename == part else z.read(item.filename)
                    )
            verified = inspect(Path(name), month)[product.output_row]
            for col in COLUMNS:
                expected_value = product.values.get(col, "")
                actual = verified.get(col, "")
                if col == "AO":
                    if str(actual or "") != expected_value:
                        raise ValueError("备注写回校验未通过")
                elif expected_value.strip() and money(actual) != money(expected_value):
                    raise ValueError(f"{col}写回值精度无法保持")
            if digest(path, fresh=True) != expected:
                raise ValueError("报价工作簿版本在保存期间变化，已取消写回")
            backup = make_backup(path, expected)
            os.replace(name, path)
            return backup
        finally:
            if os.path.exists(name):
                os.unlink(name)


def rollback(path, backup, expected):
    """Restore exact previous bytes if no third party changed the new file."""
    if backup is None:
        return
    if digest(path, fresh=True) != expected:
        raise OSError(f"保存复核记录失败且工作簿又被外部修改；请从备份恢复：{backup}")
    fd, name = tempfile.mkstemp(prefix=".review-rollback-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as out:
            out.write(backup.read_bytes())
            out.flush()
            os.fsync(out.fileno())
        if digest(path, fresh=True) != expected:
            raise OSError(f"恢复前工作簿版本变化，请使用备份：{backup}")
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)
