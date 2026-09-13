"""Read-only recovery of generated prices/evidence and exact-key source facts."""

from collections import defaultdict
from dataclasses import replace
from contextlib import closing
from datetime import date, datetime
from functools import lru_cache
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import tempfile

from openpyxl import load_workbook

from quote_app.core.normalization import normalize_code
from quote_app.desktop_state import TaskRow, _readonly_connection, read_task_rows
from quote_app.services.review_workbook import digest, money

CHANNEL_COLUMNS = {"jd": ("AI", 37), "tmall": ("AJ", 38), "official": ("AK", 39)}


def warehouse_limit(value):
    """User-approved: the upper end of an explicit suggested retail range."""
    direct = money(value)
    if direct is not None:
        return direct
    match = re.fullmatch(
        r"\s*[¥￥]?\s*(\d+(?:\.\d+)?)\s*(?:到|至|[-~～—])\s*[¥￥]?\s*(\d+(?:\.\d+)?)\s*元?\s*",
        str(value or ""),
    )
    if not match:
        return None
    lower, upper = money(match[1]), money(match[2])
    return upper if lower is not None and upper is not None and lower <= upper else None


def workbook_channels(path, rows):
    """Anchor each embedded image to its exact output row and AL/AM/AN column."""
    evidence = {}
    book = load_workbook(path, data_only=False)
    try:
        for picture in book["5G手机"]._images:
            anchor = getattr(picture.anchor, "_from", None)
            if anchor is None or anchor.col not in (37, 38, 39):
                continue
            row = anchor.row + 1
            if row not in rows:
                continue
            data = picture._data()
            key = hashlib.sha256(data).hexdigest()
            folder = Path(tempfile.gettempdir()) / "quotation-review-evidence"
            folder.mkdir(exist_ok=True)
            destination = folder / (key + "." + picture.format)
            if (
                not destination.exists()
                or hashlib.sha256(destination.read_bytes()).hexdigest() != key
            ):
                destination.write_bytes(data)
            evidence.setdefault((row, anchor.col), []).append(destination)
    finally:
        book.close()
    result = []
    for number, cells in rows.items():
        for channel, (col, image_col) in CHANNEL_COLUMNS.items():
            value = cells.get(col)
            images = evidence.get((number, image_col), [])
            if value is None and not images:
                continue
            price = money(value)
            # A workbook observation is not a newly verified live channel result.
            result.append(
                TaskRow(
                    f"workbook:{number}:{channel}",
                    str(cells.get("E") or cells.get("D") or ""),
                    channel,
                    price=str(price) if price is not None else "",
                    outcome="price_found" if price is not None else "",
                    state="workbook",
                    evidence_state="complete" if len(images) == 1 else "missing",
                    evidence_path=images[0] if len(images) == 1 else None,
                    error="WORKBOOK_EVIDENCE_AMBIGUOUS" if len(images) > 1 else "",
                    source_row_number=number,
                    material_code=str(cells.get("C", "")),
                )
            )
    return result


@lru_cache(maxsize=128)
def _thumbnail_signature(path, version, high_detail):
    from quote_app.services.web_to_excel import _excel_thumbnail

    return hashlib.sha256(
        _excel_thumbnail(Path(path).read_bytes(), high_detail=high_detail)
    ).hexdigest()


def recover_task_provenance(observations, database, *, run_id=None):
    """Join only exact image bytes + material + channel + price, never a model-name guess."""
    if not database or not Path(database).is_file():
        return observations
    try:
        runs = [{"run_id": run_id}] if run_id else []
        candidates = {}
        for run in runs:
            tasks, error = read_task_rows(Path(database), run["run_id"])
            if error:
                continue
            for task in tasks:
                signature = digest(task.evidence_path)
                if signature != "missing":
                    thumb = _thumbnail_signature(
                        str(task.evidence_path), signature, task.outcome == "no_model"
                    )
                    for image_hash in {signature, thumb}:
                        key = (task.material_code, task.channel, money(task.price), image_hash)
                        candidates.setdefault(key, []).append(task)
        result = []
        matched = 0
        for observation in observations:
            key = (
                observation.material_code,
                observation.channel,
                money(observation.price),
                digest(observation.evidence_path),
            )
            found = candidates.get(key, [])
            # Ambiguous provenance is left as a workbook observation.
            if len(found) == 1:
                result.append(replace(found[0], source_row_number=observation.source_row_number))
                matched += 1
            else:
                result.append(observation)
        if observations and matched == len(observations) and len(runs) == 1:
            # All embedded images prove this batch. Keep its explicit failures too,
            # so one missing screenshot has a concrete cause instead of disappearing.
            row_numbers = {r.material_code: r.source_row_number for r in observations}
            present = {(r.material_code, r.channel) for r in result}
            for task in tasks:
                if (
                    (task.material_code, task.channel) not in present
                    and task.material_code in row_numbers
                    and task.state == "technical_failure"
                ):
                    result.append(replace(task, source_row_number=row_numbers[task.material_code]))
        return result
    except (OSError, sqlite3.Error, ValueError, KeyError, TypeError):
        return observations


def matching_run_paths(path, month, rows, database, *, run_ids=None):
    """Recover input references only; never attach another run's channel results."""
    if not database or not Path(database).is_file():
        return {}
    try:
        with closing(_readonly_connection(Path(database))) as connection:
            candidates = connection.execute(
                "SELECT run_id, payload_json FROM quotation_runs ORDER BY created_at DESC LIMIT 100"
            ).fetchall()
        actual = [
            (str(c.get("C", "")), tuple(str(c.get(k, "") or "") for k in ("D", "E", "F", "G")))
            for c in rows.values()
            if c.get("C")
        ]
        for item in candidates:
            data = json.loads(item["payload_json"])["data"]
            if data["quote_month"] != {"year": month.year, "month": month.month}:
                continue
            if Path(data["output_dir"]).resolve() != Path(path).resolve().parent:
                continue
            snapshot = data.get("associated_rows_snapshot")
            if not snapshot or hashlib.sha256(snapshot.encode()).hexdigest() != data.get(
                "associated_rows_snapshot_sha256"
            ):
                continue
            saved = json.loads(snapshot)["rows"]
            identities = [
                (
                    str(r["material_code"]),
                    tuple(str(r["cells"].get(k, "") or "") for k in ("D", "E", "F", "G")),
                )
                for r in saved
            ]
            if identities != actual:
                continue
            if run_ids is not None:
                run_ids.append(item["run_id"])
            return {
                f["source_role"]: Path(f["path"])
                for f in data["input_fingerprints"]
                if f["source_role"] in ("marketing", "base")
                and digest(Path(f["path"])) == f["sha256"]
            }
    except (OSError, sqlite3.Error, ValueError, KeyError, TypeError):
        return {}
    return {}


@lru_cache(maxsize=8)
def _source_table(path, stamp, code_header):
    del stamp
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book:
            iterator = sheet.iter_rows(values_only=True)
            headers = tuple(str(v or "").replace("\n", "").strip() for v in next(iterator, ()))
            if code_header not in headers:
                continue
            code_index = headers.index(code_header)
            result = defaultdict(list)
            for number, values in enumerate(iterator, 2):
                code = normalize_code(values[code_index])
                if code:
                    # Only retain needed business fields; do not cache contact columns.
                    selected = {
                        key: value
                        for key, value in zip(headers, values)
                        if key
                        in (
                            "产品状态",
                            "入库时间",
                            "建议市场零售价",
                            "市场通俗名称",
                            "RAM大小(MB/GB)",
                            "ROM大小(MB/GB)",
                            "颜色",
                        )
                        or "结算报价" in key
                    }
                    result[code].append((number, selected))
            return dict(result)
        return {}
    finally:
        book.close()


def source_facts(source_paths, material_code, month):
    context = {}
    for role, header in (("marketing", "物料编码"), ("base", "集团一级库物料编码")):
        path = source_paths.get(role)
        if not path:
            continue
        path = Path(path)
        try:
            stat = path.stat()
            found = _source_table(
                str(path.resolve()), (stat.st_mtime_ns, stat.st_size, stat.st_ctime_ns), header
            ).get(normalize_code(material_code), [])
        except (OSError, ValueError, KeyError):
            context[role + "_source_error"] = f"{path.name} 无法读取，请重新关联源表"
            continue
        if len(found) != 1:
            context[role + "_source_error"] = (
                f"{path.name} 匹配到 {len(found)} 条记录，需检查物料编码或重复记录"
            )
            continue
        number, values = found[0]
        origin = f"{path.name} · 第{number}行 · 物料编码{material_code}"
        if role == "marketing":
            state = str(values.get("产品状态") or "")
            # Presence is the user's stock criterion; explicit退库仍不能自动放行。
            context["stock"] = "不在库" if any(v in state for v in ("退库", "不在库")) else "在库"
            context["stock_source"] = origin
            entry = values.get("入库时间")
            if isinstance(entry, datetime):
                entry = entry.date()
            try:
                entry = date.fromisoformat(str(entry).replace("/", "-")[:10]).isoformat()
            except (ValueError, TypeError):
                entry = ""
            context.update(
                entry_date=entry,
                entry_date_source=origin,
                warehouse_price=str(values.get("建议市场零售价") or ""),
                warehouse_source=origin,
            )
            for source, dest in (
                ("市场通俗名称", "source_title"),
                ("RAM大小(MB/GB)", "source_ram"),
                ("ROM大小(MB/GB)", "source_storage"),
                ("颜色", "source_color"),
            ):
                context[dest] = str(values.get(source) or "")
        else:
            history = []
            for title, value in values.items():
                match = re.search(r"(20\d{2})年\s*(\d{1,2})月.*结算报价", title)
                if not match:
                    continue
                year, number = int(match[1]), int(match[2])
                if not 1 <= number <= 12 or (year, number) >= (month.year, month.month):
                    continue
                history.append((f"{year}-{number:02}", str(value or "")))
            context["history_prices"] = json.dumps(sorted(history), ensure_ascii=False)
            context["history_source"] = origin
    return context
