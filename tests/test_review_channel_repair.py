import pytest
from PIL import Image
from openpyxl import load_workbook
from quote_app.services.review_support import ReviewSession

pytest_plugins = ["tests.test_review_recovery"]


def picture(tmp_path, name="basis.png", color="red"):
    path = tmp_path / name
    Image.new("RGB", (640, 480), color).save(path)
    return path


def test_ap_text_ao_images_save_reload_and_noop(session, tmp_path):
    p = session.products[0]
    p.values["AP"] = "厂家说明 <原材料上涨>"
    p.attachments = [picture(tmp_path), picture(tmp_path, "second.png", "blue")]
    session.save(p)
    book = load_workbook(session.quote_path)
    sheet = book["5G手机"]
    assert sheet["AP2"].value == p.values["AP"]
    assert sheet["AO2"].value is None
    assert (
        len([im for im in sheet._images if im.anchor._from.col == 40 and im.anchor._from.row == 1])
        == 2
    )
    assert sheet["Z2"].value == "=K2/L2-1"
    book.close()
    before = session.quote_path.read_bytes()
    session.save(p)
    assert session.quote_path.read_bytes() == before
    restored = ReviewSession.from_workbook(session.quote_path, session.month)
    assert restored.products[0].values["AP"] == p.values["AP"]
    assert all(a.is_file() for a in restored.products[0].attachments)


def test_confirmation_three_states_survive_save(session):
    p = session.products[0]
    assert session.confirmation_status(p) == "未确认"
    session.save(p)
    session._confirmed[p.id] = session._fingerprint(p)
    assert session.confirmation_status(p) == "已确认"
    p.values["K"] = "2000"
    assert session.confirmation_status(p) == "修改后待重新确认"
    session.save(p)
    assert session.confirmation_status(p) == "修改后待重新确认"


def test_channel_repair_writes_price_link_image_and_reopens(session, tmp_path):
    p = session.products[0]
    session.repair_channel(
        p,
        "official",
        picture(tmp_path),
        "1999",
        "https://www.vmall.com/product/123.html",
        "官网失败后人工补录",
    )
    t = next(t for t in p.channels if t.channel == "official")
    assert t.price == "1999" and t.state == "manual_corrected"
    assert t.url == "https://www.vmall.com/product/123.html"
    book = load_workbook(session.quote_path)
    assert book["5G手机"]["AK2"].value == 1999
    assert book["5G手机"]["AK2"].hyperlink.target == t.url
    assert any(im.anchor._from.col == 39 for im in book["5G手机"]._images)
    book.close()
    checks = session.evaluate(p)
    assert next(c for c in checks if c.code == "E09").status == "未通过"
    restored = ReviewSession.from_workbook(session.quote_path, session.month)
    assert next(t for t in restored.products[0].channels if t.channel == "official").price == "1999"
    assert restored._history[-1]["action"] == "repair_channel"


def test_channel_repair_sidecar_failure_rolls_back(session, tmp_path, monkeypatch):
    before = session.quote_path.read_bytes()
    monkeypatch.setattr(session, "_persist", lambda: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        session.repair_channel(
            session.products[0],
            "official",
            picture(tmp_path),
            "1999",
            "https://www.vmall.com/product/123.html",
            "补录",
        )
    assert session.quote_path.read_bytes() == before
    assert not session._replacements


@pytest.mark.parametrize(
    "price,url",
    [
        ("0", "https://www.vmall.com"),
        ("NaN", "https://www.vmall.com"),
        ("100", "javascript:alert(1)"),
        ("100", "file:///tmp/x"),
    ],
)
def test_channel_repair_rejects_invalid_values(session, tmp_path, price, url):
    before = session.quote_path.read_bytes()
    with pytest.raises(ValueError):
        session.repair_channel(
            session.products[0], "official", picture(tmp_path), price, url, "补录"
        )
    assert session.quote_path.read_bytes() == before


def test_channel_repair_updates_minimum_outputs_and_can_then_save(session, tmp_path):
    p = session.products[0]
    session.repair_channel(
        p, "jd", picture(tmp_path), "1999", "https://item.jd.com/123.html", "补录京东"
    )
    session.repair_channel(
        p,
        "official",
        picture(tmp_path, "official.png"),
        "2099",
        "https://www.vmall.com/product/123.html",
        "补录官网",
    )
    p.values["AP"] = "最新说明"
    session.save(p)
    book = load_workbook(session.quote_path)
    sheet = book["5G手机"]
    assert sheet["AH2"].value == sheet["R2"].value == 1999
    assert sheet["S2"].value == "https://item.jd.com/123.html"
    assert sheet["AP2"].value == "最新说明"
    assert sheet["AK2"].hyperlink.target == "https://www.vmall.com/product/123.html"
    book.close()


def test_basis_images_save_rolls_back_when_persistence_fails(session, tmp_path, monkeypatch):
    p = session.products[0]
    p.values["AP"] = "说明"
    p.attachments = [picture(tmp_path)]
    before = session.quote_path.read_bytes()
    monkeypatch.setattr(session, "_persist", lambda: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        session.save(p)
    assert session.quote_path.read_bytes() == before


def test_old_notes_migrate_without_overwriting_existing_ao(session):
    import json

    p = session.products[0]
    session.save(p)
    data = json.loads(session.store_path.read_text())
    data.pop("notes_column")
    data["products"][p.id]["values"]["AO"] = "旧版填写的说明"
    session.store_path.write_text(json.dumps(data))
    reopened = ReviewSession(session.model, session.month)
    assert reopened.products[0].values["AP"] == "旧版填写的说明"
    reopened.save(reopened.products[0])
    book = load_workbook(session.quote_path)
    assert book["5G手机"]["AP2"].value == "旧版填写的说明"
    book.close()


def test_ao_chart_collision_cannot_delete_chart(session, tmp_path):
    from openpyxl.chart import BarChart
    from quote_app.services.review_workbook import digest

    book = load_workbook(session.quote_path)
    book["5G手机"].add_chart(BarChart(), "AO2")
    book.save(session.quote_path)
    book.close()
    session._version = digest(session.quote_path, fresh=True)
    before = session.quote_path.read_bytes()
    session.products[0].attachments = [picture(tmp_path)]
    with pytest.raises(ValueError, match="非图片"):
        session.save(session.products[0])
    assert session.quote_path.read_bytes() == before


def test_channel_repair_preserves_qname_namespace_declarations(session, tmp_path):
    from zipfile import ZipFile
    from quote_app.services.review_workbook import digest

    with ZipFile(session.quote_path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b"<worksheet ",
        b'<worksheet xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" mc:Ignorable="x14ac" ',
    )
    with ZipFile(session.quote_path, "w") as z:
        for n, data in parts.items():
            z.writestr(n, data)
    session._version = digest(session.quote_path, fresh=True)
    session.repair_channel(
        session.products[0],
        "official",
        picture(tmp_path),
        "1999",
        "https://www.vmall.com/product/123.html",
        "人工补录",
    )
    with ZipFile(session.quote_path) as z:
        xml = z.read("xl/worksheets/sheet1.xml")
    assert b"xmlns:x14ac=" in xml
    assert b'mc:Ignorable="x14ac"' in xml


def test_missing_workbook_reports_load_error(session):
    session.quote_path.unlink()
    reopened = ReviewSession(session.model, session.month)
    assert reopened.load_error
