from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image
import pytest

from quote_app.domain.models import InputPaths, QuoteMonth
from quote_app.evidence.models import EvidenceRecord, EvidenceRectangle
from quote_app.evidence.platform import (
    BrowserWindowIdentity,
    CaptureContext,
    CaptureRequest,
    make_capture_error,
)
from quote_app.evidence.semantic_state import VerifiedPageStateProbe, VerifiedSemanticState
from quote_app.services.full_pipeline import FullPipelineRequest, run_full_pipeline
from quote_app.services.web_run import (
    WebsiteRunRequest,
    WebsiteRunSnapshot,
    WebsiteRunSummary,
)
from quote_app.sites.registry import AdapterRegistry
from quote_app.tasks.models import (
    SCHEMA_VERSION,
    BusinessOutcome,
    InputFingerprint,
    RunRecord,
    RunState,
    TaskState,
    WebsiteChannel,
    WebsiteObservationCheckpoint,
    WebsiteResult,
    WebsiteTask,
)
from quote_app.tasks.repository import SQLiteTaskRepository
from quote_app.tasks.retry import RetryPolicy
from quote_app.tasks.runner import WebsiteTaskRunner
from tests.contract.test_official_vivo_live import _VivoPage, _set_cards
from tests.factories.workbook_factory import save_workbook


NOW = datetime(2026, 8, 12, 9, tzinfo=timezone.utc)


def _headers(length: int, **named: str) -> list[str]:
    headers = [f"字段{index}" for index in range(1, length + 1)]
    for column, value in named.items():
        headers[column_index_from_string(column) - 1] = value
    return headers


def _row(length: int, **values: object) -> list[object]:
    row: list[object] = [None] * length
    for column, value in values.items():
        row[column_index_from_string(column) - 1] = value
    return row


def _vivo_inputs(tmp_path: Path, *, rows: int = 1) -> InputPaths:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    materials = tuple(f"VIVO-{index}" for index in range(1, rows + 1))
    base = save_workbook(
        inputs / "base.xlsx",
        _headers(
            13,
            A="集团一级库物料编码",
            C="2026年3月结算报价（元/台）",
            D="2026年7月结算报价（元/台）",
        ),
        [
            _row(13, A=material, B=f"经理{index}", C=4499, D=4399)
            for index, material in enumerate(materials, start=1)
        ],
    )
    marketing = save_workbook(
        inputs / "marketing.xlsx",
        _headers(45, I="物料编码"),
        [
            _row(
                45,
                B="智能手机",
                C="VIVO",
                E="vivo X200",
                I=material,
                M="5G手机",
                V="2026-01-01",
                X=4499,
                AQ="12GB",
                AR="256GB",
                AS="辰夜黑",
            )
            for material in materials
        ],
    )
    bop = save_workbook(
        inputs / "bop.xlsx",
        _headers(26, K="集团一级库编码"),
        [_row(26, K=material, Y="12+256", Z="已配置") for material in materials],
    )
    return InputPaths(base, marketing, bop, tmp_path / "outputs")


class _PipelineVivoPage(_VivoPage):
    """Model a settled live offer: after 4499 -> 4399, later reads stay 4399."""

    def next_sale_price(self, node: object) -> str:
        if self.selected("capacity") != "12GB+256GB" or self.selected("color") != "辰夜黑":
            return ""
        value = "4499" if self.price_poll == 0 else "4399"
        self.price_poll += 1
        node.text_parts = [value]  # type: ignore[attr-defined]
        self.events.append(f"price:{self.selected('capacity')}:{self.selected('color')}")
        return f"¥{value}"


class _VivoSession:
    def __init__(self, page: _VivoPage | None = None) -> None:
        self.page = page or _PipelineVivoPage()

    def page_for(self, _site_family: str) -> _VivoPage:
        return self.page


class _FormalCapture:
    def __init__(self, *, fail: bool = False) -> None:
        self.fail = fail
        self.requests: list[CaptureRequest] = []

    def capture(self, request: CaptureRequest) -> EvidenceRecord:
        self.requests.append(request)
        if self.fail:
            raise make_capture_error("CAPTURE_FAILED", "fixture capture failed")
        request.destination.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (320, 180), (30, 116, 92)).save(request.destination)
        return EvidenceRecord(
            state=request.state,
            path=request.destination,
            sha256=hashlib.sha256(request.destination.read_bytes()).hexdigest(),
            pixel_width=320,
            pixel_height=180,
            captured_at=NOW,
            validation_code="CAPTURE_OK",
            annotations=tuple(
                EvidenceRectangle(rect.role, 10, 10 + index * 20, 80, 18)
                for index, rect in enumerate(request.css_rectangles)
            ),
        )


class _RunnerAudit:
    def __init__(self) -> None:
        self.tasks: list[WebsiteTask] = []
        self.observations: list[WebsiteObservationCheckpoint] = []
        self.results: list[WebsiteResult] = []


def _capture_context_provider(registry: AdapterRegistry):
    def provider(
        task: WebsiteTask,
        page: _VivoPage,
        expected: VerifiedSemanticState,
    ) -> CaptureContext:
        adapter = registry.adapter_for(task.brand, task.channel)
        prepare = getattr(adapter, "prepare_capture_view")
        reader_factory = getattr(adapter, "verified_state_reader")
        rectangles_reader = getattr(adapter, "capture_rectangles_for_capture")
        prepare(task, page, expected)
        reader = reader_factory(task, page, expected)
        assert reader() == expected
        return CaptureContext(
            expected_window=BrowserWindowIdentity("fixture", 42, "vivo-window"),
            stability_probe=VerifiedPageStateProbe(reader),
            css_rectangles=rectangles_reader(task, page, expected),
        )

    return provider


def _run_real_runner(
    request: WebsiteRunRequest,
    *,
    capture: _FormalCapture,
    session: _VivoSession | None = None,
    audit: _RunnerAudit | None = None,
) -> WebsiteRunSummary:
    selected_session = session or _VivoSession()
    registry = AdapterRegistry()
    with SQLiteTaskRepository(request.database_path) as repository:
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=request.run_id,
            browser_session=selected_session,
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=capture,
            evidence_dir=request.evidence_dir,
            retry_policy=RetryPolicy(technical_retries=0, retry_delay_seconds=0),
        ).run(request.tasks)
        observations = tuple(
            checkpoint
            for task in request.tasks
            if (checkpoint := repository.load_observation(task.task_id)) is not None
        )
        saved_results = tuple(
            result
            for task in request.tasks
            if (result := repository.load_result(task.task_id)) is not None
        )
        if audit is not None:
            audit.tasks.extend(request.tasks)
            audit.observations.extend(observations)
            audit.results.extend(saved_results)
        if request.checkpoint_sink is not None:
            request.checkpoint_sink(
                WebsiteRunSnapshot(observations, saved_results, frozenset())
            )
    return WebsiteRunSummary(
        succeeded=sum(result.state is TaskState.SUCCEEDED for result in results),
        waiting_for_login=0,
        technical_failure=sum(
            result.state is TaskState.TECHNICAL_FAILURE for result in results
        ),
        evidence_paths=tuple(
            result.evidence.path for result in results if result.evidence is not None
        ),
        technical_failure_codes=tuple(
            result.error_code for result in results if result.error_code is not None
        ),
    )


def _full_request(paths: InputPaths, tmp_path: Path) -> FullPipelineRequest:
    return FullPipelineRequest(
        paths=paths,
        quote_month=QuoteMonth(2026, 8),
        browser_profile_dir=tmp_path / "profile",
        database_path=tmp_path / "tasks.sqlite3",
        evidence_dir=tmp_path / "evidence",
        selected_brand="维沃",
        selected_channels=frozenset({WebsiteChannel.OFFICIAL}),
    )


def _image_anchors(sheet: object) -> set[str]:
    return {
        f"{get_column_letter(image.anchor._from.col + 1)}{image.anchor._from.row + 1}"
        for image in sheet._images  # type: ignore[attr-defined]
    }


def _run_record(tmp_path: Path, run_id: str) -> RunRecord:
    fingerprints = tuple(
        InputFingerprint(
            source_role=role,
            path=tmp_path / f"{role}.xlsx",
            sha256=str(index) * 64,
            byte_size=1,
            modified_ns=1,
        )
        for index, role in enumerate(("base", "marketing", "bop"), start=1)
    )
    return RunRecord(
        run_id=run_id,
        schema_version=SCHEMA_VERSION,
        quote_month=QuoteMonth(2026, 8),
        input_fingerprints=fingerprints,
        output_dir=tmp_path / "output",
        browser_profile_dir=tmp_path / "profile",
        associated_rows_snapshot="[]",
        associated_rows_snapshot_path=None,
        associated_rows_snapshot_sha256=hashlib.sha256(b"[]").hexdigest(),
        quote_path=None,
        report_path=None,
        state=RunState.CREATED,
        created_at=NOW,
        updated_at=NOW,
    )


def _task(
    task_id: str,
    *,
    run_id: str,
    model: str = "vivo X200",
    output_row: int = 2,
) -> WebsiteTask:
    return WebsiteTask(
        task_id=task_id,
        run_id=run_id,
        source_row_number=output_row,
        output_row_number=output_row,
        material_code=f"VIVO-{output_row}",
        brand="维沃",
        model_name=model,
        ram="12GB",
        storage="256GB",
        color="辰夜黑",
        channel=WebsiteChannel.OFFICIAL,
    )


def test_vivo_price_capture_checkpoint_reaches_excel_ak_and_an(tmp_path: Path) -> None:
    capture = _FormalCapture()
    audit = _RunnerAudit()

    result = run_full_pipeline(
        _full_request(_vivo_inputs(tmp_path), tmp_path),
        website_runner=lambda request: _run_real_runner(
            request,
            capture=capture,
            audit=audit,
        ),
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 4399
        assert _image_anchors(sheet) == {"AN2"}
    finally:
        quote.close()
    assert len(audit.tasks) == 1
    assert audit.observations[0].price == Decimal("4399")
    assert audit.observations[0].url.startswith("https://shop.vivo.com.cn/product/")
    assert audit.results[0].state is TaskState.SUCCEEDED
    assert len(capture.requests) == 1


@pytest.mark.parametrize(
    ("scenario", "outcome", "roles"),
    [
        ("no-model", BusinessOutcome.NO_MODEL, ("search_keyword", "result_region")),
        ("detail_missing_capacity.html", BusinessOutcome.CAPACITY_UNAVAILABLE, ("capacity",)),
        ("detail_missing_color.html", BusinessOutcome.COLOR_UNAVAILABLE, ("color",)),
    ],
)
def test_vivo_each_legal_no_state_reaches_formal_capture(
    tmp_path: Path,
    scenario: str,
    outcome: BusinessOutcome,
    roles: tuple[str, ...],
) -> None:
    run_id = f"vivo-legal-no-{outcome.value}"
    task = _task("vivo-legal-no", run_id=run_id)
    page = _PipelineVivoPage() if scenario == "no-model" else _PipelineVivoPage(scenario)
    if scenario == "no-model":
        _set_cards(
            page,
            (("vivo X200 Pro", "https://shop.vivo.com.cn/product/10010281?skuId=1"),),
        )
        page.empty_after_waits = 3
    capture = _FormalCapture()
    with SQLiteTaskRepository(tmp_path / f"{outcome.value}.sqlite3") as repository:
        repository.create_run(_run_record(tmp_path, run_id))
        registry = AdapterRegistry()
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=run_id,
            browser_session=_VivoSession(page),
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=capture,
            evidence_dir=tmp_path / f"evidence-{outcome.value}",
            retry_policy=RetryPolicy(technical_retries=0, retry_delay_seconds=0),
        ).run((task,))

    assert results[0].state is TaskState.SUCCEEDED
    assert results[0].outcome is outcome
    assert results[0].evidence is not None
    assert capture.requests[0].expected_roles == roles


def test_vivo_capture_failure_keeps_checkpoint_price_url_and_excel_ak(
    tmp_path: Path,
) -> None:
    seen_tasks: list[WebsiteTask] = []

    def website_runner(request: WebsiteRunRequest) -> WebsiteRunSummary:
        seen_tasks.extend(request.tasks)
        return _run_real_runner(request, capture=_FormalCapture(fail=True))

    result = run_full_pipeline(
        _full_request(_vivo_inputs(tmp_path), tmp_path),
        website_runner=website_runner,
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 4399
        assert "AN2" not in _image_anchors(sheet)
    finally:
        quote.close()
    with SQLiteTaskRepository(tmp_path / "tasks.sqlite3") as repository:
        assert len(seen_tasks) == 1
        checkpoint = repository.load_observation(seen_tasks[0].task_id)
        saved_result = repository.load_result(seen_tasks[0].task_id)
        assert checkpoint is not None
        assert checkpoint.price == Decimal("4399")
        assert checkpoint.url.startswith("https://shop.vivo.com.cn/product/")
        assert saved_result is not None
        assert saved_result.state is TaskState.TECHNICAL_FAILURE


def test_vivo_two_rows_keep_input_order_and_report_consistent_counts(
    tmp_path: Path,
) -> None:
    page = _PipelineVivoPage()
    result = run_full_pipeline(
        _full_request(_vivo_inputs(tmp_path, rows=2), tmp_path),
        website_runner=lambda request: _run_real_runner(
            request,
            capture=_FormalCapture(),
            session=_VivoSession(page),
        ),
    )

    assert [row.material_code for row in result.rows] == ["VIVO-1", "VIVO-2"]
    assert page.fill_calls == ["vivo X200", "vivo X200"]
    assert result.summary.total_rows == 2
    assert result.summary.completed_rows == 2
    assert result.summary.failed_rows == 0
    quote = load_workbook(result.quote_path, data_only=False)
    report = load_workbook(result.report_path, data_only=True)
    try:
        sheet = quote["5G手机"]
        assert [sheet["AK2"].value, sheet["AK3"].value] == [4399, 4399]
        assert _image_anchors(sheet) == {"AN2", "AN3"}
        overview = report["运行总览"]
        totals = {
            row[0].value: row[1].value
            for row in overview.iter_rows(min_col=1, max_col=2)
        }
        assert totals["报价总行数"] == 2
        assert report["处理明细"].max_row - 1 == 2
    finally:
        quote.close()
        report.close()


def test_vivo_restart_resumes_detail_checkpoint_without_repeating_search(
    tmp_path: Path,
) -> None:
    run_id = "vivo-resume"
    task = _task("vivo-resume-task", run_id=run_id)
    database = tmp_path / "resume.sqlite3"

    class InterruptCapture:
        def capture(self, _request: CaptureRequest) -> EvidenceRecord:
            raise SystemExit("stop after checkpoint")

    first_page = _PipelineVivoPage()
    with SQLiteTaskRepository(database) as repository:
        repository.create_run(_run_record(tmp_path, run_id))
        registry = AdapterRegistry()
        runner = WebsiteTaskRunner(
            repository=repository,
            run_id=run_id,
            browser_session=_VivoSession(first_page),
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=InterruptCapture(),
            evidence_dir=tmp_path / "resume-evidence",
        )
        with pytest.raises(SystemExit, match="stop after checkpoint"):
            runner.run((task,))
        checkpoint = repository.load_observation(task.task_id)
        assert checkpoint is not None
        assert checkpoint.url.startswith("https://shop.vivo.com.cn/product/")

    resumed_page = _PipelineVivoPage()
    with SQLiteTaskRepository(database) as repository:
        registry = AdapterRegistry()
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=run_id,
            browser_session=_VivoSession(resumed_page),
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=_FormalCapture(),
            evidence_dir=tmp_path / "resume-evidence",
            retry_policy=RetryPolicy(technical_retries=0, retry_delay_seconds=0),
        ).run((task,))

    assert results[0].state is TaskState.SUCCEEDED
    assert resumed_page.fill_calls == []
    assert resumed_page.search_submissions == 0
    assert len(resumed_page.goto_calls) == 1
    assert resumed_page.goto_calls[0].startswith("https://shop.vivo.com.cn/product/")


def test_vivo_rejects_iqoo_before_browser_visit_and_does_not_save_price(
    tmp_path: Path,
) -> None:
    run_id = "vivo-iqoo"
    task = _task("vivo-iqoo-task", run_id=run_id, model="iQOO 15")
    database = tmp_path / "iqoo.sqlite3"
    page = _VivoPage()
    capture = _FormalCapture()
    with SQLiteTaskRepository(database) as repository:
        repository.create_run(_run_record(tmp_path, run_id))
        registry = AdapterRegistry()
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=run_id,
            browser_session=_VivoSession(page),
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=capture,
            evidence_dir=tmp_path / "iqoo-evidence",
            retry_policy=RetryPolicy(technical_retries=0, retry_delay_seconds=0),
        ).run((task,))
        checkpoint = repository.load_observation(task.task_id)

    assert results[0].state is TaskState.TECHNICAL_FAILURE
    assert results[0].error_code == "UNSUPPORTED_VIVO_MODEL_FAMILY"
    assert results[0].price is None
    assert checkpoint is None
    assert page.goto_calls == [] and page.fill_calls == []
    assert capture.requests == []
