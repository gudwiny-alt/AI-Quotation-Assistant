from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pytest

from quote_app.services import web_to_excel
from quote_app.services.core_pipeline import CorePipelineError
from quote_app.services.web_pipeline import (
    QueryKey,
    SameRunQueryCache,
    execute_with_same_run_cache,
)
from quote_app.tasks.models import BusinessOutcome, WebsiteChannel
from tests.factories.web_run_factory import (
    make_business_result,
    make_observation_checkpoint,
    make_quote_row,
    make_tasks,
    make_technical_result,
    results_from_executor,
    run_three_channel_fixture,
    run_web_fixture,
    with_missing_evidence,
)


def _anchor_coordinate(image: Any) -> str:
    anchor = image.anchor
    return (
        f"{get_column_letter(anchor._from.col + 1)}"
        f"{anchor._from.row + 1}"
    )


def _overview_row(sheet: Any, label: str) -> list[object]:
    return next(
        [cell.value for cell in row]
        for row in sheet.iter_rows()
        if row[0].value == label
    )


def test_checkpoint_price_is_written_when_capture_failed(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    observation = make_observation_checkpoint(
        tasks[2],
        price=Decimal("4999"),
    )
    failure = make_technical_result(
        tasks[2],
        code="CAPTURE_FOREGROUND",
    )

    fixture = run_web_fixture(
        tmp_path,
        (row,),
        tasks,
        (failure,),
        observations=(observation,),
    )

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AK2"].value == 4999
        assert sheet["AN2"].value is None
    finally:
        workbook.close()


def test_checkpoint_price_survives_later_channel_failures(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    observation = make_observation_checkpoint(
        tasks[2],
        price=Decimal("4999"),
    )
    failures = (
        make_technical_result(tasks[0]),
        make_technical_result(tasks[1]),
        make_technical_result(tasks[2], code="CAPTURE_FOREGROUND"),
    )

    fixture = run_web_fixture(
        tmp_path,
        (row,),
        tasks,
        failures,
        observations=(observation,),
    )

    workbook = load_workbook(fixture.quote_path)
    try:
        assert workbook["5G手机"]["AK2"].value == 4999
    finally:
        workbook.close()


def test_legal_no_checkpoint_is_written_without_capture(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    observation = make_observation_checkpoint(
        tasks[2],
        outcome=BusinessOutcome.NO_MODEL,
    )
    failure = make_technical_result(
        tasks[2],
        code="CAPTURE_FOREGROUND",
    )

    fixture = run_web_fixture(
        tmp_path,
        (row,),
        tasks,
        (failure,),
        observations=(observation,),
    )

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AK2"].value == "无"
        assert sheet["AN2"].value is None
    finally:
        workbook.close()


def test_successful_result_wins_over_checkpoint_for_price_url_and_image(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    result = make_business_result(
        tmp_path,
        tasks[2],
        price=Decimal("4499"),
        url="https://official.example/final",
    )
    checkpoint = make_observation_checkpoint(
        tasks[2],
        price=Decimal("4999"),
        url="https://official.example/checkpoint",
    )

    fixture = run_web_fixture(
        tmp_path,
        (row,),
        tasks,
        (result,),
        observations=(checkpoint,),
    )

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AK2"].value == 4499
        assert sheet["AH2"].value == 4499
        assert sheet["S2"].value == "https://official.example/final"
        assert {_anchor_coordinate(image) for image in sheet._images} == {"AN2"}
    finally:
        workbook.close()


def test_zero_price_checkpoint_is_published_as_minimum_without_image(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    checkpoint = make_observation_checkpoint(
        tasks[2],
        price=Decimal("0"),
    )

    fixture = run_web_fixture(
        tmp_path,
        (row,),
        tasks,
        (make_technical_result(tasks[2], code="CAPTURE_FOREGROUND"),),
        observations=(checkpoint,),
    )

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert [sheet[f"{column}2"].value for column in ("AH", "R", "AK")] == [0, 0, 0]
        assert sheet["S2"].value == "https://official.example/product"
        assert len(sheet._images) == 0
    finally:
        workbook.close()


@pytest.mark.parametrize("kind", ("duplicate", "unknown"))
def test_checkpoint_rejects_duplicate_and_unknown_task_ids(
    kind: str,
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    checkpoint = make_observation_checkpoint(tasks[2])
    observations = (
        (checkpoint, replace(checkpoint, observed_at=checkpoint.observed_at.replace(second=1)))
        if kind == "duplicate"
        else (replace(checkpoint, task_id="unknown-checkpoint"),)
    )

    with pytest.raises(ValueError, match="duplicate|does not belong"):
        run_web_fixture(tmp_path, (row,), tasks, (), observations=observations)

    assert list(tmp_path.glob("*.xlsx")) == []


def test_three_channel_results_populate_minimum_url_images_and_report(
    tmp_path: Path,
) -> None:
    fixture = run_three_channel_fixture(tmp_path)

    workbook = load_workbook(fixture.quote_path, data_only=False)
    try:
        assert workbook.sheetnames == ["5G手机"]
        sheet = workbook["5G手机"]
        assert [sheet[f"{column}2"].value for column in ("AI", "AJ", "AK")] == [
            4499,
            4399,
            4399,
        ]
        assert sheet["AH2"].value == 4399
        assert sheet["R2"].value == 4399
        assert sheet["S2"].value == "https://tmall.example/product"
        assert len(sheet._images) == 3
        assert {_anchor_coordinate(image) for image in sheet._images} == {
            "AL2",
            "AM2",
            "AN2",
        }
        assert all(
            image.width / image.height == pytest.approx(2)
            for image in sheet._images
        )
        assert sheet["K2"].value is None
        assert sheet["L2"].value is None
        assert sheet["M2"].value is None
        assert sheet["N2"].value is None
        assert sheet["P2"].value is None
        assert sheet["Q2"].value is None
        assert sheet["X2"].value == '=IF(K2="","",IF(J2="无","无",(K2-J2)/J2))'
    finally:
        workbook.close()

    report = load_workbook(fixture.report_path)
    try:
        detail = report["处理明细"]
        assert [detail[f"{column}2"].value for column in ("I", "J", "K")] == [
            "价格成功（4499）；截图成功",
            "价格成功（4399）；截图成功",
            "价格成功（4399）；截图成功",
        ]
        assert detail["L2"].value == "完成"
        overview = report["运行总览"]
        assert _overview_row(overview, "报价总行数")[1] == 1
        assert _overview_row(overview, "处理完成")[1] == 1
        assert _overview_row(overview, "部分完成")[1] == 0
    finally:
        report.close()


def test_equal_low_price_uses_jd_then_tmall_then_official_priority(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(
            tmp_path,
            task,
            price=Decimal("4399"),
            url=f"https://{task.channel.value}.example/tied",
        )
        for task in tasks
    )

    fixture = run_web_fixture(tmp_path, (row,), tasks, results)

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AH2"].value == 4399
        assert sheet["S2"].value == "https://jd.example/tied"
    finally:
        workbook.close()


def test_three_legal_no_results_write_all_no_and_embed_proof_images(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(
            tmp_path,
            task,
            outcome=BusinessOutcome.NO_MODEL,
        )
        for task in tasks
    )

    fixture = run_web_fixture(tmp_path, (row,), tasks, results)

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert [sheet[f"{column}2"].value for column in ("AI", "AJ", "AK")] == [
            "无",
            "无",
            "无",
        ]
        assert [sheet[f"{column}2"].value for column in ("AH", "R", "S")] == [
            "无",
            "无",
            "无",
        ]
        assert len(sheet._images) == 3
        assert {_anchor_coordinate(image) for image in sheet._images} == {
            "AL2",
            "AM2",
            "AN2",
        }
    finally:
        workbook.close()

    report = load_workbook(fixture.report_path)
    try:
        detail = report["处理明细"]
        assert [detail[f"{column}2"].value for column in ("I", "J", "K")] == [
            "无（无该机型）",
            "无（无该机型）",
            "无（无该机型）",
        ]
        assert detail["L2"].value == "完成"
    finally:
        report.close()


def test_one_technical_failure_preserves_two_prices_and_reports_diagnostic(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = (
        make_technical_result(
            tasks[0],
            code="PAGE_TIMEOUT",
            message="页面加载超时",
        ),
        make_business_result(tmp_path, tasks[1], price=Decimal("4299")),
        make_business_result(tmp_path, tasks[2], price=Decimal("4399")),
    )

    fixture = run_web_fixture(tmp_path, (row,), tasks, results)

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AI2"].value is None
        assert sheet["AJ2"].value == 4299
        assert sheet["AK2"].value == 4399
        assert sheet["AH2"].value == 4299
        assert sheet["R2"].value == 4299
        assert sheet["S2"].value == "https://tmall.example/product"
        assert len(sheet._images) == 2
        assert {_anchor_coordinate(image) for image in sheet._images} == {
            "AM2",
            "AN2",
        }
    finally:
        workbook.close()

    report = load_workbook(fixture.report_path)
    try:
        detail = report["处理明细"]
        assert detail["I2"].value == "技术失败（PAGE_TIMEOUT：页面加载超时）"
        assert detail["J2"].value == "价格成功（4299）；截图成功"
        assert detail["L2"].value == "部分完成"
        assert "PAGE_TIMEOUT" in str(detail["N2"].value)
        overview = report["运行总览"]
        assert _overview_row(overview, "京东")[1:4] == [0, 0, 1]
        assert _overview_row(overview, "天猫")[1:4] == [1, 0, 0]
        assert _overview_row(overview, "官网")[1:4] == [1, 0, 0]
    finally:
        report.close()


@pytest.mark.parametrize(
    ("result_kinds", "expected_status"),
    [
        (("technical", "technical", "missing"), "部分完成"),
        (("technical", "technical", "technical"), "失败"),
    ],
)
def test_all_technical_or_missing_leaves_minimum_blank(
    result_kinds: tuple[str, str, str],
    expected_status: str,
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = tuple(
        make_technical_result(task)
        for task, kind in zip(tasks, result_kinds, strict=True)
        if kind == "technical"
    )

    fixture = run_web_fixture(tmp_path, (row,), tasks, results)

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert all(
            sheet[f"{column}2"].value is None
            for column in ("AH", "R", "S", "AI", "AJ", "AK")
        )
        assert len(sheet._images) == 0
    finally:
        workbook.close()
    report = load_workbook(fixture.report_path)
    try:
        assert report["处理明细"]["L2"].value == expected_status
    finally:
        report.close()


def test_legal_no_plus_technical_and_missing_does_not_claim_all_no(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = (
        make_business_result(
            tmp_path,
            tasks[0],
            outcome=BusinessOutcome.CAPACITY_UNAVAILABLE,
        ),
        make_technical_result(tasks[1]),
    )

    fixture = run_web_fixture(tmp_path, (row,), tasks, results)

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AI2"].value == "无"
        assert sheet["AJ2"].value is None
        assert sheet["AK2"].value is None
        assert all(sheet[f"{column}2"].value is None for column in ("AH", "R", "S"))
        assert len(sheet._images) == 1
        assert _anchor_coordinate(sheet._images[0]) == "AL2"
    finally:
        workbook.close()


def test_decimal_price_is_written_as_numeric_cell(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(
            tmp_path,
            task,
            price=Decimal("4399.50"),
        )
        for task in tasks
    )

    fixture = run_web_fixture(tmp_path, (row,), tasks, results)

    workbook = load_workbook(fixture.quote_path)
    try:
        cell = workbook["5G手机"]["AI2"]
        assert cell.value == 4399.5
        assert cell.data_type == "n"
    finally:
        workbook.close()


@pytest.mark.parametrize(
    "damage",
    [
        "missing",
        "hash",
        "dimensions",
        "unreadable",
        "diagnostic",
        "diagnostic_hardlink",
    ],
)
def test_invalid_or_missing_formal_evidence_is_not_embedded_and_is_reported(
    damage: str,
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    valid = make_business_result(tmp_path, tasks[0], price=Decimal("4499"))
    assert valid.evidence is not None
    if damage == "missing":
        damaged = with_missing_evidence(valid, tmp_path / "missing.png")
    elif damage == "hash":
        damaged = replace(
            valid,
            evidence=replace(valid.evidence, sha256="0" * 64),
        )
    else:
        if damage == "unreadable":
            valid.evidence.path.write_bytes(b"not-an-image")
            damaged = replace(
                valid,
                evidence=replace(
                    valid.evidence,
                    sha256=sha256(b"not-an-image").hexdigest(),
                ),
            )
        elif damage == "diagnostic":
            damaged = replace(
                valid,
                diagnostic_path=valid.evidence.path,
            )
        elif damage == "diagnostic_hardlink":
            diagnostic_alias = tmp_path / "diagnostic-alias.png"
            diagnostic_alias.hardlink_to(valid.evidence.path)
            damaged = replace(
                valid,
                diagnostic_path=diagnostic_alias,
            )
        else:
            damaged = replace(
                valid,
                evidence=replace(valid.evidence, pixel_width=999),
            )

    fixture = run_web_fixture(tmp_path, (row,), tasks, (damaged,))

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AI2"].value is None
        assert sheet["AL2"].value is None
        assert all(sheet[f"{column}2"].value is None for column in ("AH", "R", "S"))
        assert len(sheet._images) == 0
    finally:
        workbook.close()
    report = load_workbook(fixture.report_path)
    try:
        assert str(report["处理明细"]["I2"].value).startswith(
            "技术失败（EVIDENCE_"
        )
    finally:
        report.close()


def test_legal_no_with_invalid_evidence_is_not_published(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    valid = make_business_result(
        tmp_path,
        tasks[0],
        outcome=BusinessOutcome.NO_MODEL,
    )
    damaged = with_missing_evidence(valid, tmp_path / "missing-no-proof.png")

    fixture = run_web_fixture(tmp_path, (row,), tasks, (damaged,))

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AI2"].value is None
        assert sheet["AL2"].value is None
        assert all(sheet[f"{column}2"].value is None for column in ("AH", "R", "S"))
        assert len(sheet._images) == 0
    finally:
        workbook.close()

    report = load_workbook(fixture.report_path)
    try:
        assert str(report["处理明细"]["I2"].value).startswith(
            "技术失败（EVIDENCE_"
        )
    finally:
        report.close()


@pytest.mark.parametrize(
    ("field_name", "wrong_value"),
    [
        ("brand", "华为"),
        ("model_name", "小米 15 Pro"),
        ("ram", "16GB"),
        ("storage", "512GB"),
        ("color", "白色"),
        ("run_id", "different-run"),
    ],
)
def test_task_batch_must_bind_full_row_query_and_one_run_id_before_output(
    field_name: str,
    wrong_value: str,
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(tmp_path, task)
        for task in tasks
    )
    task_index = 1 if field_name == "run_id" else 0
    tampered = list(tasks)
    tampered[task_index] = replace(
        tampered[task_index],
        **{field_name: wrong_value},
    )

    with pytest.raises(ValueError, match="查询|run_id"):
        run_web_fixture(tmp_path, (row,), tuple(tampered), results)

    assert list(tmp_path.glob("*.xlsx")) == []


@pytest.mark.parametrize(
    ("query_field", "query_value", "identity_label"),
    [
        ("brand", "华为", "品牌"),
        ("model_name", "小米 15 Pro", "型号"),
    ],
)
def test_displayed_product_identity_must_match_web_query_before_output(
    query_field: str,
    query_value: str,
    identity_label: str,
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    row.web_query = replace(
        row.web_query,
        **{query_field: query_value},
    )
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(tmp_path, task)
        for task in tasks
    )

    with pytest.raises(ValueError, match=identity_label):
        run_web_fixture(tmp_path, (row,), tasks, results)

    assert list(tmp_path.glob("*.xlsx")) == []


def test_displayed_identity_accepts_canonical_brand_alias_and_model_whitespace(
    tmp_path: Path,
) -> None:
    row = make_quote_row(brand="HONOR", model_name="HONOR　500")
    row.cells["B"] = " honor "
    row.web_query = replace(
        row.web_query,
        brand="荣耀",
        model_name="honor   500",
    )
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(tmp_path, task)
        for task in tasks
    )

    fixture = run_web_fixture(tmp_path, (row,), tasks, results)

    workbook = load_workbook(fixture.quote_path)
    try:
        assert workbook["5G手机"]["AI2"].value == 4399
    finally:
        workbook.close()


def test_exact_query_cache_reuses_only_identical_normalized_key(
    tmp_path: Path,
) -> None:
    row = make_quote_row()
    first = make_tasks((row,), run_id="same-run")[0]
    identical = replace(
        first,
        task_id="same-run-identical",
        output_row_number=3,
        source_row_number=3,
        brand="　小米　",
        model_name="小米   15",
    )
    other_storage = replace(
        first,
        task_id="same-run-other-storage",
        storage="512GB",
    )
    other_channel = replace(
        first,
        task_id="same-run-other-channel",
        channel=WebsiteChannel.TMALL,
    )
    calls: list[str] = []

    def execute(task: Any) -> Any:
        calls.append(task.task_id)
        return make_business_result(tmp_path, task, price=Decimal("4399"))

    cache = SameRunQueryCache()
    results = tuple(
        cache.get_or_execute(task, execute)
        for task in (first, identical, other_storage, other_channel)
    )

    assert calls == [
        first.task_id,
        other_storage.task_id,
        other_channel.task_id,
    ]
    assert [result.task_id for result in results] == [
        first.task_id,
        identical.task_id,
        other_storage.task_id,
        other_channel.task_id,
    ]
    assert QueryKey.from_task(first) == QueryKey.from_task(identical)
    assert QueryKey.from_task(first) != QueryKey.from_task(other_storage)
    assert QueryKey.from_task(first) != QueryKey.from_task(other_channel)


def test_exact_query_cache_reuses_completed_technical_failure(
    tmp_path: Path,
) -> None:
    del tmp_path
    row = make_quote_row()
    first = make_tasks((row,), run_id="technical-run")[0]
    identical = replace(
        first,
        task_id="technical-run-identical",
        output_row_number=3,
        source_row_number=3,
    )
    calls: list[str] = []

    def execute(task: Any) -> Any:
        calls.append(task.task_id)
        return make_technical_result(task)

    cache = SameRunQueryCache()
    first_result = cache.get_or_execute(first, execute)
    second_result = cache.get_or_execute(identical, execute)

    assert calls == [first.task_id]
    assert first_result.task_id == first.task_id
    assert second_result.task_id == identical.task_id
    assert second_result.error_code == first_result.error_code
    with pytest.raises(ValueError, match="run IDs"):
        cache.get_or_execute(
            replace(first, task_id="other-run", run_id="other-run"),
            execute,
        )


@pytest.mark.parametrize("brand", ("其他品牌", "ZTE中兴"))
def test_unsupported_brand_report_marks_each_channel_unsupported(
    tmp_path: Path,
    brand: str,
) -> None:
    row = make_quote_row(brand=brand, model_name="其他手机")

    fixture = run_web_fixture(tmp_path, (row,), (), ())

    report = load_workbook(fixture.report_path)
    try:
        detail = report["处理明细"]
        assert [detail[f"{column}2"].value for column in ("I", "J", "K")] == [
            "不支持",
            "不支持",
            "不支持",
        ]
        assert detail["L2"].value == "不支持"
    finally:
        report.close()


def test_unsupported_brand_tasks_are_rejected_before_workbook_publication(
    tmp_path: Path,
) -> None:
    row = make_quote_row(brand="其他品牌", model_name="其他手机")
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(tmp_path, task)
        for task in tasks
    )

    with pytest.raises(ValueError, match="不支持"):
        run_web_fixture(tmp_path, (row,), tasks, results)

    assert list(tmp_path.glob("*.xlsx")) == []


def test_report_failure_rolls_back_quote_with_stable_output_pair_error(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(tmp_path, task)
        for task in tasks
    )

    def fail_report(*args: object, **kwargs: object) -> Path:
        raise ValueError("internal report failure")

    monkeypatch.setattr(web_to_excel, "write_execution_report", fail_report)

    with pytest.raises(CorePipelineError) as caught:
        run_web_fixture(tmp_path, (row,), tasks, results)

    assert caught.value.issues[0].code == "OUTPUT_PAIR_FAILED"
    assert "internal report failure" not in str(caught.value)
    assert list(tmp_path.glob("*.xlsx")) == []


def test_report_failure_with_unlink_failure_reports_exact_orphan_quote(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    row = make_quote_row()
    tasks = make_tasks((row,))
    results = tuple(
        make_business_result(tmp_path, task)
        for task in tasks
    )
    real_unlink = Path.unlink

    def fail_report(*args: object, **kwargs: object) -> Path:
        raise ValueError("internal report failure")

    def fail_quote_unlink(
        self: Path,
        missing_ok: bool = False,
    ) -> None:
        if (
            "终端供货价报价表" in self.name
            and not self.name.startswith(".")
        ):
            raise OSError("internal unlink failure")
        real_unlink(self, missing_ok=missing_ok)

    monkeypatch.setattr(web_to_excel, "write_execution_report", fail_report)
    monkeypatch.setattr(Path, "unlink", fail_quote_unlink)

    with pytest.raises(CorePipelineError) as caught:
        run_web_fixture(tmp_path, (row,), tasks, results)

    orphan_quotes = list(tmp_path.glob("*.xlsx"))
    assert len(orphan_quotes) == 1
    assert caught.value.issues[0].code == "ROLLBACK_FAILED"
    assert str(orphan_quotes[0]) in str(caught.value)
    assert "internal report failure" not in str(caught.value)
    assert "internal unlink failure" not in str(caught.value)


def test_two_rows_with_same_exact_query_execute_once_per_channel_but_each_gets_images(
    tmp_path: Path,
) -> None:
    rows = (
        make_quote_row(2, material_code="9101"),
        make_quote_row(3, material_code="9102"),
    )
    tasks = make_tasks(rows)
    calls: list[str] = []
    execute = results_from_executor(
        tmp_path,
        {
            WebsiteChannel.JD: Decimal("4499"),
            WebsiteChannel.TMALL: Decimal("4399"),
            WebsiteChannel.OFFICIAL: Decimal("4599"),
        },
    )

    def counted(task: Any) -> Any:
        calls.append(task.task_id)
        return execute(task)

    results = execute_with_same_run_cache(tasks, counted)
    fixture = run_web_fixture(tmp_path, rows, tasks, results)

    assert len(calls) == 3
    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert [sheet[f"AH{row}"].value for row in (2, 3)] == [4399, 4399]
        assert len(sheet._images) == 6
        assert {_anchor_coordinate(image) for image in sheet._images} == {
            "AL2",
            "AM2",
            "AN2",
            "AL3",
            "AM3",
            "AN3",
        }
    finally:
        workbook.close()


def test_rerun_of_same_result_set_does_not_duplicate_images(
    tmp_path: Path,
) -> None:
    first = run_three_channel_fixture(tmp_path)
    second = run_web_fixture(
        tmp_path,
        first.rows,
        first.tasks,
        first.results,
    )

    assert second.quote_path != first.quote_path
    for path in (first.quote_path, second.quote_path):
        workbook = load_workbook(path)
        try:
            assert len(workbook["5G手机"]._images) == 3
        finally:
            workbook.close()
