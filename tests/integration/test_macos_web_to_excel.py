from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image
import pytest

from quote_app.browser.worker import WorkerEvent
from quote_app.domain.models import QuoteMonth
from quote_app.evidence.models import (
    EvidenceRecord,
    EvidenceState,
    MacCapturePolicy,
)
from quote_app.evidence.platform import (
    BrowserWindowIdentity,
    CaptureRequest,
    PlatformEvidenceCapture,
)
from quote_app.services.incremental_publication import IncrementalExcelPublisher
from quote_app.services.web_run import WebsiteRunSnapshot
from quote_app.tasks.builder import create_run_record
from quote_app.tasks.models import (
    BusinessOutcome,
    TaskState,
    WebsiteChannel,
    WebsiteResult,
)
from quote_app.tasks.repository import SQLiteTaskRepository
from quote_app.tasks.runner import FixtureObservation, WebsiteTaskRunner
from tests.factories.web_run_factory import (
    TEMPLATE_PATH,
    make_quote_row,
    make_tasks,
    run_web_fixture,
)


HONOR_DETAIL_URL = (
    "https://www.honor.com/cn/shop/product/10086164863190.html"
)
_LOCAL_FILE_REFERENCE = re.compile(
    rb"(?i)(?:file:/+|(?<![A-Za-z])[A-Za-z]:[\\/]|/(?:Users|home)/)"
)


def _anchor_coordinate(image: Any) -> str:
    anchor = image.anchor
    return (
        f"{get_column_letter(anchor._from.col + 1)}"
        f"{anchor._from.row + 1}"
    )


def _honor_official_fixture(
    tmp_path: Path,
) -> tuple[Any, Any, WebsiteResult]:
    row = make_quote_row(
        brand="HONOR",
        model_name="荣耀Magic8",
        ram="12GB",
        storage="256GB",
        color="绒黑色",
    )
    task = next(
        task
        for task in make_tasks((row,), run_id="mac-honor")
        if task.channel is WebsiteChannel.OFFICIAL
    )
    evidence_path = tmp_path / "honor-macos-formal.png"
    Image.new("RGB", (1512, 982), (34, 48, 71)).save(evidence_path)
    evidence = EvidenceRecord(
        state=EvidenceState.NORMAL,
        path=evidence_path,
        sha256=sha256(evidence_path.read_bytes()).hexdigest(),
        pixel_width=1512,
        pixel_height=982,
        captured_at=datetime(2026, 7, 27, 9, 30, tzinfo=timezone.utc),
        validation_code="CAPTURE_OK",
    )
    result = WebsiteResult(
        task_id=task.task_id,
        state=TaskState.SUCCEEDED,
        outcome=BusinessOutcome.PRICE_FOUND,
        price=Decimal("4499"),
        url=HONOR_DETAIL_URL,
        evidence=evidence,
        diagnostic_path=None,
        error_code=None,
        error_message=None,
    )
    return row, task, result


def test_validated_honor_mac_evidence_is_standard_embedded_image(
    tmp_path: Path,
) -> None:
    row, task, result = _honor_official_fixture(tmp_path)

    fixture = run_web_fixture(tmp_path, (row,), (task,), (result,))

    workbook = load_workbook(fixture.quote_path, data_only=False)
    try:
        assert workbook.sheetnames == ["5G手机"]
        sheet = workbook["5G手机"]
        assert sheet["AK2"].value == 4499
        assert sheet["S2"].value == HONOR_DETAIL_URL
        assert sheet["AN2"].value is None
        assert len(sheet._images) == 1
        assert _anchor_coordinate(sheet._images[0]) == "AN2"
    finally:
        workbook.close()

    with ZipFile(fixture.quote_path) as archive:
        members = set(archive.namelist())
        drawing = archive.read("xl/drawings/drawing1.xml")
        relationships = archive.read("xl/drawings/_rels/drawing1.xml.rels")
        xml_payload = b"".join(
            archive.read(name)
            for name in members
            if name.endswith((".xml", ".rels"))
        )
    assert "xl/media/image1.png" in members
    assert b"<a:blip" in drawing
    assert b"r:embed=" in drawing
    assert b"/relationships/image" in relationships
    assert b"_xlfn.DISPIMG" not in xml_payload
    assert str(tmp_path).encode() not in xml_payload

    report = load_workbook(fixture.report_path)
    try:
        detail = report["处理明细"]
        assert detail["K2"].value == "价格成功（4499）；截图成功"
        assert detail["L2"].value == "部分完成"
        assert "官网：价格成功（4499）；截图成功" in str(
            detail["N2"].value
        )
    finally:
        report.close()
    with ZipFile(fixture.report_path) as archive:
        report_xml = b"".join(
            archive.read(name)
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        )
    assert str(tmp_path).encode() not in report_xml


def test_quote_output_removes_template_external_local_path_links(
    tmp_path: Path,
) -> None:
    row, task, result = _honor_official_fixture(tmp_path)

    fixture = run_web_fixture(tmp_path, (row,), (task,), (result,))

    with ZipFile("resources/templates/quote_template.xlsx") as template:
        template_xml = b"".join(
            template.read(name)
            for name in template.namelist()
            if name.endswith((".xml", ".rels"))
        )
    assert b"/Users/" in template_xml
    assert b"/home/" in template_xml
    assert b"C:/" in template_xml
    assert _LOCAL_FILE_REFERENCE.search(template_xml) is not None

    with ZipFile(fixture.quote_path) as archive:
        members = archive.namelist()
        xml_payload = b"".join(
            archive.read(name)
            for name in members
            if name.endswith((".xml", ".rels"))
        )
    assert not any(name.startswith("xl/externalLinks/") for name in members)
    assert b"/Users/" not in xml_payload
    assert b"/home/" not in xml_payload
    assert _LOCAL_FILE_REFERENCE.search(xml_payload) is None


class _StableRecoveryProbe:
    def semantic_hash(self) -> str:
        return "honor-recovery-stable"


class _RecoverySession:
    def page_for(self, _site_family: str) -> object:
        return object()


class _StoppedAfterObservationCapture:
    def capture(self, _request: CaptureRequest) -> EvidenceRecord:
        raise SystemExit("simulated process stop after official observation")


class _SuccessfulRecoveryCapture:
    def capture(self, request: CaptureRequest) -> EvidenceRecord:
        request.destination.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (240, 120), (34, 48, 71)).save(request.destination)
        return EvidenceRecord(
            state=request.state,
            path=request.destination,
            sha256=sha256(request.destination.read_bytes()).hexdigest(),
            pixel_width=240,
            pixel_height=120,
            captured_at=datetime(2026, 7, 30, 9, 30, tzinfo=timezone.utc),
            validation_code="CAPTURE_OK",
        )


def _recovery_observation(
    _task: object,
    _page: object,
) -> FixtureObservation:
    return FixtureObservation(
        outcome=BusinessOutcome.PRICE_FOUND,
        price=Decimal("4999"),
        url=HONOR_DETAIL_URL,
        css_rectangles=(),
        expected_roles=(),
        expected_window=BrowserWindowIdentity("macos", 42, "honor-window"),
        stability_probe=_StableRecoveryProbe(),
    )


def test_official_observation_recovery_reuses_generation_and_updates_same_partial(
    tmp_path: Path,
) -> None:
    """Break caught: a restart discards the saved price or publishes a new partial."""
    row = make_quote_row(
        brand="HONOR",
        model_name="荣耀Magic8",
        ram="12GB",
        storage="256GB",
        color="绒黑色",
    )
    inputs = tuple(tmp_path / f"{role}.xlsx" for role in ("base", "marketing", "bop"))
    for index, path in enumerate(inputs, start=1):
        path.write_bytes(f"input-{index}".encode())
    output_dir = tmp_path / "outputs"
    run = create_run_record(
        quote_month=QuoteMonth(2026, 8),
        base_path=inputs[0],
        marketing_path=inputs[1],
        bop_path=inputs[2],
        output_dir=output_dir,
        browser_profile_dir=tmp_path / "profile",
        rows=(row,),
    )
    task = next(
        task
        for task in make_tasks((row,), run_id=run.run_id)
        if task.channel is WebsiteChannel.OFFICIAL
    )
    database = tmp_path / "state.sqlite3"
    publisher = IncrementalExcelPublisher(
        quote_month=run.quote_month,
        rows=(row,),
        tasks=(task,),
        output_dir=output_dir,
        template_path=TEMPLATE_PATH,
        input_paths=None,
        capture_acceptance_policy=MacCapturePolicy.STRICT,
    )

    def runner_for(
        repository: SQLiteTaskRepository,
        capture: PlatformEvidenceCapture,
    ) -> WebsiteTaskRunner:
        def publish_stage(event: WorkerEvent) -> None:
            if event.event not in {"observation", "result"}:
                return
            observation = repository.load_observation(task.task_id)
            result = repository.load_result(task.task_id)
            publisher.publish(
                WebsiteRunSnapshot(
                    observations=(observation,) if observation is not None else (),
                    results=(result,) if result is not None else (),
                    waiting_task_ids=frozenset(),
                )
            )

        return WebsiteTaskRunner(
            repository=repository,
            run_id=run.run_id,
            browser_session=_RecoverySession(),
            adapter=_recovery_observation,
            evidence_capture=capture,
            evidence_dir=tmp_path / "evidence",
            event_sink=publish_stage,
        )

    with SQLiteTaskRepository(database) as repository:
        repository.create_run(run)
        interrupted_runner = runner_for(
            repository,
            _StoppedAfterObservationCapture(),
        )
        with pytest.raises(SystemExit, match="simulated process stop"):
            interrupted_runner.run((task,))
        generation_before_restart = repository.task_generation(task.task_id)
        assert repository.task_state(task.task_id) is TaskState.RUNNING

    partial_path = publisher.paths.quote_path
    assert partial_path.is_file()
    price_only = load_workbook(partial_path)
    try:
        sheet = price_only["5G手机"]
        assert sheet["AK2"].value == 4999
        assert len(sheet._images) == 0
    finally:
        price_only.close()

    with SQLiteTaskRepository(database) as repository:
        assert repository.task_state(task.task_id) is TaskState.PENDING
        assert repository.task_generation(task.task_id) == generation_before_restart
        assert repository.load_observation(task.task_id) is not None
        recovered_runner = runner_for(repository, _SuccessfulRecoveryCapture())
        recovered_runner.run((task,))
        assert repository.task_state(task.task_id) is TaskState.SUCCEEDED
        assert repository.task_generation(task.task_id) == generation_before_restart

    assert publisher.paths.quote_path == partial_path
    price_plus_image = load_workbook(partial_path)
    try:
        sheet = price_plus_image["5G手机"]
        assert sheet["AK2"].value == 4999
        assert len(sheet._images) == 1
        assert _anchor_coordinate(sheet._images[0]) == "AN2"
    finally:
        price_plus_image.close()


@pytest.mark.parametrize(
    "damage",
    ("missing", "unvalidated", "hash_changed", "symlink_replacement"),
)
def test_invalid_honor_evidence_is_not_embedded_and_is_reported_technical(
    damage: str,
    tmp_path: Path,
) -> None:
    row, task, result = _honor_official_fixture(tmp_path)
    assert result.evidence is not None
    evidence_path = result.evidence.path
    if damage == "missing":
        evidence_path.unlink()
    elif damage == "unvalidated":
        object.__setattr__(
            result.evidence,
            "validation_code",
            "CAPTURE_UNVALIDATED",
        )
    elif damage == "hash_changed":
        Image.new("RGB", (1512, 982), (90, 30, 20)).save(evidence_path)
    else:
        target = tmp_path / "matching-payload-target.png"
        evidence_path.rename(target)
        evidence_path.symlink_to(target)

    fixture = run_web_fixture(tmp_path, (row,), (task,), (result,))

    workbook = load_workbook(fixture.quote_path)
    try:
        sheet = workbook["5G手机"]
        assert sheet["AK2"].value is None
        assert sheet["S2"].value is None
        assert sheet["AN2"].value is None
        assert len(sheet._images) == 0
    finally:
        workbook.close()

    report = load_workbook(fixture.report_path)
    try:
        detail = report["处理明细"]
        assert str(detail["K2"].value).startswith(
            "技术失败（EVIDENCE_"
        )
        assert detail["L2"].value == "部分完成"
        assert "人工" in str(detail["O2"].value)
    finally:
        report.close()
