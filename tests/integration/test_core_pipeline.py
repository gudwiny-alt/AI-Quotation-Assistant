from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]
import pytest

from quote_app.domain.models import InputPaths, QuoteMonth
from quote_app.services.core_pipeline import CorePipelineError, run_core_pipeline
from tests.factories.workbook_factory import save_workbook


TEMPLATE_PATH = Path("resources/templates/quote_template.xlsx")


def _headers(length: int, **named: str) -> list[str]:
    from openpyxl.utils import column_index_from_string  # type: ignore[import-untyped]

    headers = [f"字段{index}" for index in range(1, length + 1)]
    for column, value in named.items():
        headers[column_index_from_string(column) - 1] = value
    return headers


def _row(length: int, **values: object) -> list[object]:
    from openpyxl.utils import column_index_from_string  # type: ignore[import-untyped]

    row: list[object] = [None] * length
    for column, value in values.items():
        row[column_index_from_string(column) - 1] = value
    return row


def _inputs(
    tmp_path: Path,
    *,
    base_rows: list[list[object]] | None = None,
    marketing_rows: list[list[object]] | None = None,
    bop_rows: list[list[object]] | None = None,
) -> InputPaths:
    input_dir = tmp_path / "inputs"
    output_dir = tmp_path / "outputs"
    input_dir.mkdir()
    base = save_workbook(
        input_dir / "base.xlsx",
        _headers(
            13,
            A="集团一级库物料编码",
            C="2026年3月结算报价（元/台）",
            D="2026年7月结算报价（元/台）",
        ),
        base_rows
        if base_rows is not None
        else [
            _row(13, A="9101", B="基地甲", C=3999, D=3899, I="经理甲"),
            _row(13, A="9102", B="基地乙", C="无", D="无", I="经理乙"),
        ],
    )
    marketing = save_workbook(
        input_dir / "marketing.xlsx",
        _headers(45, I="物料编码"),
        marketing_rows
        if marketing_rows is not None
        else [
            _row(
                45,
                B="智能手机",
                C="小米",
                E="小米17",
                I="9101",
                M="5G手机",
                V="2025-01-01",
                X=3999,
                AQ="12GB",
                AR="256GB",
                AS="黑色",
            ),
            _row(
                45,
                B="智能手机",
                C="其他品牌",
                E="其他手机",
                I="9102",
                M="5G手机",
                V="2025-02-01",
                X=2999,
                AQ="8GB",
                AR="128GB",
                AS="白色",
            ),
        ],
    )
    bop = save_workbook(
        input_dir / "bop.xlsx",
        _headers(26, K="集团一级库编码"),
        bop_rows
        if bop_rows is not None
        else [
            _row(26, K="9101", Y="12+256", Z="已配置"),
            _row(26, K="9102", Y="8+128", Z="已配置"),
        ],
    )
    return InputPaths(base, marketing, bop, output_dir)


def _digests(paths: InputPaths) -> dict[Path, str]:
    files = (paths.base, paths.marketing, paths.bop, TEMPLATE_PATH)
    return {path: sha256(path.read_bytes()).hexdigest() for path in files}


def test_core_pipeline_outputs_ordered_rows_reconciled_report_and_keeps_sources(
    tmp_path: Path,
) -> None:
    paths = _inputs(tmp_path)
    before = _digests(paths)

    result = run_core_pipeline(paths, QuoteMonth(2026, 8), TEMPLATE_PATH)

    assert result.quote_path.exists()
    assert result.report_path.exists()
    assert result.summary.total_rows == 2
    assert result.summary.completed_rows == 0
    assert result.summary.partial_rows == 1
    assert result.summary.failed_rows == 0
    assert result.summary.unsupported_rows == 1
    assert result.summary.manual_supplement_rows == 2
    assert [row.material_code for row in result.rows] == ["9101", "9102"]
    assert _digests(paths) == before

    quote = load_workbook(result.quote_path, data_only=False)
    report = load_workbook(result.report_path, data_only=False)
    try:
        assert quote.sheetnames == ["5G手机"]
        sheet = quote["5G手机"]
        assert sheet.max_row == 3
        assert [sheet[f"C{row}"].value for row in (2, 3)] == ["9101", "9102"]
        assert all(
            sheet[f"{column}{row}"].value is None
            for column in ("K", "L", "M", "N", "P", "Q", "AI", "AJ", "AK", "AL", "AM", "AN")
            for row in (2, 3)
        )
        assert [sheet[f"{column}1"].value for column in ("I", "J", "K")] == [
            "2026年3月结算报价（元/台）",
            "2026年7月结算报价（元/台）",
            "2026年8月结算报价（元/台）",
        ]
        assert sheet["X2"].value == '=IF(K2="","",IF(J2="无","无",(K2-J2)/J2))'

        detail = report["处理明细"]
        detail_statuses = [detail[f"L{row}"].value for row in (2, 3)]
        assert detail_statuses == ["部分完成", "不支持"]
        assert result.summary.total_rows == detail.max_row - 1
    finally:
        report.close()
        quote.close()


def test_blank_base_rows_are_excluded_and_duplicate_nonblank_rows_are_preserved(
    tmp_path: Path,
) -> None:
    paths = _inputs(
        tmp_path,
        base_rows=[
            _row(13, A="9101", C=3999, D=3899, I="经理甲"),
            _row(13, A="", C=3888, D=3788, I="经理空"),
            _row(13, A=" 9101 ", C=3777, D=3677, I="经理乙"),
        ],
        marketing_rows=[
            _row(
                45,
                B="智能手机",
                C="小米",
                E="小米17",
                I="9101",
                M="5G手机",
                V="2025-01-01",
                X=3999,
                AQ="12GB",
                AR="256GB",
                AS="黑色",
            )
        ],
        bop_rows=[_row(26, K="9101", Y="12+256", Z="已配置")],
    )

    result = run_core_pipeline(paths, QuoteMonth(2026, 8), TEMPLATE_PATH)

    assert [row.material_code for row in result.rows] == ["9101", "9101"]
    assert [row.source_row_number for row in result.rows] == [2, 4]
    assert all(
        any(issue.code == "DUPLICATE_MATERIAL_CODE" for issue in row.issues)
        for row in result.rows
    )
    quote = load_workbook(result.quote_path)
    try:
        assert [quote["5G手机"][f"C{row}"].value for row in (2, 3)] == [
            "9101",
            " 9101 ",
        ]
    finally:
        quote.close()


def test_row_association_issue_still_emits_both_outputs_and_is_failed(
    tmp_path: Path,
) -> None:
    paths = _inputs(tmp_path, marketing_rows=[])

    result = run_core_pipeline(paths, QuoteMonth(2026, 8), TEMPLATE_PATH)

    assert result.quote_path.exists()
    assert result.report_path.exists()
    assert result.summary.total_rows == 2
    assert result.summary.failed_rows == 2
    assert result.summary.partial_rows == 0
    report = load_workbook(result.report_path)
    try:
        assert [report["处理明细"][f"L{row}"].value for row in (2, 3)] == [
            "失败",
            "失败",
        ]
        assert all(
            "MARKETING_NOT_FOUND" in str(report["处理明细"][f"N{row}"].value)
            for row in (2, 3)
        )
    finally:
        report.close()


def test_fatal_precheck_writes_no_outputs_and_exposes_only_business_issues(
    tmp_path: Path,
) -> None:
    paths = _inputs(tmp_path)
    paths.base.write_bytes(b"not-an-xlsx")

    with pytest.raises(CorePipelineError) as caught:
        run_core_pipeline(paths, QuoteMonth(2026, 8), TEMPLATE_PATH)

    assert caught.value.issues
    assert caught.value.issues[0].code == "UNREADABLE_WORKBOOK"
    assert "预检查未通过" in str(caught.value)
    assert "BadZipFile" not in str(caught.value)
    assert not paths.output_dir.exists() or list(paths.output_dir.iterdir()) == []


def test_all_blank_base_codes_stop_without_outputs(tmp_path: Path) -> None:
    paths = _inputs(
        tmp_path,
        base_rows=[
            _row(13, A="", C=3999, D=3899, I="经理甲"),
            _row(13, A=None, C=3888, D=3788, I="经理乙"),
        ],
    )

    with pytest.raises(CorePipelineError) as caught:
        run_core_pipeline(paths, QuoteMonth(2026, 8), TEMPLATE_PATH)

    assert caught.value.issues[0].code == "NO_QUOTABLE_ROWS"
    assert not paths.output_dir.exists() or list(paths.output_dir.iterdir()) == []
