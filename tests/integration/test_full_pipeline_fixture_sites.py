from __future__ import annotations

from decimal import Decimal
import json
from pathlib import Path
import sqlite3

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pytest

from quote_app.domain.models import InputPaths, QuoteMonth
from quote_app.services.core_pipeline import CorePipelineError
from quote_app.services.full_pipeline import (
    FullPipelineRequest,
    RuntimeReadinessError,
    run_full_pipeline,
)
from quote_app.services.readiness import ReadinessCheck, ReadinessItem
from quote_app.services.web_run import (
    WebsiteRunRequest,
    WebsiteRunSnapshot,
    WebsiteRunSummary,
)
from quote_app.tasks.models import WebsiteChannel, WebsiteTask
from quote_app.tasks.repository import SQLiteTaskRepository
from tests.factories.web_run_factory import (
    make_business_result,
    make_observation_checkpoint,
    make_technical_result,
)
from tests.factories.workbook_factory import save_workbook


def _headers(length: int, **named: str) -> list[str]:
    from openpyxl.utils import column_index_from_string

    headers = [f"字段{index}" for index in range(1, length + 1)]
    for column, value in named.items():
        headers[column_index_from_string(column) - 1] = value
    return headers


def _row(length: int, **values: object) -> list[object]:
    from openpyxl.utils import column_index_from_string

    row: list[object] = [None] * length
    for column, value in values.items():
        row[column_index_from_string(column) - 1] = value
    return row


def _input_paths(tmp_path: Path) -> InputPaths:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    base = save_workbook(
        inputs / "base.xlsx",
        _headers(
            13,
            A="集团一级库物料编码",
            C="2026年3月结算报价（元/台）",
            D="2026年7月结算报价（元/台）",
        ),
        [
            _row(13, A="9101", B="经理甲", C=3999, D=3899),
            _row(13, A="9102", B="经理乙", C="无", D="无"),
        ],
    )
    marketing = save_workbook(
        inputs / "marketing.xlsx",
        _headers(45, I="物料编码"),
        [
            _row(
                45,
                B="智能手机",
                C="小米",
                E="小米17",
                I="9101",
                M="5G手机",
                V="2025-01-01",
                X=3999,
                AQ="12GB",
                AR="256GB",
                AS="黑色",
            ),
            _row(
                45,
                B="智能手机",
                C="其他品牌",
                E="其他手机",
                I="9102",
                M="5G手机",
                V="2025-02-01",
                X=2999,
                AQ="8GB",
                AR="128GB",
                AS="白色",
            ),
        ],
    )
    bop = save_workbook(
        inputs / "bop.xlsx",
        _headers(26, K="集团一级库编码"),
        [
            _row(26, K="9101", Y="12+256", Z="已配置"),
            _row(26, K="9102", Y="8+128", Z="已配置"),
        ],
    )
    return InputPaths(base, marketing, bop, tmp_path / "outputs")


def _fixture_website_runner(request: WebsiteRunRequest) -> WebsiteRunSummary:
    request.evidence_dir.mkdir(parents=True, exist_ok=True)
    evidence_paths: list[Path] = []
    with SQLiteTaskRepository(request.database_path) as repository:
        for task in request.tasks:
            repository.upsert_task(task)
            token = repository.start_attempt(task.task_id)
            result = make_business_result(
                request.evidence_dir,
                task,
                price=Decimal("4399"),
            )
            repository.save_result(result, token=token)
            assert result.evidence is not None
            evidence_paths.append(result.evidence.path)
    return WebsiteRunSummary(
        succeeded=len(request.tasks),
        waiting_for_login=0,
        technical_failure=0,
        evidence_paths=tuple(evidence_paths),
    )


def _honor_fixture_inputs(tmp_path: Path, *, include_honor: bool) -> InputPaths:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    base = save_workbook(
        inputs / "base.xlsx",
        _headers(
            13,
            A="集团一级库物料编码",
            C="2026年3月结算报价（元/台）",
            D="2026年7月结算报价（元/台）",
        ),
        [
            _row(13, A="HONOR-FIRST", B="经理甲", C=3999, D=3899),
            _row(13, A="HONOR-SECOND", B="经理乙", C=2999, D=2899),
            _row(13, A="NON-HONOR", B="经理丙", C=2899, D=2799),
        ],
    )
    marketing = save_workbook(
        inputs / "marketing.xlsx",
        _headers(45, I="物料编码"),
        [
            _row(
                45,
                B="智能手机",
                C="荣耀" if include_honor else "小米",
                E="HONOR 400" if include_honor else "小米17",
                I="HONOR-FIRST",
                M="5G手机",
                V="2025-01-01",
                X=3999,
                AQ="12GB",
                AR="256GB",
                AS="黑色",
            ),
            _row(
                45,
                B="智能手机",
                C=" HONOR " if include_honor else "小米",
                E="HONOR 400 Pro" if include_honor else "小米17",
                I="HONOR-SECOND",
                M="5G手机",
                V="2025-02-01",
                X=2999,
                AQ="8GB",
                AR="128GB",
                AS="白色",
            ),
            _row(
                45,
                B="智能手机",
                C="小米",
                E="小米17",
                I="NON-HONOR",
                M="5G手机",
                V="2025-03-01",
                X=2799,
                AQ="8GB",
                AR="256GB",
                AS="蓝色",
            ),
        ],
    )
    bop = save_workbook(
        inputs / "bop.xlsx",
        _headers(26, K="集团一级库编码"),
        [
            _row(26, K="HONOR-FIRST", Y="12+256", Z="已配置"),
            _row(26, K="HONOR-SECOND", Y="8+128", Z="已配置"),
            _row(26, K="NON-HONOR", Y="8+256", Z="已配置"),
        ],
    )
    return InputPaths(base, marketing, bop, tmp_path / "outputs")


@pytest.fixture
def fixture_inputs(tmp_path: Path) -> InputPaths:
    return _honor_fixture_inputs(tmp_path, include_honor=True)


@pytest.fixture
def fixture_inputs_without_honor(tmp_path: Path) -> InputPaths:
    return _honor_fixture_inputs(tmp_path, include_honor=False)


def _request(
    paths: InputPaths,
    tmp_path: Path,
    *,
    selected_brand: str | None = None,
    selected_channels: frozenset[WebsiteChannel] | None = None,
) -> FullPipelineRequest:
    return FullPipelineRequest(
        paths=paths,
        quote_month=QuoteMonth(2026, 8),
        browser_profile_dir=tmp_path / "app-data" / "browser-profile",
        database_path=tmp_path / "tasks.sqlite3",
        evidence_dir=tmp_path / "app-data" / "evidence",
        selected_brand=selected_brand,
        selected_channels=selected_channels,
    )


def _saved_task_brands(database_path: Path) -> set[str]:
    return {task["brand"] for task in _saved_tasks(database_path)}


def _saved_tasks(database_path: Path) -> list[dict[str, object]]:
    with sqlite3.connect(database_path) as connection:
        payloads = connection.execute("SELECT payload_json FROM website_tasks").fetchall()
    return [json.loads(payload)["data"] for (payload,) in payloads]


def _overview_value(workbook: object, label: str) -> object:
    overview = workbook["运行总览"]  # type: ignore[index]
    return next(
        row[1].value
        for row in overview.iter_rows(min_col=1, max_col=2)  # type: ignore[union-attr]
        if row[0].value == label
    )


def _detail_rows(workbook: object) -> int:
    detail = workbook["处理明细"]  # type: ignore[index]
    return detail.max_row - 1  # type: ignore[union-attr]


def test_honor_test_mode_outputs_all_honor_rows_in_base_order(
    fixture_inputs: InputPaths,
    tmp_path: Path,
) -> None:
    """Break caught: selected runs drop later normalized HONOR rows."""
    result = run_full_pipeline(
        _request(fixture_inputs, tmp_path, selected_brand="HONOR"),
        website_runner=_fixture_website_runner,
    )

    assert [row.material_code for row in result.rows] == [
        "HONOR-FIRST",
        "HONOR-SECOND",
    ]
    assert result.summary.total_rows == 2
    assert _saved_task_brands(tmp_path / "tasks.sqlite3") == {"HONOR"}
    assert len(_saved_tasks(tmp_path / "tasks.sqlite3")) == 6


def test_honor_official_scope_runs_only_official_tasks_for_all_honor_rows(
    fixture_inputs: InputPaths,
    tmp_path: Path,
) -> None:
    """Official-only acceptance must never schedule JD or Tmall."""
    seen: list[WebsiteTask] = []

    def runner(request: WebsiteRunRequest) -> WebsiteRunSummary:
        seen.extend(request.tasks)
        return _fixture_website_runner(request)

    result = run_full_pipeline(
        _request(
            fixture_inputs,
            tmp_path,
            selected_brand="HONOR",
            selected_channels=frozenset({WebsiteChannel.OFFICIAL}),
        ),
        website_runner=runner,
    )

    assert [row.material_code for row in result.rows] == [
        "HONOR-FIRST",
        "HONOR-SECOND",
    ]
    assert {task.channel for task in seen} == {WebsiteChannel.OFFICIAL}
    assert len(seen) == 2
    assert len(_saved_tasks(tmp_path / "tasks.sqlite3")) == 2


def test_honor_official_scope_writes_multiple_prices_and_images(
    fixture_inputs: InputPaths,
    tmp_path: Path,
) -> None:
    """Official-only completion writes both HONOR rows without JD/Tmall debt."""
    def runner(request: WebsiteRunRequest) -> WebsiteRunSummary:
        request.evidence_dir.mkdir(parents=True, exist_ok=True)
        evidence_paths: list[Path] = []
        prices = {"HONOR-FIRST": Decimal("4499"), "HONOR-SECOND": Decimal("4599")}
        with SQLiteTaskRepository(request.database_path) as repository:
            for task in request.tasks:
                assert task.channel is WebsiteChannel.OFFICIAL
                repository.upsert_task(task)
                token = repository.start_attempt(task.task_id)
                result = make_business_result(
                    request.evidence_dir,
                    task,
                    price=prices[task.material_code],
                )
                repository.save_result(result, token=token)
                assert result.evidence is not None
                evidence_paths.append(result.evidence.path)
        return WebsiteRunSummary(
            succeeded=2,
            waiting_for_login=0,
            technical_failure=0,
            evidence_paths=tuple(evidence_paths),
        )

    result = run_full_pipeline(
        _request(
            fixture_inputs,
            tmp_path,
            selected_brand="HONOR",
            selected_channels=frozenset({WebsiteChannel.OFFICIAL}),
        ),
        website_runner=runner,
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 4499
        assert sheet["AK3"].value == 4599
        assert sheet["AI2"].value is None
        assert sheet["AJ2"].value is None
        assert len(sheet._images) == 2
        assert {
            f"{get_column_letter(image.anchor._from.col + 1)}"
            f"{image.anchor._from.row + 1}"
            for image in sheet._images
        } == {"AN2", "AN3"}
    finally:
        quote.close()

    assert result.summary.completed_rows == 2
    assert result.summary.failed_rows == 0


def test_honor_test_mode_writes_all_honor_rows_and_three_evidence_images_per_row(
    fixture_inputs: InputPaths,
    tmp_path: Path,
) -> None:
    """Break caught: a multi-row selected run loses output rows or evidence."""
    result = run_full_pipeline(
        _request(fixture_inputs, tmp_path, selected_brand="HONOR"),
        website_runner=_fixture_website_runner,
    )

    assert len(_saved_tasks(tmp_path / "tasks.sqlite3")) == 6
    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet.max_row == 3
        assert [sheet[f"B{row}"].value for row in (2, 3)] == ["HONOR", "HONOR"]
        assert [sheet[f"C{row}"].value for row in (2, 3)] == [
            "HONOR-FIRST",
            "HONOR-SECOND",
        ]
        assert all(
            sheet[f"{column}{row}"].value is not None
            for row in (2, 3)
            for column in ("C", "D", "E", "F", "G", "AI", "AJ", "AK")
        )
        assert len(sheet._images) == 6
        assert {
            f"{get_column_letter(image.anchor._from.col + 1)}"
            f"{image.anchor._from.row + 1}"
            for image in sheet._images
        } == {"AL2", "AM2", "AN2", "AL3", "AM3", "AN3"}
    finally:
        quote.close()

    report = load_workbook(result.report_path, data_only=True)
    try:
        assert _overview_value(report, "报价总行数") == 2
        assert _detail_rows(report) == 2
    finally:
        report.close()


def test_honor_test_mode_rejects_input_without_honor_row(
    fixture_inputs_without_honor: InputPaths,
    tmp_path: Path,
) -> None:
    """Break caught: selected runs silently execute another brand when HONOR is absent."""
    with pytest.raises(CorePipelineError, match="未找到可用于荣耀闭环穿测的 HONOR 记录"):
        run_full_pipeline(
            _request(
                fixture_inputs_without_honor,
                tmp_path,
                selected_brand="HONOR",
            ),
            website_runner=_fixture_website_runner,
        )


def test_full_pipeline_generates_final_workbooks_from_three_inputs_and_saved_web_results(
    tmp_path: Path,
) -> None:
    """Break caught: omitting a workflow handoff leaves final AI:AN/summary unfilled."""
    result = run_full_pipeline(
        FullPipelineRequest(
            paths=_input_paths(tmp_path),
            quote_month=QuoteMonth(2026, 8),
            browser_profile_dir=tmp_path / "app-data" / "browser-profile",
            database_path=tmp_path / "app-data" / "tasks.sqlite3",
            evidence_dir=tmp_path / "app-data" / "evidence",
        ),
        website_runner=_fixture_website_runner,
    )

    assert result.quote_path.name == "2026年08月终端供货价报价表.xlsx"
    assert result.report_path.name == "2026年08月报价执行报告.xlsx"
    assert result.summary.total_rows == 2
    assert result.summary.completed_rows == 1
    assert result.summary.unsupported_rows == 1

    quote = load_workbook(result.quote_path, data_only=False)
    report = load_workbook(result.report_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert [sheet[f"{column}2"].value for column in ("AI", "AJ", "AK")] == [
            4399,
            4399,
            4399,
        ]
        assert sheet["AH2"].value == 4399
        assert sheet["S2"].value == "https://jd.example/product"
        assert len(sheet._images) == 3
        assert all(sheet[f"{column}3"].value is None for column in ("AI", "AJ", "AK"))
        assert report["处理明细"]["L2"].value == "完成"
        assert report["处理明细"]["L3"].value == "不支持"
    finally:
        quote.close()
        report.close()


def test_full_pipeline_replaces_one_partial_excel_pair_after_each_stage(
    tmp_path: Path,
) -> None:
    """Break caught: checkpoints do not replace the same readable partial pair."""
    observed_stages: list[dict[str, object]] = []

    def checkpoint_and_record(
        request: WebsiteRunRequest,
        snapshot: WebsiteRunSnapshot,
    ) -> None:
        assert request.checkpoint_sink is not None
        request.checkpoint_sink(snapshot)
        quote = load_workbook(
            tmp_path / "outputs" / "2026年08月终端供货价报价表-处理中.xlsx",
            data_only=False,
        )
        report = load_workbook(
            tmp_path / "outputs" / "2026年08月报价执行报告-处理中.xlsx",
            data_only=False,
        )
        try:
            sheet = quote["5G手机"]
            detail = report["处理明细"]
            observed_stages.append(
                {
                    "AK2": sheet["AK2"].value,
                    "AI2": sheet["AI2"].value,
                    "AN2_images": sum(
                        image.anchor._from.col == 39 and image.anchor._from.row == 1
                        for image in sheet._images
                    ),
                    "official": detail["K2"].value,
                    "jd": detail["I2"].value,
                    "status": detail["L2"].value,
                    "quote_reference": _overview_value(report, "报价表输出"),
                }
            )
        finally:
            quote.close()
            report.close()

    def staged_runner(request: WebsiteRunRequest) -> WebsiteRunSummary:
        request.evidence_dir.mkdir(parents=True, exist_ok=True)
        official = next(
            task for task in request.tasks if task.channel is WebsiteChannel.OFFICIAL
        )
        jd = next(task for task in request.tasks if task.channel is WebsiteChannel.JD)
        observation = make_observation_checkpoint(official, price=Decimal("4999"))
        official_result = make_business_result(
            request.evidence_dir,
            official,
            price=Decimal("4999"),
        )
        jd_failure = make_technical_result(jd)

        with SQLiteTaskRepository(request.database_path) as repository:
            for task in request.tasks:
                repository.upsert_task(task)
            official_token = repository.start_attempt(official.task_id)
            repository.save_observation(observation, token=official_token)
            checkpoint_and_record(
                request,
                WebsiteRunSnapshot((observation,), (), frozenset()),
            )
            repository.save_result(official_result, token=official_token)
            checkpoint_and_record(
                request,
                WebsiteRunSnapshot((observation,), (official_result,), frozenset()),
            )
            jd_token = repository.start_attempt(jd.task_id)
            repository.save_result(jd_failure, token=jd_token)
            checkpoint_and_record(
                request,
                WebsiteRunSnapshot(
                    (observation,),
                    (official_result, jd_failure),
                    frozenset(),
                ),
            )
        return WebsiteRunSummary(
            succeeded=1,
            waiting_for_login=0,
            technical_failure=1,
            evidence_paths=(official_result.evidence.path,),
        )

    request = FullPipelineRequest(
        paths=_input_paths(tmp_path),
        quote_month=QuoteMonth(2026, 8),
        browser_profile_dir=tmp_path / "app-data" / "browser-profile",
        database_path=tmp_path / "app-data" / "tasks.sqlite3",
        evidence_dir=tmp_path / "app-data" / "evidence",
    )
    result = run_full_pipeline(request, website_runner=staged_runner)

    assert observed_stages == [
        {
            "AK2": 4999,
            "AI2": None,
            "AN2_images": 0,
            "official": "价格成功（4999）；截图待补（CAPTURE_PENDING）",
            "jd": "待运行",
            "status": "部分完成",
            "quote_reference": "2026年08月终端供货价报价表-处理中.xlsx",
        },
        {
            "AK2": 4999,
            "AI2": None,
            "AN2_images": 1,
            "official": "价格成功（4999）；截图成功",
            "jd": "待运行",
            "status": "部分完成",
            "quote_reference": "2026年08月终端供货价报价表-处理中.xlsx",
        },
        {
            "AK2": 4999,
            "AI2": None,
            "AN2_images": 1,
            "official": "价格成功（4999）；截图成功",
            "jd": "技术失败（PAGE_TIMEOUT：页面在限定时间内未稳定）",
            "status": "部分完成",
            "quote_reference": "2026年08月终端供货价报价表-处理中.xlsx",
        },
    ]
    assert result.quote_path.is_file()
    assert len(list(request.paths.output_dir.glob("*终端供货价报价表-处理中.xlsx"))) == 1
    assert len(list(request.paths.output_dir.glob("*报价执行报告-处理中.xlsx"))) == 1


def test_full_pipeline_blocks_before_creating_run_when_runtime_is_not_ready(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "app-data" / "tasks.sqlite3"
    request = FullPipelineRequest(
        paths=_input_paths(tmp_path),
        quote_month=QuoteMonth(2026, 8),
        browser_profile_dir=tmp_path / "app-data" / "browser-profile",
        database_path=database_path,
        evidence_dir=tmp_path / "app-data" / "evidence",
        runtime_readiness=lambda: ReadinessCheck(
            (ReadinessItem("screen_capture", "未开启屏幕与系统音频录制权限", False),)
        ),
    )

    with pytest.raises(RuntimeReadinessError, match="未开启屏幕与系统音频录制权限"):
        run_full_pipeline(request, website_runner=_fixture_website_runner)

    assert not database_path.exists()
