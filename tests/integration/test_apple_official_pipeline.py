from __future__ import annotations

import hashlib
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image

from quote_app.domain.models import InputPaths, QuoteMonth
from quote_app.evidence.models import EvidenceRecord, EvidenceRectangle
from quote_app.evidence.platform import (
    BrowserWindowIdentity,
    CaptureContext,
    CaptureRequest,
    make_capture_error,
)
from quote_app.evidence.semantic_state import VerifiedPageStateProbe, VerifiedSemanticState
from quote_app.services.full_pipeline import FullPipelineRequest, run_full_pipeline
from quote_app.services.web_run import (
    WebsiteRunRequest,
    WebsiteRunSnapshot,
    WebsiteRunSummary,
)
from quote_app.sites.registry import AdapterRegistry
from quote_app.tasks.models import (
    SCHEMA_VERSION,
    BusinessOutcome,
    InputFingerprint,
    RunRecord,
    RunState,
    TaskState,
    WebsiteChannel,
    WebsiteTask,
)
from quote_app.tasks.repository import SQLiteTaskRepository
from quote_app.tasks.retry import RetryPolicy
from quote_app.tasks.runner import WebsiteTaskRunner
from tests.contract.test_official_apple_live import _AppleFixturePage
from tests.factories.workbook_factory import save_workbook


NOW = datetime(2026, 8, 19, 9, tzinfo=timezone.utc)
_FIXTURES = Path(__file__).parents[1] / "fixtures" / "sites" / "official_live" / "apple"


def _headers(length: int, **named: str) -> list[str]:
    headers = [f"字段{index}" for index in range(1, length + 1)]
    for column, value in named.items():
        headers[column_index_from_string(column) - 1] = value
    return headers


def _row(length: int, **values: object) -> list[object]:
    row: list[object] = [None] * length
    for column, value in values.items():
        row[column_index_from_string(column) - 1] = value
    return row


def _apple_inputs(tmp_path: Path) -> InputPaths:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    base = save_workbook(
        inputs / "base.xlsx",
        _headers(13, A="集团一级库物料编码", C="2026年3月结算报价（元/台）", D="2026年7月结算报价（元/台）"),
        [_row(13, A="APPLE-1", B="经理1", C=9999, D=9999)],
    )
    marketing = save_workbook(
        inputs / "marketing.xlsx",
        _headers(45, I="物料编码"),
        [
            _row(
                45,
                B="智能手机",
                C="苹果",
                E="iPhone 17",
                I="APPLE-1",
                M="5G手机",
                V="2026-01-01",
                X=9999,
                AQ="12GB",
                AR="512GB",
                AS="白色",
            )
        ],
    )
    bop = save_workbook(
        inputs / "bop.xlsx",
        _headers(26, K="集团一级库编码"),
        [_row(26, K="APPLE-1", Y="12+512", Z="已配置")],
    )
    return InputPaths(base, marketing, bop, tmp_path / "outputs")


class _AppleSession:
    def __init__(self, page: _AppleFixturePage | None = None) -> None:
        self.page = page or _AppleFixturePage(
            (_FIXTURES / "normal_flow.html").read_text(encoding="utf-8")
        )

    def page_for(self, _site_family: str) -> _AppleFixturePage:
        return self.page


class _FormalCapture:
    def __init__(self, *, fail: bool = False) -> None:
        self.fail = fail
        self.requests: list[CaptureRequest] = []

    def capture(self, request: CaptureRequest) -> EvidenceRecord:
        self.requests.append(request)
        if self.fail:
            raise make_capture_error("CAPTURE_FAILED", "fixture capture failed")
        request.destination.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (320, 180), (82, 82, 82)).save(request.destination)
        return EvidenceRecord(
            state=request.state,
            path=request.destination,
            sha256=hashlib.sha256(request.destination.read_bytes()).hexdigest(),
            pixel_width=320,
            pixel_height=180,
            captured_at=NOW,
            validation_code="CAPTURE_OK",
            annotations=tuple(
                EvidenceRectangle(rect.role, 10, 10 + index * 20, 80, 18)
                for index, rect in enumerate(request.css_rectangles)
            ),
        )


def _capture_context_provider(registry: AdapterRegistry):
    def provider(
        task: WebsiteTask,
        page: _AppleFixturePage,
        expected: VerifiedSemanticState,
    ) -> CaptureContext:
        adapter = registry.adapter_for(task.brand, task.channel)
        prepare = getattr(adapter, "prepare_capture_view")
        reader_factory = getattr(adapter, "verified_state_reader")
        rectangles_reader = getattr(adapter, "capture_rectangles_for_capture")
        prepare(task, page, expected)
        reader = reader_factory(task, page, expected)
        final_state = reader()
        if expected.outcome is BusinessOutcome.PRICE_FOUND:
            assert replace(final_state, css_rectangles=()) == expected
        else:
            assert final_state == expected
        rectangles = rectangles_reader(task, page, expected)
        assert final_state.css_rectangles == rectangles
        return CaptureContext(
            expected_window=BrowserWindowIdentity("fixture", 42, "apple-window"),
            stability_probe=VerifiedPageStateProbe(reader),
            css_rectangles=rectangles,
        )

    return provider


def _run_real_runner(
    request: WebsiteRunRequest,
    *,
    capture: _FormalCapture,
    session: _AppleSession | None = None,
) -> WebsiteRunSummary:
    registry = AdapterRegistry()
    with SQLiteTaskRepository(request.database_path) as repository:
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=request.run_id,
            browser_session=session or _AppleSession(),
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=capture,
            evidence_dir=request.evidence_dir,
            retry_policy=RetryPolicy(technical_retries=0, retry_delay_seconds=0),
        ).run(request.tasks)
        observations = tuple(
            checkpoint
            for task in request.tasks
            if (checkpoint := repository.load_observation(task.task_id)) is not None
        )
        saved_results = tuple(
            result
            for task in request.tasks
            if (result := repository.load_result(task.task_id)) is not None
        )
        if request.checkpoint_sink is not None:
            request.checkpoint_sink(
                WebsiteRunSnapshot(observations, saved_results, frozenset())
            )
    return WebsiteRunSummary(
        succeeded=sum(result.state is TaskState.SUCCEEDED for result in results),
        waiting_for_login=0,
        technical_failure=sum(
            result.state is TaskState.TECHNICAL_FAILURE for result in results
        ),
        evidence_paths=tuple(
            result.evidence.path for result in results if result.evidence is not None
        ),
        technical_failure_codes=tuple(
            result.error_code for result in results if result.error_code is not None
        ),
    )


def _full_request(paths: InputPaths, tmp_path: Path) -> FullPipelineRequest:
    return FullPipelineRequest(
        paths=paths,
        quote_month=QuoteMonth(2026, 8),
        browser_profile_dir=tmp_path / "profile",
        database_path=tmp_path / "tasks.sqlite3",
        evidence_dir=tmp_path / "evidence",
        selected_brand="苹果",
        selected_channels=frozenset({WebsiteChannel.OFFICIAL}),
    )


def _image_anchors(sheet: object) -> set[str]:
    return {
        f"{get_column_letter(image.anchor._from.col + 1)}{image.anchor._from.row + 1}"
        for image in sheet._images  # type: ignore[attr-defined]
    }


def _task(task_id: str, *, run_id: str, model: str = "iPhone 17") -> WebsiteTask:
    return WebsiteTask(
        task_id=task_id,
        run_id=run_id,
        source_row_number=2,
        output_row_number=2,
        material_code="APPLE-1",
        brand="苹果",
        model_name=model,
        ram="12GB",
        storage="512GB",
        color="白色",
        channel=WebsiteChannel.OFFICIAL,
    )


def _run_record(tmp_path: Path, run_id: str) -> RunRecord:
    return RunRecord(
        run_id=run_id,
        schema_version=SCHEMA_VERSION,
        quote_month=QuoteMonth(2026, 8),
        input_fingerprints=tuple(
            InputFingerprint(
                source_role=role,
                path=tmp_path / f"{role}.xlsx",
                sha256=str(index) * 64,
                byte_size=1,
                modified_ns=1,
            )
            for index, role in enumerate(("base", "marketing", "bop"), start=1)
        ),
        output_dir=tmp_path / "output",
        browser_profile_dir=tmp_path / "profile",
        associated_rows_snapshot="[]",
        associated_rows_snapshot_path=None,
        associated_rows_snapshot_sha256=hashlib.sha256(b"[]").hexdigest(),
        quote_path=None,
        report_path=None,
        state=RunState.CREATED,
        created_at=NOW,
        updated_at=NOW,
    )


def test_apple_price_capture_checkpoint_reaches_excel_ak_and_an(tmp_path: Path) -> None:
    capture = _FormalCapture()
    result = run_full_pipeline(
        _full_request(_apple_inputs(tmp_path), tmp_path),
        website_runner=lambda request: _run_real_runner(request, capture=capture),
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 9999
        assert _image_anchors(sheet) == {"AN2"}
    finally:
        quote.close()
    assert len(capture.requests) == 1
    assert tuple(rect.role for rect in capture.requests[0].css_rectangles) == (
        "title",
        "price",
        "capacity",
        "color",
    )


def test_apple_legal_no_generates_formal_search_evidence(tmp_path: Path) -> None:
    run_id = "apple-legal-no"
    page = _AppleFixturePage((_FIXTURES / "no_model.html").read_text(encoding="utf-8"))
    capture = _FormalCapture()
    with SQLiteTaskRepository(tmp_path / "legal-no.sqlite3") as repository:
        repository.create_run(_run_record(tmp_path, run_id))
        registry = AdapterRegistry()
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=run_id,
            browser_session=_AppleSession(page),
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=capture,
            evidence_dir=tmp_path / "evidence",
            retry_policy=RetryPolicy(technical_retries=0, retry_delay_seconds=0),
        ).run((_task("apple-no-model", run_id=run_id, model="iPhone 99"),))

    assert results[0].state is TaskState.SUCCEEDED
    assert results[0].outcome is BusinessOutcome.NO_MODEL
    assert capture.requests[0].expected_roles == ("search_keyword", "result_region")


def test_apple_capture_failure_keeps_checkpoint_price_and_excel_ak(tmp_path: Path) -> None:
    capture = _FormalCapture(fail=True)
    result = run_full_pipeline(
        _full_request(_apple_inputs(tmp_path), tmp_path),
        website_runner=lambda request: _run_real_runner(request, capture=capture),
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 9999
        assert "AN2" not in _image_anchors(sheet)
    finally:
        quote.close()
    assert result.summary.partial_rows == 1
    assert len(capture.requests) == 3
