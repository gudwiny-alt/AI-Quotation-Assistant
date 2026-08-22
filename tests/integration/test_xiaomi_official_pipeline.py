from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlsplit

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image
import pytest

from quote_app.domain.models import InputPaths, QuoteMonth
from quote_app.evidence.models import EvidenceRecord, EvidenceRectangle
from quote_app.evidence.platform import (
    BrowserWindowIdentity,
    CaptureContext,
    CaptureRequest,
    make_capture_error,
)
from quote_app.evidence.semantic_state import (
    VerifiedPageStateProbe,
    VerifiedSemanticState,
)
from quote_app.services.full_pipeline import FullPipelineRequest, run_full_pipeline
from quote_app.services.web_run import (
    WebsiteRunRequest,
    WebsiteRunSnapshot,
    WebsiteRunSummary,
)
from quote_app.sites.registry import AdapterRegistry
from quote_app.tasks.models import (
    SCHEMA_VERSION,
    BusinessOutcome,
    InputFingerprint,
    RunRecord,
    RunState,
    TaskState,
    WebsiteChannel,
    WebsiteTask,
)
from quote_app.tasks.repository import SQLiteTaskRepository
from quote_app.tasks.retry import RetryPolicy
from quote_app.tasks.runner import WebsiteTaskRunner
from tests.contract.test_official_xiaomi_live import _XiaomiFixturePage
from tests.factories.workbook_factory import save_workbook


NOW = datetime(2026, 8, 9, 9, tzinfo=timezone.utc)


def _headers(length: int, **named: str) -> list[str]:
    from openpyxl.utils import column_index_from_string

    headers = [f"字段{index}" for index in range(1, length + 1)]
    for column, value in named.items():
        headers[column_index_from_string(column) - 1] = value
    return headers


def _row(length: int, **values: object) -> list[object]:
    from openpyxl.utils import column_index_from_string

    row: list[object] = [None] * length
    for column, value in values.items():
        row[column_index_from_string(column) - 1] = value
    return row


def _xiaomi_inputs(tmp_path: Path, *, rows: int = 1) -> InputPaths:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    materials = tuple(f"XM-{index}" for index in range(1, rows + 1))
    colors = ("黑色", "白色")
    base = save_workbook(
        inputs / "base.xlsx",
        _headers(
            13,
            A="集团一级库物料编码",
            C="2026年3月结算报价（元/台）",
            D="2026年7月结算报价（元/台）",
        ),
        [
            _row(13, A=material, B=f"经理{index}", C=4599, D=4499)
            for index, material in enumerate(materials, start=1)
        ],
    )
    marketing = save_workbook(
        inputs / "marketing.xlsx",
        _headers(45, I="物料编码"),
        [
            _row(
                45,
                B="智能手机",
                C="小米",
                E="Xiaomi 17 Max",
                I=material,
                M="5G手机",
                V="2026-01-01",
                X=4599,
                AQ="12GB",
                AR="256GB",
                AS=colors[(index - 1) % len(colors)],
            )
            for index, material in enumerate(materials, start=1)
        ],
    )
    bop = save_workbook(
        inputs / "bop.xlsx",
        _headers(26, K="集团一级库编码"),
        [
            _row(26, K=material, Y="12+256", Z="已配置")
            for material in materials
        ],
    )
    return InputPaths(base, marketing, bop, tmp_path / "outputs")


class _ScenarioPage(_XiaomiFixturePage):
    """One controlled page that swaps only fixture HTML at each search."""

    def __init__(self) -> None:
        super().__init__("normal.html")
        self.search_keywords: list[str] = []

    def goto(self, url: str, **kwargs: object) -> None:
        parsed = urlsplit(url)
        if parsed.path == "/shop/search":
            keyword = unquote(parse_qs(parsed.query).get("keyword", [""])[0])
            self.search_keywords.append(keyword)
            fixture = {
                "Xiaomi 20": "no_model.html",
                "Xiaomi 18": "no_capacity.html",
                "Xiaomi 19": "no_color.html",
            }.get(keyword, "normal.html")
            replacement = _XiaomiFixturePage(fixture)
            self.root = replacement.root
            self.generation += 1
            self.price_reads = 0
            self._retitle_fixture(keyword, fixture)
        super().goto(url, **kwargs)

    def _retitle_fixture(self, keyword: str, fixture: str) -> None:
        for node in self.root.iter():
            if node.get("data-xiaomi-role") == "search-keyword":
                node.set("value", keyword)
        if fixture not in {"no_capacity.html", "no_color.html"}:
            return
        for node in self.root.iter():
            if node.get("data-xiaomi-role") in {
                "product-title",
                "detail-title",
            }:
                node.text = keyword


class _Session:
    def __init__(self, page: _ScenarioPage | None = None) -> None:
        self.page = page or _ScenarioPage()

    def page_for(self, _site_family: str) -> _ScenarioPage:
        return self.page


class _FormalCapture:
    def __init__(self, *, fail: bool = False) -> None:
        self.fail = fail
        self.requests: list[CaptureRequest] = []

    def capture(self, request: CaptureRequest) -> EvidenceRecord:
        self.requests.append(request)
        if self.fail:
            raise make_capture_error("CAPTURE_FAILED", "fixture capture failed")
        request.destination.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (320, 180), (35, 90, 170)).save(request.destination)
        return EvidenceRecord(
            state=request.state,
            path=request.destination,
            sha256=hashlib.sha256(request.destination.read_bytes()).hexdigest(),
            pixel_width=320,
            pixel_height=180,
            captured_at=NOW,
            validation_code="CAPTURE_OK",
            annotations=tuple(
                EvidenceRectangle(rect.role, 10, 10 + index * 20, 80, 18)
                for index, rect in enumerate(request.css_rectangles)
            ),
        )


def _capture_context_provider(registry: AdapterRegistry):
    def provider(
        task: WebsiteTask,
        page: _ScenarioPage,
        expected: VerifiedSemanticState,
    ) -> CaptureContext:
        adapter = registry.adapter_for(task.brand, task.channel)
        prepare = getattr(adapter, "prepare_capture_view")
        reader_factory = getattr(adapter, "verified_state_reader")
        rectangles_reader = getattr(adapter, "capture_rectangles_for_capture")
        prepare(task, page, expected)
        reader = reader_factory(task, page, expected)
        # This is the same final semantic reread required by the formal runtime.
        assert reader() == expected
        return CaptureContext(
            expected_window=BrowserWindowIdentity("fixture", 42, "xiaomi-window"),
            stability_probe=VerifiedPageStateProbe(reader),
            css_rectangles=rectangles_reader(task, page, expected),
        )

    return provider


def _run_real_runner(
    request: WebsiteRunRequest,
    *,
    capture: _FormalCapture,
    session: _Session | None = None,
) -> WebsiteRunSummary:
    selected_session = session or _Session()
    registry = AdapterRegistry()
    with SQLiteTaskRepository(request.database_path) as repository:
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=request.run_id,
            browser_session=selected_session,
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=capture,
            evidence_dir=request.evidence_dir,
            retry_policy=RetryPolicy(
                technical_retries=0,
                retry_delay_seconds=0,
            ),
        ).run(request.tasks)
        observations = tuple(
            checkpoint
            for task in request.tasks
            if (checkpoint := repository.load_observation(task.task_id)) is not None
        )
        saved_results = tuple(
            result
            for task in request.tasks
            if (result := repository.load_result(task.task_id)) is not None
        )
        if request.checkpoint_sink is not None:
            request.checkpoint_sink(
                WebsiteRunSnapshot(observations, saved_results, frozenset())
            )
    return WebsiteRunSummary(
        succeeded=sum(result.state is TaskState.SUCCEEDED for result in results),
        waiting_for_login=0,
        technical_failure=sum(
            result.state is TaskState.TECHNICAL_FAILURE for result in results
        ),
        evidence_paths=tuple(
            result.evidence.path
            for result in results
            if result.evidence is not None
        ),
        technical_failure_codes=tuple(
            result.error_code
            for result in results
            if result.error_code is not None
        ),
    )


def _full_request(paths: InputPaths, tmp_path: Path) -> FullPipelineRequest:
    return FullPipelineRequest(
        paths=paths,
        quote_month=QuoteMonth(2026, 8),
        browser_profile_dir=tmp_path / "profile",
        database_path=tmp_path / "tasks.sqlite3",
        evidence_dir=tmp_path / "evidence",
        selected_brand="小米",
        selected_channels=frozenset({WebsiteChannel.OFFICIAL}),
    )


def _image_anchors(sheet: object) -> set[str]:
    return {
        f"{get_column_letter(image.anchor._from.col + 1)}"
        f"{image.anchor._from.row + 1}"
        for image in sheet._images  # type: ignore[attr-defined]
    }


def _run_record(tmp_path: Path, run_id: str) -> RunRecord:
    fingerprints = tuple(
        InputFingerprint(
            source_role=role,
            path=tmp_path / f"{role}.xlsx",
            sha256=str(index) * 64,
            byte_size=1,
            modified_ns=1,
        )
        for index, role in enumerate(("base", "marketing", "bop"), start=1)
    )
    return RunRecord(
        run_id=run_id,
        schema_version=SCHEMA_VERSION,
        quote_month=QuoteMonth(2026, 8),
        input_fingerprints=fingerprints,
        output_dir=tmp_path / "output",
        browser_profile_dir=tmp_path / "profile",
        associated_rows_snapshot="[]",
        associated_rows_snapshot_path=None,
        associated_rows_snapshot_sha256=hashlib.sha256(b"[]").hexdigest(),
        quote_path=None,
        report_path=None,
        state=RunState.CREATED,
        created_at=NOW,
        updated_at=NOW,
    )


def _task(
    task_id: str,
    *,
    run_id: str,
    output_row: int,
    model: str = "Xiaomi 17 Max",
    color: str = "黑色",
) -> WebsiteTask:
    return WebsiteTask(
        task_id=task_id,
        run_id=run_id,
        source_row_number=output_row,
        output_row_number=output_row,
        material_code=f"XM-{output_row}",
        brand="小米",
        model_name=model,
        ram="12GB",
        storage="256GB",
        color=color,
        channel=WebsiteChannel.OFFICIAL,
    )


def test_xiaomi_price_capture_checkpoint_reaches_excel_ak_and_an(
    tmp_path: Path,
) -> None:
    """Break caught: a valid official result is persisted but never published to AK/AN."""
    capture = _FormalCapture()
    seen_tasks: list[WebsiteTask] = []

    def website_runner(request: WebsiteRunRequest) -> WebsiteRunSummary:
        seen_tasks.extend(request.tasks)
        return _run_real_runner(request, capture=capture)

    result = run_full_pipeline(
        _full_request(_xiaomi_inputs(tmp_path), tmp_path),
        website_runner=website_runner,
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 4399
        assert _image_anchors(sheet) == {"AN2"}
    finally:
        quote.close()
    with SQLiteTaskRepository(tmp_path / "tasks.sqlite3") as repository:
        # The durable observation is intentionally checked independently of Excel.
        assert len(seen_tasks) == 1
        checkpoint = repository.load_observation(seen_tasks[0].task_id)
        assert checkpoint is not None
        assert checkpoint.price == Decimal("4399")
        assert checkpoint.url.endswith("product_id=24648")
    assert len(capture.requests) == 1


@pytest.mark.parametrize(
    ("model", "outcome", "roles"),
    [
        ("Xiaomi 20", BusinessOutcome.NO_MODEL, ("search_keyword", "result_region")),
        ("Xiaomi 18", BusinessOutcome.CAPACITY_UNAVAILABLE, ("capacity",)),
        ("Xiaomi 19", BusinessOutcome.COLOR_UNAVAILABLE, ("color",)),
    ],
)
def test_xiaomi_each_legal_no_state_reaches_formal_capture(
    tmp_path: Path,
    model: str,
    outcome: BusinessOutcome,
    roles: tuple[str, ...],
) -> None:
    """Break caught: a legal-no observation is accepted but formal proof is skipped."""
    run_id = f"legal-no-{outcome.value}"
    task = _task("task-legal-no", run_id=run_id, output_row=2, model=model)
    capture = _FormalCapture()
    with SQLiteTaskRepository(tmp_path / f"{outcome.value}.sqlite3") as repository:
        repository.create_run(_run_record(tmp_path, run_id))
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=run_id,
            browser_session=_Session(),
            adapter_registry=(registry := AdapterRegistry()),
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=capture,
            evidence_dir=tmp_path / f"evidence-{outcome.value}",
            retry_policy=RetryPolicy(technical_retries=0, retry_delay_seconds=0),
        ).run((task,))

    assert results[0].state is TaskState.SUCCEEDED
    assert results[0].outcome is outcome
    assert results[0].evidence is not None
    assert capture.requests[0].expected_roles == roles


def test_xiaomi_capture_failure_keeps_checkpoint_price_url_and_excel_ak(
    tmp_path: Path,
) -> None:
    """Break caught: screenshot failure clears an already persisted price and URL."""
    seen_tasks: list[WebsiteTask] = []

    def website_runner(request: WebsiteRunRequest) -> WebsiteRunSummary:
        seen_tasks.extend(request.tasks)
        return _run_real_runner(request, capture=_FormalCapture(fail=True))

    result = run_full_pipeline(
        _full_request(_xiaomi_inputs(tmp_path), tmp_path),
        website_runner=website_runner,
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 4399
        assert "AN2" not in _image_anchors(sheet)
    finally:
        quote.close()
    with SQLiteTaskRepository(tmp_path / "tasks.sqlite3") as repository:
        assert len(seen_tasks) == 1
        checkpoint = repository.load_observation(seen_tasks[0].task_id)
        saved_result = repository.load_result(seen_tasks[0].task_id)
        assert checkpoint is not None
        assert checkpoint.price == Decimal("4399")
        assert checkpoint.url.endswith("product_id=24648")
        assert saved_result is not None
        assert saved_result.state is TaskState.TECHNICAL_FAILURE


def test_xiaomi_two_rows_run_in_base_order_and_report_consistent_counts(
    tmp_path: Path,
) -> None:
    """Break caught: multi-row official work is reordered, dropped, or miscounted."""
    page = _ScenarioPage()
    result = run_full_pipeline(
        _full_request(_xiaomi_inputs(tmp_path, rows=2), tmp_path),
        website_runner=lambda request: _run_real_runner(
            request,
            capture=_FormalCapture(),
            session=_Session(page),
        ),
    )

    assert [row.material_code for row in result.rows] == ["XM-1", "XM-2"]
    assert page.search_keywords == ["Xiaomi 17 Max", "Xiaomi 17 Max"]
    assert result.summary.total_rows == 2
    assert result.summary.completed_rows == 2
    assert result.summary.failed_rows == 0
    quote = load_workbook(result.quote_path, data_only=False)
    report = load_workbook(result.report_path, data_only=True)
    try:
        sheet = quote["5G手机"]
        assert [sheet["AK2"].value, sheet["AK3"].value] == [4399, 4399]
        assert _image_anchors(sheet) == {"AN2", "AN3"}
        overview = report["运行总览"]
        totals = {
            row[0].value: row[1].value
            for row in overview.iter_rows(min_col=1, max_col=2)
        }
        assert totals["报价总行数"] == 2
        assert report["处理明细"].max_row - 1 == 2
    finally:
        quote.close()
        report.close()


def test_xiaomi_restart_resumes_checkpoint_without_repeating_search(
    tmp_path: Path,
) -> None:
    """Break caught: restart discards a durable detail checkpoint and searches again."""
    run_id = "xiaomi-resume"
    task = _task("xiaomi-resume-task", run_id=run_id, output_row=2)
    database = tmp_path / "resume.sqlite3"

    class InterruptCapture:
        def capture(self, _request: CaptureRequest) -> EvidenceRecord:
            raise SystemExit("stop after checkpoint")

    first_page = _ScenarioPage()
    with SQLiteTaskRepository(database) as repository:
        repository.create_run(_run_record(tmp_path, run_id))
        registry = AdapterRegistry()
        runner = WebsiteTaskRunner(
            repository=repository,
            run_id=run_id,
            browser_session=_Session(first_page),
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=InterruptCapture(),
            evidence_dir=tmp_path / "resume-evidence",
        )
        with pytest.raises(SystemExit, match="stop after checkpoint"):
            runner.run((task,))
        assert repository.load_observation(task.task_id) is not None

    resumed_page = _ScenarioPage()
    with SQLiteTaskRepository(database) as repository:
        registry = AdapterRegistry()
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=run_id,
            browser_session=_Session(resumed_page),
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=_FormalCapture(),
            evidence_dir=tmp_path / "resume-evidence",
            retry_policy=RetryPolicy(technical_retries=0, retry_delay_seconds=0),
        ).run((task,))

    assert results[0].state is TaskState.SUCCEEDED
    assert resumed_page.search_keywords == []
    assert all("/shop/search" not in url for url in resumed_page.goto_calls)
    assert resumed_page.goto_calls[0].endswith("product_id=24648")
