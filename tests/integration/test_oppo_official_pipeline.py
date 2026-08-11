from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image

from quote_app.domain.models import InputPaths, QuoteMonth
from quote_app.evidence.models import EvidenceRecord, EvidenceRectangle
from quote_app.evidence.platform import (
    BrowserWindowIdentity,
    CaptureContext,
    CaptureRequest,
    make_capture_error,
)
from quote_app.evidence.semantic_state import (
    VerifiedPageStateProbe,
    VerifiedSemanticState,
)
from quote_app.services.full_pipeline import FullPipelineRequest, run_full_pipeline
from quote_app.services.web_run import (
    WebsiteRunRequest,
    WebsiteRunSnapshot,
    WebsiteRunSummary,
)
from quote_app.sites.registry import AdapterRegistry
from quote_app.tasks.models import (
    TaskState,
    WebsiteChannel,
    WebsiteObservationCheckpoint,
    WebsiteResult,
    WebsiteTask,
)
from quote_app.tasks.repository import AttemptRecord, SQLiteTaskRepository
from quote_app.tasks.retry import RetryPolicy
from quote_app.tasks.runner import WebsiteTaskRunner
from tests.contract.test_official_oppo_live import (
    _OppoFixturePage,
    _set_detail_product,
    _set_result_cards,
)
from tests.factories.workbook_factory import save_workbook


NOW = datetime(2026, 8, 12, 9, tzinfo=timezone.utc)

OPPO_ROWS = (
    ("OPPO-A5M", "OPPO A5m 5G", "8GB", "256GB", "钻石白"),
    ("OPPO-A6T", "OPPO A6t", "6GB", "128GB", "墨竹黑"),
)


def _headers(length: int, **named: str) -> list[str]:
    headers = [f"字段{index}" for index in range(1, length + 1)]
    for column, value in named.items():
        headers[column_index_from_string(column) - 1] = value
    return headers


def _row(length: int, **values: object) -> list[object]:
    row: list[object] = [None] * length
    for column, value in values.items():
        row[column_index_from_string(column) - 1] = value
    return row


def _oppo_inputs(tmp_path: Path, *, rows: int = 1) -> InputPaths:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    materials = tuple(f"OPPO-{index}" for index in range(1, rows + 1))
    base = save_workbook(
        inputs / "base.xlsx",
        _headers(
            13,
            A="集团一级库物料编码",
            C="2026年3月结算报价（元/台）",
            D="2026年7月结算报价（元/台）",
        ),
        [
            _row(13, A=material, B=f"经理{index}", C=2199, D=1999)
            for index, material in enumerate(materials, start=1)
        ],
    )
    marketing = save_workbook(
        inputs / "marketing.xlsx",
        _headers(45, I="物料编码"),
        [
            _row(
                45,
                B="智能手机",
                C="欧珀",
                E="OPPO A6 5G",
                I=material,
                M="5G手机",
                V="2026-01-01",
                X=2199,
                AQ="12GB",
                AR="256GB",
                AS="蓝海浮光",
            )
            for material in materials
        ],
    )
    bop = save_workbook(
        inputs / "bop.xlsx",
        _headers(26, K="集团一级库编码"),
        [_row(26, K=material, Y="12+256", Z="已配置") for material in materials],
    )
    return InputPaths(base, marketing, bop, tmp_path / "outputs")


def _oppo_model_only_inputs(tmp_path: Path) -> InputPaths:
    inputs = tmp_path / "model-only-inputs"
    inputs.mkdir()
    base = save_workbook(
        inputs / "base.xlsx",
        _headers(
            13,
            A="集团一级库物料编码",
            C="2026年3月结算报价（元/台）",
            D="2026年7月结算报价（元/台）",
        ),
        [
            _row(13, A=material, B=f"经理{index}", C=2199, D=1999)
            for index, (material, _model, _ram, _storage, _color) in enumerate(
                OPPO_ROWS, start=1
            )
        ],
    )
    marketing = save_workbook(
        inputs / "marketing.xlsx",
        _headers(45, I="物料编码"),
        [
            _row(
                45,
                B="智能手机",
                C="欧珀",
                E=model,
                I=material,
                M="5G手机",
                V="2026-01-01",
                X=2199,
                AQ=ram,
                AR=storage,
                AS=color,
            )
            for material, model, ram, storage, color in OPPO_ROWS
        ],
    )
    bop = save_workbook(
        inputs / "bop.xlsx",
        _headers(26, K="集团一级库编码"),
        [
            _row(
                26,
                K=material,
                Y=f"{ram.removesuffix('GB')}+{storage.removesuffix('GB')}",
                Z="已配置",
            )
            for material, _model, ram, storage, _color in OPPO_ROWS
        ],
    )
    return InputPaths(base, marketing, bop, tmp_path / "outputs")


class _OppoScenarioPage(_OppoFixturePage):
    def __init__(self) -> None:
        super().__init__("normal.html")
        self.search_keywords: list[str] = []

    def activate_results(self) -> None:
        keyword = self.locator('[data-oppo-role="search-input"]').input_value()
        self.search_keywords.append(keyword)
        super().activate_results()


def _set_inactive_target_options(
    page: _OppoFixturePage,
    *,
    model_name: str,
    capacity: str,
    color: str,
) -> None:
    _set_detail_product(
        page,
        model_name=model_name,
        capacity=capacity,
        color=color,
    )
    for kind, target, wrong in (
        ("capacity", capacity, "4GB+64GB"),
        ("color", color, "星夜黑"),
    ):
        options = [
            node
            for node in page.root.descendants()
            if node.attrs.get("data-option-kind") == kind
        ]
        assert len(options) >= 2
        active, inactive = options[:2]
        active.text_parts = [wrong]
        active.attrs["aria-selected"] = "true"
        active_classes = set(active.attrs.get("class", "").split())
        active_classes.add("active-btn")
        active.attrs["class"] = " ".join(sorted(active_classes))
        inactive.text_parts = [target]
        inactive.attrs["aria-selected"] = "false"
        inactive_classes = set(inactive.attrs.get("class", "").split())
        inactive_classes.discard("active-btn")
        inactive.attrs["class"] = " ".join(sorted(inactive_classes))


class _OppoModelOnlyScenarioPage(_OppoScenarioPage):
    def activate_results(self) -> None:
        keyword = self.locator('[data-oppo-role="search-input"]').input_value()
        if keyword == "OPPO A5m 5G":
            _set_result_cards(
                self,
                (("OPPO A5m 水晶粉 6GB+128GB", "/cn/web/products/38672.html?us=search"),),
            )
            _set_inactive_target_options(
                self,
                model_name="OPPO A5m",
                capacity="8GB+256GB",
                color="钻石白",
            )
        elif keyword == "OPPO A6t":
            _set_result_cards(
                self,
                (("OPPO A6t 青出于蓝 6GB+128GB", "/cn/web/products/41956.html?us=search"),),
            )
            _set_inactive_target_options(
                self,
                model_name="OPPO A6t",
                capacity="6GB+128GB",
                color="墨竹黑",
            )
        super().activate_results()


class _OppoSession:
    def __init__(self, page: _OppoScenarioPage | None = None) -> None:
        self.page = page or _OppoScenarioPage()

    def page_for(self, _site_family: str) -> _OppoScenarioPage:
        return self.page


class _FormalCapture:
    def __init__(self, *, fail: bool = False) -> None:
        self.fail = fail
        self.requests: list[CaptureRequest] = []

    def capture(self, request: CaptureRequest) -> EvidenceRecord:
        self.requests.append(request)
        if self.fail:
            raise make_capture_error("CAPTURE_FAILED", "fixture capture failed")
        request.destination.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (320, 180), (35, 90, 170)).save(request.destination)
        return EvidenceRecord(
            state=request.state,
            path=request.destination,
            sha256=hashlib.sha256(request.destination.read_bytes()).hexdigest(),
            pixel_width=320,
            pixel_height=180,
            captured_at=NOW,
            validation_code="CAPTURE_OK",
            annotations=tuple(
                EvidenceRectangle(rect.role, 10, 10 + index * 20, 80, 18)
                for index, rect in enumerate(request.css_rectangles)
            ),
        )


class _RunnerAudit:
    def __init__(self) -> None:
        self.request_tasks: list[WebsiteTask] = []
        self.observations: list[WebsiteObservationCheckpoint] = []
        self.results: list[WebsiteResult] = []
        self.attempts: dict[str, tuple[AttemptRecord, ...]] = {}


def _capture_context_provider(registry: AdapterRegistry):
    def provider(
        task: WebsiteTask,
        page: _OppoScenarioPage,
        expected: VerifiedSemanticState,
    ) -> CaptureContext:
        adapter = registry.adapter_for(task.brand, task.channel)
        prepare = getattr(adapter, "prepare_capture_view")
        reader_factory = getattr(adapter, "verified_state_reader")
        rectangles_reader = getattr(adapter, "capture_rectangles_for_capture")
        prepare(task, page, expected)
        reader = reader_factory(task, page, expected)
        assert reader() == expected
        return CaptureContext(
            expected_window=BrowserWindowIdentity("fixture", 42, "oppo-window"),
            stability_probe=VerifiedPageStateProbe(reader),
            css_rectangles=rectangles_reader(task, page, expected),
        )

    return provider


def _run_real_runner(
    request: WebsiteRunRequest,
    *,
    capture: _FormalCapture,
    session: _OppoSession | None = None,
    audit: _RunnerAudit | None = None,
) -> WebsiteRunSummary:
    selected_session = session or _OppoSession()
    registry = AdapterRegistry()
    with SQLiteTaskRepository(request.database_path) as repository:
        results = WebsiteTaskRunner(
            repository=repository,
            run_id=request.run_id,
            browser_session=selected_session,
            adapter_registry=registry,
            capture_context_provider=_capture_context_provider(registry),
            evidence_capture=capture,
            evidence_dir=request.evidence_dir,
            retry_policy=RetryPolicy(
                technical_retries=0,
                retry_delay_seconds=0,
            ),
        ).run(request.tasks)
        observations = tuple(
            checkpoint
            for task in request.tasks
            if (checkpoint := repository.load_observation(task.task_id)) is not None
        )
        saved_results = tuple(
            result
            for task in request.tasks
            if (result := repository.load_result(task.task_id)) is not None
        )
        if audit is not None:
            audit.request_tasks.extend(request.tasks)
            audit.observations.extend(observations)
            audit.results.extend(saved_results)
            audit.attempts.update(
                {
                    task.task_id: repository.list_attempts(task.task_id)
                    for task in request.tasks
                }
            )
        if request.checkpoint_sink is not None:
            request.checkpoint_sink(
                WebsiteRunSnapshot(observations, saved_results, frozenset())
            )
    return WebsiteRunSummary(
        succeeded=sum(result.state is TaskState.SUCCEEDED for result in results),
        waiting_for_login=0,
        technical_failure=sum(
            result.state is TaskState.TECHNICAL_FAILURE for result in results
        ),
        evidence_paths=tuple(
            result.evidence.path
            for result in results
            if result.evidence is not None
        ),
        technical_failure_codes=tuple(
            result.error_code
            for result in results
            if result.error_code is not None
        ),
    )


def _image_anchors(sheet: object) -> set[str]:
    return {
        f"{get_column_letter(image.anchor._from.col + 1)}"
        f"{image.anchor._from.row + 1}"
        for image in sheet._images  # type: ignore[attr-defined]
    }


def _full_request(paths: InputPaths, tmp_path: Path) -> FullPipelineRequest:
    return FullPipelineRequest(
        paths=paths,
        quote_month=QuoteMonth(2026, 8),
        browser_profile_dir=tmp_path / "profile",
        database_path=tmp_path / "tasks.sqlite3",
        evidence_dir=tmp_path / "evidence",
        selected_brand="欧珀",
        selected_channels=frozenset({WebsiteChannel.OFFICIAL}),
    )


def test_oppo_price_and_capture_reach_excel_ak_and_an(tmp_path: Path) -> None:
    capture = _FormalCapture()

    def website_runner(request: WebsiteRunRequest) -> WebsiteRunSummary:
        return _run_real_runner(
            request,
            capture=capture,
            session=_OppoSession(),  # type: ignore[arg-type]
        )

    result = run_full_pipeline(
        _full_request(_oppo_inputs(tmp_path), tmp_path),
        website_runner=website_runner,
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 1899
        assert _image_anchors(sheet) == {"AN2"}
    finally:
        quote.close()
    assert len(capture.requests) == 1


def test_oppo_capture_failure_keeps_price_in_excel(tmp_path: Path) -> None:
    result = run_full_pipeline(
        _full_request(_oppo_inputs(tmp_path), tmp_path),
        website_runner=lambda request: _run_real_runner(
            request,
            capture=_FormalCapture(fail=True),
            session=_OppoSession(),  # type: ignore[arg-type]
        ),
    )

    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert sheet["AK2"].value == 1899
        assert "AN2" not in _image_anchors(sheet)
    finally:
        quote.close()


def test_oppo_two_rows_keep_base_order_and_complete_once_each(tmp_path: Path) -> None:
    page = _OppoScenarioPage()
    result = run_full_pipeline(
        _full_request(_oppo_inputs(tmp_path, rows=2), tmp_path),
        website_runner=lambda request: _run_real_runner(
            request,
            capture=_FormalCapture(),
            session=_OppoSession(page),  # type: ignore[arg-type]
        ),
    )

    assert [row.material_code for row in result.rows] == ["OPPO-1", "OPPO-2"]
    assert page.search_keywords == ["OPPO A6 5G", "OPPO A6 5G"]
    assert result.summary.total_rows == 2
    assert result.summary.completed_rows == 2
    assert result.summary.failed_rows == 0
    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert [sheet["AK2"].value, sheet["AK3"].value] == [1899, 1899]
        assert _image_anchors(sheet) == {"AN2", "AN3"}
    finally:
        quote.close()


def test_oppo_a5m_a6t_keep_input_order_and_write_ak_an(tmp_path: Path) -> None:
    page = _OppoModelOnlyScenarioPage()
    session = _OppoSession(page)
    capture = _FormalCapture()
    audit = _RunnerAudit()
    result = run_full_pipeline(
        _full_request(_oppo_model_only_inputs(tmp_path), tmp_path),
        website_runner=lambda request: _run_real_runner(
            request,
            capture=capture,
            session=session,  # type: ignore[arg-type]
            audit=audit,
        ),
    )

    assert [row.material_code for row in result.rows] == ["OPPO-A5M", "OPPO-A6T"]
    assert page.search_keywords == ["OPPO A5m 5G", "OPPO A6t"]
    assert list(zip(page.option_clicks, page.option_labels, strict=True)) == [
        ("capacity", "8GB+256GB"),
        ("color", "钻石白"),
        ("capacity", "6GB+128GB"),
        ("color", "墨竹黑"),
    ]
    task_ids = [task.task_id for task in audit.request_tasks]
    assert len(task_ids) == 2
    assert len(set(task_ids)) == 2
    assert [item.task_id for item in audit.observations] == task_ids
    assert [item.task_id for item in audit.results] == task_ids
    assert all(item.state is TaskState.SUCCEEDED for item in audit.results)
    assert set(audit.attempts) == set(task_ids)
    for task_id in task_ids:
        assert len(audit.attempts[task_id]) == 1
        attempt = audit.attempts[task_id][0]
        assert attempt.attempt_number == 1
        assert attempt.finished_at is not None
        assert attempt.error_code is None
        assert attempt.error_message is None
    assert len(capture.requests) == 2
    assert len({request.destination for request in capture.requests}) == 2
    quote = load_workbook(result.quote_path, data_only=False)
    try:
        sheet = quote["5G手机"]
        assert [sheet["AK2"].value, sheet["AK3"].value] == [1899, 1899]
        assert _image_anchors(sheet) == {"AN2", "AN3"}
    finally:
        quote.close()
