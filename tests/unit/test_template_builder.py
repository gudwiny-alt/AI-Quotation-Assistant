from hashlib import sha256
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.drawing.image import Image as OpenpyxlImage  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]
from openpyxl.xml.functions import tostring  # type: ignore[import-untyped]
from PIL import Image as PillowImage


TEMPLATE_PATH = Path("resources/templates/quote_template.xlsx")
EXPECTED_HEADERS = (
    "终端类型（二级）",
    "品牌",
    "集团一级库物料编码",
    "集团一级库\n产品名称（型号）",
    "集团一级库\n市场通俗名称",
    "终端资源编码（B域）",
    "终端资源名称（B域）",
    "入中国移动总部库时间",
    "2026年3月结算报价（元/台）",
    "2026年7月结算报价（元/台）",
    "2026年8月结算报价（元/台）",
    "终端公司采购价",
    "终端公司全国采购系统均价",
    "货源情况（货源充足\\货源紧缺\\新品上市\\尾货期）",
    "集团一级库价格",
    "渠道买断价",
    "分销零售价",
    "京东自营或官网目前价格",
    "网址",
    "同配置机型",
    "同配置价格",
    "是否新增",
    "是否达到6个月需调价条件",
    "较上个月降价幅度",
    "6个月内降价幅度",
    "铺货价较采购价上浮比例（不高于4.5%）",
    "政企统谈价",
    "特殊情况备注",
    "采购类型\n（终端总部集采、自采）",
    "相关依据\n（总部文件及文号、合同名称及编码等）",
    "终端总部集采、省代自采合同价格",
    "总部文件/合同价格是否调价\n（时间、调整后价格）",
    "产品经理",
    "三网最低价",
    "京东自营/官旗价",
    "天猫官旗价价格",
    "品牌官网价价格",
    "京东自营/官旗价截图",
    "天猫官旗价截图",
    "品牌官网价截图",
)
EXPECTED_COLUMN_WIDTHS = {
    "A": 16.5,
    "H": 17.0803571428571,
    "K": 9.61607142857143,
    "N": 10.5089285714286,
    "X": 13.4196428571429,
    "AI": 27.5803571428571,
    "AL": 23.3303571428571,
    "AN": 13.0,
}
EXPECTED_A_TO_AN_STYLE_DIGEST = (
    "7f58e0c3f7013f61c7ebb7c69bae6d5ab8b0fb076aa10748f0735fae679f0b9f"
)


def _style_digest_for_first_two_rows(sheet: Worksheet) -> str:
    digest = sha256()
    for row in (1, 2):
        for column in range(1, 41):
            cell = sheet.cell(row, column)
            for style_component in (
                cell.font,
                cell.fill,
                cell.border,
                cell.alignment,
                cell.protection,
            ):
                digest.update(tostring(style_component.to_tree()))
            digest.update(cell.number_format.encode("utf-8"))
            digest.update(str(cell.pivotButton).encode())
            digest.update(str(cell.quotePrefix).encode())
    return digest.hexdigest()


def test_build_script_cleans_explicit_source_into_explicit_destination(
    tmp_path: Path,
) -> None:
    source = tmp_path / "dirty-source.xlsx"
    destination = tmp_path / "nested" / "clean-template.xlsx"
    image_path = tmp_path / "cell-image.png"
    PillowImage.new("RGB", (2, 2), "red").save(image_path)

    source_workbook = Workbook()
    source_sheet = source_workbook.active
    source_sheet.title = "5G手机"
    source_sheet.append(["字段A", "字段B", "字段C"])
    source_sheet.append(["示例", "=1+1", '=_xlfn.DISPIMG("dirty",1)'])
    source_sheet.append(["说明", "必须删除", "必须删除"])
    source_sheet["A2"].font = Font(name="宋体", bold=True)
    source_sheet["A2"].fill = PatternFill("solid", fgColor="FFFF00")
    source_sheet["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    source_sheet["A2"].number_format = "0.00%"
    source_sheet.column_dimensions["A"].width = 21.5
    source_sheet.row_dimensions[1].height = 31.0
    source_sheet.row_dimensions[2].height = 27.0
    source_sheet.freeze_panes = "A2"
    source_sheet.sheet_view.showGridLines = False
    source_sheet.page_setup.orientation = "landscape"
    source_sheet.print_title_rows = "1:2"
    source_sheet.print_area = "A1:C3"

    hidden_sheet = source_workbook.create_sheet("隐藏数据")
    hidden_sheet.sheet_state = "hidden"
    wps_sheet = source_workbook.create_sheet("WpsReserved_CellImgList")
    wps_sheet.sheet_state = "veryHidden"
    wps_sheet.add_image(OpenpyxlImage(image_path), "A1")
    source_workbook.save(source)
    source_workbook.close()

    subprocess.run(
        [
            sys.executable,
            "scripts/build_quote_template.py",
            str(source),
            str(destination),
        ],
        check=True,
    )

    clean_workbook = load_workbook(destination, data_only=False)
    try:
        assert clean_workbook.sheetnames == ["5G手机"]
        clean_sheet = clean_workbook["5G手机"]
        assert clean_sheet.max_row == 2
        assert [clean_sheet.cell(1, column).value for column in range(1, 4)] == [
            "字段A",
            "字段B",
            "字段C",
        ]
        assert all(clean_sheet.cell(2, column).value is None for column in range(1, 4))
        assert clean_sheet["A2"].font.name == "宋体"
        assert clean_sheet["A2"].font.bold is True
        assert clean_sheet["A2"].fill.fill_type == "solid"
        assert clean_sheet["A2"].fill.fgColor.rgb == "00FFFF00"
        assert clean_sheet["A2"].alignment.horizontal == "center"
        assert clean_sheet["A2"].alignment.wrap_text is True
        assert clean_sheet["A2"].number_format == "0.00%"
        assert clean_sheet.column_dimensions["A"].width == 21.5
        assert clean_sheet.row_dimensions[1].height == 31.0
        assert clean_sheet.row_dimensions[2].height == 27.0
        assert clean_sheet.freeze_panes == "A2"
        assert clean_sheet.sheet_view.showGridLines is False
        assert clean_sheet.page_setup.orientation == "landscape"
        assert clean_sheet.print_title_rows == "$1:$2"
        assert str(clean_sheet.print_area) == "'5G手机'!$A$1:$C$3"
        assert not clean_sheet._images
    finally:
        clean_workbook.close()

    unchanged_source = load_workbook(source, data_only=False)
    try:
        assert unchanged_source.sheetnames == [
            "5G手机",
            "隐藏数据",
            "WpsReserved_CellImgList",
        ]
        assert unchanged_source["5G手机"]["A2"].value == "示例"
        assert unchanged_source["5G手机"]["A3"].value == "说明"
    finally:
        unchanged_source.close()

    with ZipFile(destination) as archive:
        assert not any(
            part == "xl/cellimages.xml"
            or part.startswith("xl/media/")
            or part.startswith("xl/drawings/")
            for part in archive.namelist()
        )


def test_load_clean_template_returns_editable_workbook() -> None:
    from quote_app.excel import load_clean_template

    workbook = load_clean_template(TEMPLATE_PATH)
    try:
        assert workbook.read_only is False
        assert workbook["5G手机"]["A1"].value == EXPECTED_HEADERS[0]
    finally:
        workbook.close()


def test_clean_template_has_only_header_and_empty_style_row() -> None:
    workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    try:
        assert workbook.sheetnames == ["5G手机"]
        sheet = workbook["5G手机"]
        assert sheet.max_row == 2
        assert tuple(sheet.cell(1, column).value for column in range(1, 41)) == EXPECTED_HEADERS
        assert all(sheet.cell(2, column).value is None for column in range(1, 41))
        assert sheet["A3"].value is None
        assert not sheet._images
    finally:
        workbook.close()


def test_clean_template_preserves_approved_format_and_page_settings() -> None:
    workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    try:
        sheet = workbook["5G手机"]

        assert sheet.row_dimensions[1].height == 58.0
        assert sheet.row_dimensions[2].height == 38.0
        assert {
            column: sheet.column_dimensions[column].width
            for column in EXPECTED_COLUMN_WIDTHS
        } == EXPECTED_COLUMN_WIDTHS
        assert _style_digest_for_first_two_rows(sheet) == EXPECTED_A_TO_AN_STYLE_DIGEST
        assert sheet.freeze_panes == "A2"
        assert sheet.sheet_view.showGridLines is None
        assert not sheet.merged_cells.ranges
        assert sheet.page_margins.left == 0.75
        assert sheet.page_margins.right == 0.75
        assert sheet.page_margins.top == 1.0
        assert sheet.page_margins.bottom == 1.0
        assert sheet.page_setup.orientation is None
        assert sheet.page_setup.paperSize is None
        assert sheet.print_title_rows is None
        assert sheet.print_title_cols is None
    finally:
        workbook.close()


def test_clean_template_package_has_no_images_or_wps_dispimg_content() -> None:
    with ZipFile(TEMPLATE_PATH) as archive:
        package_parts = set(archive.namelist())
        assert not any(
            part == "xl/cellimages.xml"
            or part.startswith("xl/media/")
            or part.startswith("xl/drawings/")
            for part in package_parts
        )
        worksheet_xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"DISPIMG" not in worksheet_xml.upper()
