from dataclasses import replace
from hashlib import sha256
import os
from pathlib import Path
import re

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.workbook.workbook import Workbook  # type: ignore[import-untyped]
from openpyxl.xml.functions import tostring  # type: ignore[import-untyped]
import pytest

from quote_app.domain.models import QuoteMonth, QuoteRow
from quote_app.excel import quote_writer
from quote_app.excel.quote_writer import QuoteWriteRequest, write_quote_workbook


TEMPLATE_PATH = Path("resources/templates/quote_template.xlsx")
MANUAL_COLUMNS = ("K", "L", "M", "N", "P", "Q")
EXPECTED_VALIDATION_FORMULA = '"货源充足,货源紧缺,新品上市,尾货期"'


def _rows() -> list[QuoteRow]:
    return [
        QuoteRow(
            source_row_number=2,
            material_code="9101",
            cells={
                "A": "智能手机",
                "B": "小米",
                "C": "9101",
                "F": "已配置",
                "I": 3999,
                "J": 3899,
                "K": 9999,
                "L": 8888,
                "M": 7777,
                "N": "货源充足",
                "P": 6666,
                "Q": 5555,
                "T": "不得保留",
                "U": 123,
                "AB": "备注一",
                "AO": "AO一",
                "AP": "AP一",
            },
        ),
        QuoteRow(
            source_row_number=3,
            material_code="9101",
            cells={
                "A": "智能手机",
                "B": "荣耀",
                "C": "9101",
                "F": "申请配置",
                "I": "无",
                "J": "无",
                "AO": "AO二",
                "AP": "AP二",
            },
        ),
    ]


def _serialize(value: object) -> bytes:
    assert hasattr(value, "to_tree")
    return tostring(value.to_tree())


def test_writer_keeps_rows_order_formulas_decisions_and_manual_blanks(
    tmp_path: Path,
) -> None:
    template_digest = sha256(TEMPLATE_PATH.read_bytes()).hexdigest()

    output = write_quote_workbook(
        QuoteWriteRequest(
            quote_month=QuoteMonth(2026, 8),
            rows=_rows(),
            template_path=TEMPLATE_PATH,
            output_dir=tmp_path,
        )
    )

    assert output == tmp_path / "2026年08月终端供货价报价表.xlsx"
    assert sha256(TEMPLATE_PATH.read_bytes()).hexdigest() == template_digest

    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.sheetnames == ["5G手机"]
        sheet = workbook["5G手机"]
        assert sheet.max_row == 3
        assert [sheet["C2"].value, sheet["C3"].value] == ["9101", "9101"]
        assert [sheet["B2"].value, sheet["B3"].value] == ["小米", "荣耀"]
        assert sheet["AO2"].value == "AO一"
        assert sheet["AP3"].value == "AP二"

        for row_number in (2, 3):
            assert all(
                sheet[f"{column}{row_number}"].value is None
                for column in MANUAL_COLUMNS
            )
            assert sheet[f"T{row_number}"].value == "无"
            assert sheet[f"U{row_number}"].value == "无"

        assert (sheet["V2"].value, sheet["W2"].value) == ("否", "是")
        assert (sheet["V3"].value, sheet["W3"].value) == ("是", "否")
        assert {
            column: sheet[f"{column}2"].value
            for column in ("X", "Y", "Z", "AA")
        } == {
            "X": '=IF(K2="","",IF(J2="无","无",(K2-J2)/J2))',
            "Y": '=IF(K2="","",IF(I2="无","无",(K2-I2)/I2))',
            "Z": '=IF(OR(K2="",L2=""),"",(K2-L2)/L2)',
            "AA": '=IF(K2="","",K2-5)',
        }
        assert {
            column: sheet[f"{column}3"].value
            for column in ("X", "Y", "Z", "AA")
        } == {
            "X": '=IF(K3="","",IF(J3="无","无",(K3-J3)/J3))',
            "Y": '=IF(K3="","",IF(I3="无","无",(K3-I3)/I3))',
            "Z": '=IF(OR(K3="",L3=""),"",(K3-L3)/L3)',
            "AA": '=IF(K3="","",K3-5)',
        }
        assert all(
            sheet[f"{column}{row_number}"].number_format == "0.00%"
            for column in ("X", "Y", "Z")
            for row_number in (2, 3)
        )
    finally:
        workbook.close()


def test_writer_uses_normalized_material_code_as_text_instead_of_raw_cell_c(
    tmp_path: Path,
) -> None:
    row = QuoteRow(
        source_row_number=2,
        material_code="　0000 9101　",
        cells={"B": "小米", "C": "　0000 9101　"},
    )

    output = write_quote_workbook(
        QuoteWriteRequest(
            quote_month=QuoteMonth(2026, 8),
            rows=[row],
            template_path=TEMPLATE_PATH,
            output_dir=tmp_path,
        )
    )

    workbook = load_workbook(output)
    try:
        cell = workbook["5G手机"]["C2"]
        assert cell.value == "00009101"
        assert cell.number_format == "@"
        assert cell.data_type == "s"
    finally:
        workbook.close()


def test_writer_replicates_complete_a_to_ap_format_and_page_properties(
    tmp_path: Path,
) -> None:
    output = write_quote_workbook(
        QuoteWriteRequest(
            quote_month=QuoteMonth(2026, 8),
            rows=_rows(),
            template_path=TEMPLATE_PATH,
            output_dir=tmp_path,
        )
    )

    template = load_workbook(TEMPLATE_PATH, data_only=False)
    workbook = load_workbook(output, data_only=False)
    try:
        source = template["5G手机"]
        sheet = workbook["5G手机"]

        for column in range(1, 43):
            letter = source.cell(1, column).column_letter
            assert sheet.column_dimensions[letter].width == (
                source.column_dimensions[letter].width
            ), letter
            assert sheet.cell(1, column)._style == source.cell(1, column)._style
            for row_number in (2, 3):
                target_cell = sheet.cell(row_number, column)
                source_cell = source.cell(2, column)
                if letter in ("X", "Y", "Z"):
                    for component in (
                        "font",
                        "fill",
                        "border",
                        "alignment",
                        "protection",
                    ):
                        assert _serialize(getattr(target_cell, component)) == (
                            _serialize(getattr(source_cell, component))
                        ), (target_cell.coordinate, component)
                else:
                    assert target_cell._style == source_cell._style, (
                        target_cell.coordinate
                    )

        assert sheet.row_dimensions[1].height == source.row_dimensions[1].height
        assert sheet.row_dimensions[2].height == source.row_dimensions[2].height
        assert sheet.row_dimensions[3].height == source.row_dimensions[2].height
        assert sheet.freeze_panes == source.freeze_panes
        assert _serialize(sheet.sheet_view) == _serialize(source.sheet_view)
        assert _serialize(sheet.sheet_format) == _serialize(source.sheet_format)
        assert _serialize(sheet.sheet_properties) == _serialize(
            source.sheet_properties
        )
        assert _serialize(sheet.page_margins) == _serialize(source.page_margins)
        assert _serialize(sheet.page_setup) == _serialize(source.page_setup)
        assert _serialize(sheet.print_options) == _serialize(source.print_options)
        assert sheet.print_title_rows == source.print_title_rows
        assert sheet.print_title_cols == source.print_title_cols
        assert str(sheet.print_area) == str(source.print_area)
    finally:
        workbook.close()
        template.close()


def test_writer_adds_n_validation_for_every_and_only_generated_row(
    tmp_path: Path,
) -> None:
    output = write_quote_workbook(
        QuoteWriteRequest(
            quote_month=QuoteMonth(2026, 8),
            rows=_rows(),
            template_path=TEMPLATE_PATH,
            output_dir=tmp_path,
        )
    )

    workbook = load_workbook(output, data_only=False)
    try:
        validations = workbook["5G手机"].data_validations.dataValidation
        matching = [
            validation
            for validation in validations
            if validation.type == "list"
            and validation.formula1 == EXPECTED_VALIDATION_FORMULA
        ]
        assert len(matching) == 1
        assert str(matching[0].sqref) == "N2:N3"
        assert matching[0].allow_blank is True
    finally:
        workbook.close()


def test_writer_rolls_headers_across_year_boundary_without_changing_other_headers(
    tmp_path: Path,
) -> None:
    template = load_workbook(TEMPLATE_PATH, data_only=False)
    try:
        original_headers = [
            template["5G手机"].cell(1, column).value for column in range(1, 43)
        ]
    finally:
        template.close()

    output = write_quote_workbook(
        QuoteWriteRequest(
            quote_month=QuoteMonth(2026, 1),
            rows=_rows()[:1],
            template_path=TEMPLATE_PATH,
            output_dir=tmp_path,
        )
    )

    workbook = load_workbook(output, data_only=False)
    try:
        headers = [
            workbook["5G手机"].cell(1, column).value for column in range(1, 43)
        ]
        assert headers[8:11] == [
            "2025年8月结算报价（元/台）",
            "2025年12月结算报价（元/台）",
            "2026年1月结算报价（元/台）",
        ]
        assert headers[:8] == original_headers[:8]
        assert headers[11:] == original_headers[11:]
        assert headers[40:] == ["备注", "超6个月价格说明"]
    finally:
        workbook.close()


def test_writer_uses_safe_timestamped_name_when_exact_target_exists(
    tmp_path: Path,
) -> None:
    request = QuoteWriteRequest(
        quote_month=QuoteMonth(2026, 8),
        rows=_rows()[:1],
        template_path=TEMPLATE_PATH,
        output_dir=tmp_path,
    )
    first = write_quote_workbook(request)
    first_digest = sha256(first.read_bytes()).hexdigest()

    second = write_quote_workbook(request)

    assert second != first
    assert re.fullmatch(
        r"2026年08月终端供货价报价表-\d{8}-\d{6}(?:-\d+)?\.xlsx",
        second.name,
    )
    assert sha256(first.read_bytes()).hexdigest() == first_digest
    assert second.is_file()


def test_explicit_destination_is_atomically_replaced(
    tmp_path: Path,
) -> None:
    """Break caught: partial publication creates extra files or leaves stale bytes."""
    request = QuoteWriteRequest(
        quote_month=QuoteMonth(2026, 8),
        rows=_rows()[:1],
        template_path=TEMPLATE_PATH,
        output_dir=tmp_path,
    )
    destination = tmp_path / "2026年08月终端供货价报价表-处理中.xlsx"
    destination.write_bytes(b"old-snapshot")

    path = write_quote_workbook(replace(request, destination_path=destination))

    assert path == destination
    assert load_workbook(path)["5G手机"]["C2"].value == "9101"
    assert not tuple(tmp_path.glob(".quote-*.tmp.xlsx"))


def test_writer_never_overwrites_file_created_during_exclusive_publish(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    competitor_bytes = b"competitor-created-after-candidate-selection"
    raced_paths: list[Path] = []
    real_link = os.link

    def link_after_competitor_wins(source: str | Path, destination: str | Path) -> None:
        destination_path = Path(destination)
        if not raced_paths:
            destination_path.write_bytes(competitor_bytes)
            raced_paths.append(destination_path)
        real_link(source, destination)

    monkeypatch.setattr(quote_writer.os, "link", link_after_competitor_wins)

    output = write_quote_workbook(
        QuoteWriteRequest(
            quote_month=QuoteMonth(2026, 8),
            rows=_rows()[:1],
            template_path=TEMPLATE_PATH,
            output_dir=tmp_path,
        )
    )

    assert raced_paths == [tmp_path / "2026年08月终端供货价报价表.xlsx"]
    assert raced_paths[0].read_bytes() == competitor_bytes
    assert output != raced_paths[0]
    assert re.fullmatch(
        r"2026年08月终端供货价报价表-\d{8}-\d{6}\.xlsx",
        output.name,
    )
    workbook = load_workbook(output, data_only=False)
    workbook.close()
    assert list(tmp_path.glob(".*.tmp.xlsx")) == []


def test_writer_retries_multiple_exclusive_candidate_conflicts(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    destinations: list[Path] = []
    real_link = os.link

    def link_with_two_competitors(
        source: str | Path,
        destination: str | Path,
    ) -> None:
        destination_path = Path(destination)
        destinations.append(destination_path)
        if len(destinations) <= 2:
            destination_path.write_bytes(f"competitor-{len(destinations)}".encode())
        real_link(source, destination)

    monkeypatch.setattr(quote_writer.os, "link", link_with_two_competitors)

    output = write_quote_workbook(
        QuoteWriteRequest(
            quote_month=QuoteMonth(2026, 8),
            rows=_rows()[:1],
            template_path=TEMPLATE_PATH,
            output_dir=tmp_path,
        )
    )

    assert output == destinations[2]
    assert destinations[0].read_bytes() == b"competitor-1"
    assert destinations[1].read_bytes() == b"competitor-2"
    assert re.fullmatch(
        r"2026年08月终端供货价报价表-\d{8}-\d{6}-2\.xlsx",
        output.name,
    )
    assert list(tmp_path.glob(".*.tmp.xlsx")) == []


def test_writer_cleans_temporary_file_when_workbook_save_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_save(self: Workbook, filename: str | Path) -> None:
        raise OSError("injected save failure")

    monkeypatch.setattr(Workbook, "save", fail_save)

    with pytest.raises(ValueError, match="报价表无法写入输出目录"):
        write_quote_workbook(
            QuoteWriteRequest(
                quote_month=QuoteMonth(2026, 8),
                rows=_rows()[:1],
                template_path=TEMPLATE_PATH,
                output_dir=tmp_path,
            )
        )

    assert list(tmp_path.iterdir()) == []


def test_writer_cleans_temporary_file_when_exclusive_publish_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    loaded_workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    original_close = loaded_workbook.close
    close_calls = 0

    def tracked_close() -> None:
        nonlocal close_calls
        close_calls += 1
        original_close()

    monkeypatch.setattr(loaded_workbook, "close", tracked_close)
    monkeypatch.setattr(
        quote_writer,
        "load_clean_template",
        lambda path: loaded_workbook,
    )

    def fail_publish(source: str | Path, destination: str | Path) -> None:
        assert close_calls == 1
        raise PermissionError("injected publish failure")

    monkeypatch.setattr(quote_writer.os, "link", fail_publish)

    with pytest.raises(ValueError, match="报价表无法安全发布到输出目录"):
        write_quote_workbook(
            QuoteWriteRequest(
                quote_month=QuoteMonth(2026, 8),
                rows=_rows()[:1],
                template_path=TEMPLATE_PATH,
                output_dir=tmp_path,
            )
        )

    assert list(tmp_path.iterdir()) == []
    assert close_calls == 1


def test_writer_returns_published_quote_when_post_publish_cleanup_keeps_failing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    real_unlink = Path.unlink

    def fail_private_temp_unlink(
        self: Path,
        missing_ok: bool = False,
    ) -> None:
        if self.name.startswith(".quote-") and self.name.endswith(".tmp.xlsx"):
            raise OSError("persistent post-publish cleanup failure")
        real_unlink(self, missing_ok=missing_ok)

    monkeypatch.setattr(Path, "unlink", fail_private_temp_unlink)

    with pytest.warns(RuntimeWarning, match="正式报价表已生成.*临时文件"):
        output = write_quote_workbook(
            QuoteWriteRequest(
                quote_month=QuoteMonth(2026, 8),
                rows=_rows()[:1],
                template_path=TEMPLATE_PATH,
                output_dir=tmp_path,
            )
        )

    assert output == tmp_path / "2026年08月终端供货价报价表.xlsx"
    workbook = load_workbook(output)
    workbook.close()
    assert len(list(tmp_path.glob(".quote-*.tmp.xlsx"))) == 1


def test_writer_save_and_cleanup_failure_names_private_temp_without_internal_detail(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_save(self: Workbook, filename: str | Path) -> None:
        raise OSError("secret serialization failure")

    def fail_unlink(self: Path, missing_ok: bool = False) -> None:
        raise OSError("secret cleanup failure")

    monkeypatch.setattr(Workbook, "save", fail_save)
    monkeypatch.setattr(Path, "unlink", fail_unlink)

    with pytest.raises(ValueError) as caught:
        write_quote_workbook(
            QuoteWriteRequest(
                quote_month=QuoteMonth(2026, 8),
                rows=_rows()[:1],
                template_path=TEMPLATE_PATH,
                output_dir=tmp_path,
            )
        )

    temporary_files = list(tmp_path.glob(".quote-*.tmp.xlsx"))
    assert len(temporary_files) == 1
    assert str(caught.value) == f"报价表临时文件无法清理：{temporary_files[0]}"
    assert "secret serialization failure" not in str(caught.value)
    assert "secret cleanup failure" not in str(caught.value)
    assert list(tmp_path.glob("*终端供货价报价表*.xlsx")) == []


def test_writer_with_empty_rows_outputs_header_only_and_no_validation(
    tmp_path: Path,
) -> None:
    output = write_quote_workbook(
        QuoteWriteRequest(
            quote_month=QuoteMonth(2026, 8),
            rows=[],
            template_path=TEMPLATE_PATH,
            output_dir=tmp_path,
        )
    )

    workbook = load_workbook(output, data_only=False)
    try:
        sheet = workbook["5G手机"]
        assert sheet.max_row == 1
        assert sheet["A1"].value == "终端类型（二级）"
        assert len(sheet.data_validations.dataValidation) == 0
    finally:
        workbook.close()


def test_writer_reports_unusable_output_directory_as_business_error(
    tmp_path: Path,
) -> None:
    output_file = tmp_path / "not-a-directory"
    output_file.write_text("occupied", encoding="utf-8")

    with pytest.raises(ValueError, match="输出目录无法创建或使用"):
        write_quote_workbook(
            QuoteWriteRequest(
                quote_month=QuoteMonth(2026, 8),
                rows=[],
                template_path=TEMPLATE_PATH,
                output_dir=output_file,
            )
        )
