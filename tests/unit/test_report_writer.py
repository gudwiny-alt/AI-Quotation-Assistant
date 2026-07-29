from dataclasses import replace
from datetime import datetime, timezone
from decimal import Decimal
from hashlib import sha256
import os
from pathlib import Path
import re

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.workbook.workbook import Workbook  # type: ignore[import-untyped]
import pytest

from quote_app.domain.models import InputPaths, Issue, QuoteMonth, QuoteRow, WebQuery
from quote_app.evidence.models import EvidenceRecord, EvidenceState, MacCapturePolicy
from quote_app.excel import report_writer
from quote_app.excel.report_writer import ReportWriteRequest, write_execution_report
from quote_app.tasks.models import (
    BusinessOutcome,
    TaskState,
    WebsiteChannel,
    WebsiteObservationCheckpoint,
    WebsiteResult,
    WebsiteTask,
)


ORANGE = "F4B183"
RED = "F8696B"
GREEN = "70AD47"


def _request(tmp_path: Path) -> ReportWriteRequest:
    inputs = InputPaths(
        base=tmp_path / "base.xlsx",
        marketing=tmp_path / "marketing.xlsx",
        bop=tmp_path / "bop.xlsx",
        output_dir=tmp_path,
    )
    rows = [
        QuoteRow(
            2,
            "9101",
            {
                "B": "小米",
                "C": "9101",
                "E": "手机甲",
                "G": "12+256",
                "AB": "特殊情况备注甲",
                "AG": "经理甲",
            },
        ),
        QuoteRow(
            3,
            "9102",
            {
                "B": "其他品牌",
                "C": "9102",
                "E": "手机乙",
                "G": "8+128",
                "AB": "特殊情况备注乙",
                "AG": "经理乙",
            },
        ),
        QuoteRow(
            4,
            "9103",
            {"C": "9103", "AB": "特殊情况备注丙", "AG": "经理甲"},
            [Issue("MARKETING_NOT_FOUND", "未找到营销记录", False, 4)],
        ),
    ]
    return ReportWriteRequest(
        quote_month=QuoteMonth(2026, 8),
        rows=rows,
        output_dir=tmp_path,
        input_paths=inputs,
        quote_path=tmp_path / "2026年08月终端供货价报价表.xlsx",
        run_at=datetime(2026, 7, 25, 9, 30, 0),
    )


def _overview_values(path: Path) -> list[list[object]]:
    workbook = load_workbook(path, data_only=False)
    try:
        return [
            [cell.value for cell in row]
            for row in workbook["运行总览"].iter_rows()
        ]
    finally:
        workbook.close()


def _find_row(values: list[list[object]], label: str) -> list[object]:
    return next(row for row in values if row and row[0] == label)


def _website_row() -> QuoteRow:
    return QuoteRow(
        2,
        "9101",
        {"B": "HONOR", "C": "9101", "E": "Magic8", "AG": "经理甲"},
        web_query=WebQuery(
            brand="HONOR",
            model_name="Magic8",
            ram="16GB",
            storage="512GB",
            color="天青釉",
        ),
    )


def _website_tasks(row: QuoteRow, *, run_id: str = "run-1") -> tuple[WebsiteTask, ...]:
    return tuple(
        WebsiteTask(
            task_id=f"{run_id}-{channel.value}",
            run_id=run_id,
            source_row_number=row.source_row_number,
            output_row_number=2,
            material_code=row.material_code,
            brand="HONOR",
            model_name="Magic8",
            ram="16GB",
            storage="512GB",
            color="天青釉",
            channel=channel,
        )
        for channel in (
            WebsiteChannel.JD,
            WebsiteChannel.TMALL,
            WebsiteChannel.OFFICIAL,
        )
    )


def _checkpoint(
    task: WebsiteTask,
    *,
    outcome: BusinessOutcome = BusinessOutcome.PRICE_FOUND,
    price: Decimal | None = Decimal("4999"),
) -> WebsiteObservationCheckpoint:
    return WebsiteObservationCheckpoint(
        task_id=task.task_id,
        outcome=outcome,
        price=price if outcome is BusinessOutcome.PRICE_FOUND else None,
        url=f"https://{task.channel.value}.example/product",
        observed_at=datetime(2026, 7, 26, tzinfo=timezone.utc),
    )


def _capture_failure(task: WebsiteTask) -> WebsiteResult:
    return WebsiteResult(
        task_id=task.task_id,
        state=TaskState.TECHNICAL_FAILURE,
        outcome=None,
        price=None,
        url=None,
        evidence=None,
        diagnostic_path=None,
        error_code="CAPTURE_FOREGROUND",
        error_message="截图前台校验失败",
    )


def test_report_has_reconciled_overview_breakdowns_detail_and_formatting(
    tmp_path: Path,
) -> None:
    output = write_execution_report(_request(tmp_path))

    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.sheetnames == ["运行总览", "处理明细"]
        overview = workbook["运行总览"]
        detail = workbook["处理明细"]
        values = [[cell.value for cell in row] for row in overview.iter_rows()]

        assert _find_row(values, "报价月份")[1] == "2026年08月"
        assert _find_row(values, "报价总行数")[1] == 3
        assert _find_row(values, "处理完成")[1] == 0
        assert _find_row(values, "部分完成")[1] == 1
        assert _find_row(values, "处理失败")[1] == 1
        assert _find_row(values, "不支持品牌")[1] == 1
        assert _find_row(values, "待人工补充")[1] == 2

        brand_header = next(
            index for index, row in enumerate(values) if row[:6] == [
                "品牌",
                "总行数",
                "完成",
                "部分完成",
                "失败",
                "不支持",
            ]
        )
        brand_rows = values[brand_header + 1 : brand_header + 4]
        assert {tuple(row[:7]) for row in brand_rows} == {
            ("小米", 1, 0, 1, 0, 0, 1),
            ("其他品牌", 1, 0, 0, 0, 1, 1),
            ("未识别", 1, 0, 0, 1, 0, 0),
        }
        assert sum(int(row[1]) for row in brand_rows) == 3

        manager_header = next(
            index for index, row in enumerate(values) if row[:2] == ["产品经理", "总行数"]
        )
        manager_rows = values[manager_header + 1 : manager_header + 3]
        assert {tuple(row[:7]) for row in manager_rows} == {
            ("经理甲", 2, 0, 1, 1, 0, 1),
            ("经理乙", 1, 0, 0, 0, 1, 1),
        }
        assert sum(int(row[1]) for row in manager_rows) == 3

        channel_header = next(
            index for index, row in enumerate(values) if row[:4] == [
                "渠道",
                "已完成",
                "待人工补充",
                "失败",
            ]
        )
        assert values[channel_header + 1 : channel_header + 4] == [
            ["京东", 0, 3, 0, None, None, None],
            ["天猫", 0, 3, 0, None, None, None],
            ["官网", 0, 3, 0, None, None, None],
        ]

        expected_headers = [
            "产品经理",
            "品牌",
            "集团一级库物料编码",
            "产品名称",
            "配置",
            "基础表关联状态",
            "营销表关联状态",
            "BOP关联状态",
            "京东渠道状态",
            "天猫渠道状态",
            "官网渠道状态",
            "最终状态",
            "失败步骤",
            "原因",
            "建议操作",
        ]
        assert [cell.value for cell in detail[1]] == expected_headers
        assert detail.freeze_panes == "A2"
        assert detail.auto_filter.ref == "A1:O4"
        assert all(cell.alignment.wrap_text for row in detail.iter_rows() for cell in row)
        assert all(
            detail.column_dimensions[letter].width is not None
            and detail.column_dimensions[letter].width >= 12
            for letter in ("A", "B", "C", "D", "E", "L", "N", "O")
        )
        assert [detail[f"L{row}"].value for row in range(2, 5)] == [
            "部分完成",
            "不支持",
            "失败",
        ]
        assert detail["L2"].fill.fgColor.rgb[-6:] == ORANGE
        assert detail["L3"].fill.fgColor.rgb[-6:] == ORANGE
        assert detail["L4"].fill.fgColor.rgb[-6:] == RED
        assert all(
            detail[f"{column}{row}"].value == "待人工补充"
            for column in ("I", "J", "K")
            for row in range(2, 5)
        )
        assert "AI:AN" in str(detail["O2"].value)
        assert "营销" in str(detail["M4"].value)
    finally:
        workbook.close()


def test_report_manager_uses_ag_and_never_treats_ab_special_note_as_manager(
    tmp_path: Path,
) -> None:
    special_note = "这是特殊情况备注，不是产品经理；内容很长，绝不能进入经理汇总。"
    row = QuoteRow(
        2,
        "00009101",
        {
            "B": "小米",
            "C": "00009101",
            "E": "手机甲",
            "AB": special_note,
            "AG": "产品经理甲",
        },
    )

    output = write_execution_report(
        ReportWriteRequest(QuoteMonth(2026, 8), [row], tmp_path)
    )

    workbook = load_workbook(output)
    try:
        overview_values = [
            [cell.value for cell in overview_row]
            for overview_row in workbook["运行总览"].iter_rows()
        ]
        detail = workbook["处理明细"]
        assert detail["A2"].value == "产品经理甲"
        assert detail["C2"].value == "00009101"
        assert detail["C2"].number_format == "@"
        assert _find_row(overview_values, "产品经理甲")[:2] == ["产品经理甲", 1]
        assert all(
            special_note not in str(cell)
            for overview_row in overview_values
            for cell in overview_row
        )
    finally:
        workbook.close()


def test_missing_web_fields_are_source_data_partial_not_unsupported_brand(
    tmp_path: Path,
) -> None:
    row = QuoteRow(
        7,
        "9107",
        {"C": "9107", "E": "待补资料机型", "AG": "产品经理庚"},
        [
            Issue(
                "WEB_FIELDS_MISSING",
                "基础表第7行缺少网站查询字段：brand, color",
                False,
                7,
            )
        ],
    )

    output = write_execution_report(
        ReportWriteRequest(QuoteMonth(2026, 8), [row], tmp_path)
    )

    workbook = load_workbook(output)
    try:
        detail = workbook["处理明细"]
        overview = [
            [cell.value for cell in overview_row]
            for overview_row in workbook["运行总览"].iter_rows()
        ]
        assert detail["L2"].value == "部分完成"
        assert detail["M2"].value == "营销商品信息完整性"
        assert "补齐营销商品信息" in str(detail["O2"].value)
        assert _find_row(overview, "不支持品牌")[1] == 0
        assert _find_row(overview, "部分完成")[1] == 1
        assert _find_row(overview, "待人工补充")[1] == 1
    finally:
        workbook.close()


def test_report_naming_is_collision_safe_and_existing_file_is_unchanged(
    tmp_path: Path,
) -> None:
    first = write_execution_report(_request(tmp_path))
    digest = sha256(first.read_bytes()).hexdigest()

    second = write_execution_report(_request(tmp_path))

    assert second != first
    assert re.fullmatch(
        r"2026年08月报价执行报告-\d{8}-\d{6}(?:-\d+)?\.xlsx",
        second.name,
    )
    assert sha256(first.read_bytes()).hexdigest() == digest


def test_report_never_overwrites_concurrent_winner(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    competitor = b"competitor-report"
    raced: list[Path] = []
    real_link = os.link

    def link_after_competitor(source: str | Path, destination: str | Path) -> None:
        destination_path = Path(destination)
        if not raced:
            destination_path.write_bytes(competitor)
            raced.append(destination_path)
        real_link(source, destination)

    monkeypatch.setattr(report_writer.os, "link", link_after_competitor)

    output = write_execution_report(_request(tmp_path))

    assert raced == [tmp_path / "2026年08月报价执行报告.xlsx"]
    assert raced[0].read_bytes() == competitor
    assert output != raced[0]
    workbook = load_workbook(output)
    workbook.close()
    assert list(tmp_path.glob(".report-*.tmp.xlsx")) == []


def test_report_marks_fully_populated_channels_completed_and_green(
    tmp_path: Path,
) -> None:
    row = QuoteRow(
        2,
        "9101",
        {
            "B": "小米",
            "C": "9101",
            "E": "手机甲",
            "AB": "特殊情况备注",
            "AG": "经理甲",
            "AI": 3999,
            "AJ": 3998,
            "AK": 3997,
            "AL": "京东截图",
            "AM": "天猫截图",
            "AN": "官网截图",
        },
    )
    output = write_execution_report(
        ReportWriteRequest(QuoteMonth(2026, 8), [row], tmp_path)
    )

    workbook = load_workbook(output)
    try:
        detail = workbook["处理明细"]
        assert detail["L2"].value == "完成"
        assert detail["L2"].fill.fgColor.rgb[-6:] == GREEN
        assert [detail[f"{column}2"].value for column in ("I", "J", "K")] == [
            "已完成",
            "已完成",
            "已完成",
        ]
        values = [[cell.value for cell in row] for row in workbook["运行总览"].iter_rows()]
        assert _find_row(values, "处理完成")[1] == 1
        assert _find_row(values, "待人工补充")[1] == 0
    finally:
        workbook.close()


def test_report_marks_checkpoint_price_as_screenshot_pending(
    tmp_path: Path,
) -> None:
    row = _website_row()
    task = WebsiteTask(
        task_id="official-checkpoint",
        run_id="run-1",
        source_row_number=2,
        output_row_number=2,
        material_code="9101",
        brand="HONOR",
        model_name="Magic8",
        ram="16GB",
        storage="512GB",
        color="天青釉",
        channel=WebsiteChannel.OFFICIAL,
    )
    observation = WebsiteObservationCheckpoint(
        task_id=task.task_id,
        outcome=BusinessOutcome.PRICE_FOUND,
        price=Decimal("4999"),
        url="https://official.example/product",
        observed_at=datetime(2026, 7, 26, tzinfo=timezone.utc),
    )
    failure = WebsiteResult(
        task_id=task.task_id,
        state=TaskState.TECHNICAL_FAILURE,
        outcome=None,
        price=None,
        url=None,
        evidence=None,
        diagnostic_path=None,
        error_code="CAPTURE_FOREGROUND",
        error_message="截图前台校验失败",
    )

    output = write_execution_report(
        ReportWriteRequest(
            QuoteMonth(2026, 8),
            [row],
            tmp_path,
            website_tasks=(task,),
            website_results=(failure,),
            website_observations=(observation,),
            website_run=True,
        )
    )

    workbook = load_workbook(output)
    try:
        detail = workbook["处理明细"]
        assert detail["K2"].value == "价格成功（4999）；截图待补（CAPTURE_FOREGROUND）"
        assert detail["L2"].value == "部分完成"
    finally:
        workbook.close()


def test_report_all_price_checkpoints_remain_partial_and_pending(
    tmp_path: Path,
) -> None:
    row = _website_row()
    tasks = _website_tasks(row)
    observations = tuple(
        _checkpoint(task, price=Decimal(4999 - index))
        for index, task in enumerate(tasks)
    )

    output = write_execution_report(
        ReportWriteRequest(
            QuoteMonth(2026, 8),
            [row],
            tmp_path,
            website_tasks=tasks,
            website_results=tuple(_capture_failure(task) for task in tasks),
            website_observations=observations,
            website_run=True,
        )
    )

    workbook = load_workbook(output)
    try:
        detail = workbook["处理明细"]
        overview = [[cell.value for cell in row] for row in workbook["运行总览"].iter_rows()]
        assert detail["L2"].value == "部分完成"
        assert detail["O2"].value != "无需操作"
        assert [detail[f"{column}2"].value for column in ("I", "J", "K")] == [
            "价格成功（4999）；截图待补（CAPTURE_FOREGROUND）",
            "价格成功（4998）；截图待补（CAPTURE_FOREGROUND）",
            "价格成功（4997）；截图待补（CAPTURE_FOREGROUND）",
        ]
        assert _find_row(overview, "处理完成")[1] == 0
        assert _find_row(overview, "部分完成")[1] == 1
        assert _find_row(overview, "待人工补充")[1] == 1
        assert [_find_row(overview, label)[1:4] for label in ("京东", "天猫", "官网")] == [
            [0, 1, 0],
            [0, 1, 0],
            [0, 1, 0],
        ]
    finally:
        workbook.close()


def test_report_all_legal_no_checkpoints_remain_partial_and_pending(
    tmp_path: Path,
) -> None:
    row = _website_row()
    tasks = _website_tasks(row)
    observations = tuple(
        _checkpoint(task, outcome=outcome, price=None)
        for task, outcome in zip(
            tasks,
            (
                BusinessOutcome.NO_MODEL,
                BusinessOutcome.CAPACITY_UNAVAILABLE,
                BusinessOutcome.COLOR_UNAVAILABLE,
            ),
            strict=True,
        )
    )

    output = write_execution_report(
        ReportWriteRequest(
            QuoteMonth(2026, 8),
            [row],
            tmp_path,
            website_tasks=tasks,
            website_results=tuple(_capture_failure(task) for task in tasks),
            website_observations=observations,
            website_run=True,
        )
    )

    workbook = load_workbook(output)
    try:
        detail = workbook["处理明细"]
        overview = [[cell.value for cell in row] for row in workbook["运行总览"].iter_rows()]
        assert detail["L2"].value == "部分完成"
        assert detail["O2"].value != "无需操作"
        assert [detail[f"{column}2"].value for column in ("I", "J", "K")] == [
            "无（无该机型）；截图待补",
            "无（容量不可用）；截图待补",
            "无（颜色不可用）；截图待补",
        ]
        assert [_find_row(overview, label)[1:4] for label in ("京东", "天猫", "官网")] == [
            [0, 1, 0],
            [0, 1, 0],
            [0, 1, 0],
        ]
    finally:
        workbook.close()


def test_report_mixed_checkpoints_remain_partial_and_pending(
    tmp_path: Path,
) -> None:
    row = _website_row()
    tasks = _website_tasks(row)
    observations = (
        _checkpoint(tasks[0]),
        _checkpoint(
            tasks[1],
            outcome=BusinessOutcome.SOLD_OUT,
            price=None,
        ),
    )

    output = write_execution_report(
        ReportWriteRequest(
            QuoteMonth(2026, 8),
            [row],
            tmp_path,
            website_tasks=tasks,
            website_results=tuple(_capture_failure(task) for task in tasks),
            website_observations=observations,
            website_run=True,
        )
    )

    workbook = load_workbook(output)
    try:
        detail = workbook["处理明细"]
        overview = [[cell.value for cell in row] for row in workbook["运行总览"].iter_rows()]
        assert detail["L2"].value == "部分完成"
        assert detail["O2"].value != "无需操作"
        assert [_find_row(overview, label)[1:4] for label in ("京东", "天猫", "官网")] == [
            [0, 1, 0],
            [0, 1, 0],
            [0, 0, 1],
        ]
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("outcome", "expected"),
    [
        (BusinessOutcome.NO_MODEL, "无（无该机型）；截图待补"),
        (BusinessOutcome.CAPACITY_UNAVAILABLE, "无（容量不可用）；截图待补"),
        (BusinessOutcome.COLOR_UNAVAILABLE, "无（颜色不可用）；截图待补"),
        (BusinessOutcome.SOLD_OUT, "无（已售罄）；截图待补"),
    ],
)
def test_report_legal_no_checkpoint_uses_exact_pending_text(
    outcome: BusinessOutcome,
    expected: str,
    tmp_path: Path,
) -> None:
    row = _website_row()
    task = _website_tasks(row)[2]

    output = write_execution_report(
        ReportWriteRequest(
            QuoteMonth(2026, 8),
            [row],
            tmp_path,
            website_tasks=(task,),
            website_results=(_capture_failure(task),),
            website_observations=(_checkpoint(task, outcome=outcome, price=None),),
            website_run=True,
        )
    )

    workbook = load_workbook(output)
    try:
        detail = workbook["处理明细"]
        assert detail["K2"].value == expected
        assert detail["L2"].value == "部分完成"
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("tamper", "message"),
    [
        (lambda task: replace(task, output_row_number=3), "outside quotation rows"),
        (lambda task: replace(task, source_row_number=3), "does not bind"),
        (lambda task: replace(task, material_code="wrong"), "does not bind"),
        (lambda task: replace(task, storage="1TB"), "查询字段"),
    ],
)
def test_report_rejects_checkpoint_task_with_invalid_row_binding(
    tamper: object,
    message: str,
    tmp_path: Path,
) -> None:
    row = _website_row()
    task = _website_tasks(row)[2]
    invalid_task = tamper(task)  # type: ignore[operator]

    with pytest.raises(ValueError, match=message):
        write_execution_report(
            ReportWriteRequest(
                QuoteMonth(2026, 8),
                [row],
                tmp_path,
                website_tasks=(invalid_task,),
                website_observations=(_checkpoint(invalid_task),),
                website_run=True,
            )
        )


def test_report_rejects_mixed_task_run_ids_before_checkpoint_publication(
    tmp_path: Path,
) -> None:
    row = _website_row()
    first, second, _third = _website_tasks(row)
    mixed = replace(second, run_id="other-run")

    with pytest.raises(ValueError, match="share one run_id"):
        write_execution_report(
            ReportWriteRequest(
                QuoteMonth(2026, 8),
                [row],
                tmp_path,
                website_tasks=(first, mixed),
                website_observations=(_checkpoint(first), _checkpoint(mixed)),
                website_run=True,
            )
        )


def test_report_identifies_waiting_manual_verification_and_recovery_action(
    tmp_path: Path,
) -> None:
    row = _website_row()
    task = WebsiteTask(
        task_id="waiting-jd",
        run_id="run-1",
        source_row_number=2,
        output_row_number=2,
        material_code="9101",
        brand="HONOR",
        model_name="Magic8",
        ram="16GB",
        storage="512GB",
        color="天青釉",
        channel=report_writer.WebsiteChannel.JD,
    )

    output = write_execution_report(
        ReportWriteRequest(
            QuoteMonth(2026, 8),
            [row],
            tmp_path,
            website_tasks=(task,),
            website_run=True,
            waiting_task_ids=frozenset((task.task_id,)),
        )
    )

    workbook = load_workbook(output)
    try:
        detail = workbook["处理明细"]
        assert detail["I2"].value == "等待人工验证"
        assert detail["L2"].value == "部分完成"
        assert "京东：等待人工验证" in str(detail["N2"].value)
        assert detail["O2"].value == "完成人工验证后继续当前任务"
    finally:
        workbook.close()


def test_report_cleans_private_temporary_file_when_save_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_save(self: Workbook, filename: str | Path) -> None:
        raise OSError("injected report save failure")

    monkeypatch.setattr(Workbook, "save", fail_save)

    with pytest.raises(ValueError, match="执行报告无法写入输出目录"):
        write_execution_report(_request(tmp_path))

    assert list(tmp_path.glob(".report-*.tmp.xlsx")) == []


def test_report_cleans_private_temporary_file_when_publish_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_link(source: str | Path, destination: str | Path) -> None:
        raise OSError("injected report publication failure")

    monkeypatch.setattr(report_writer.os, "link", fail_link)

    with pytest.raises(ValueError, match="执行报告无法安全发布到输出目录"):
        write_execution_report(_request(tmp_path))

    assert list(tmp_path.glob(".report-*.tmp.xlsx")) == []
    assert list(tmp_path.glob("*报价执行报告*.xlsx")) == []


def test_report_returns_published_path_when_post_publish_cleanup_keeps_failing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    real_unlink = Path.unlink

    def fail_private_temp_unlink(
        self: Path,
        missing_ok: bool = False,
    ) -> None:
        if self.name.startswith(".report-") and self.name.endswith(".tmp.xlsx"):
            raise OSError("persistent post-publish cleanup failure")
        real_unlink(self, missing_ok=missing_ok)

    monkeypatch.setattr(Path, "unlink", fail_private_temp_unlink)

    with pytest.warns(RuntimeWarning, match="正式执行报告已生成.*临时文件"):
        output = write_execution_report(_request(tmp_path))

    assert output == tmp_path / "2026年08月报价执行报告.xlsx"
    workbook = load_workbook(output)
    workbook.close()
    assert len(list(tmp_path.glob(".report-*.tmp.xlsx"))) == 1


def test_report_marks_only_authorized_beta_evidence_for_mac_visual_review(
    tmp_path: Path,
) -> None:
    evidence = EvidenceRecord(
        state=EvidenceState.NORMAL,
        path=tmp_path / "formal.png",
        sha256="a" * 64,
        pixel_width=800,
        pixel_height=600,
        captured_at=datetime(2026, 7, 27, tzinfo=timezone.utc),
        validation_code="CAPTURE_OK_MAC_VISUAL_REVIEW",
    )
    result = WebsiteResult(
        task_id="mac-beta",
        state=TaskState.SUCCEEDED,
        outcome=BusinessOutcome.PRICE_FOUND,
        price=Decimal("4499"),
        url="https://example.test/product",
        evidence=evidence,
        diagnostic_path=None,
        error_code=None,
        error_message=None,
    )

    with pytest.raises(ValueError, match="capture validation"):
        report_writer._website_channel_state(
            result,
            capture_acceptance_policy=MacCapturePolicy.STRICT,
        )

    assert report_writer._website_channel_state(
        result,
        capture_acceptance_policy=MacCapturePolicy.MAC_VISUAL_REVIEW_BETA,
    ) == "价格成功（4499）；截图成功（Mac 视觉复核）"


def test_report_save_and_cleanup_failure_names_private_temp_without_internal_detail(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_save(self: Workbook, filename: str | Path) -> None:
        raise OSError("secret serialization failure")

    def fail_unlink(self: Path, missing_ok: bool = False) -> None:
        raise OSError("secret cleanup failure")

    monkeypatch.setattr(Workbook, "save", fail_save)
    monkeypatch.setattr(Path, "unlink", fail_unlink)

    with pytest.raises(ValueError) as caught:
        write_execution_report(_request(tmp_path))

    temporary_files = list(tmp_path.glob(".report-*.tmp.xlsx"))
    assert len(temporary_files) == 1
    assert str(caught.value) == f"执行报告临时文件无法清理：{temporary_files[0]}"
    assert "secret serialization failure" not in str(caught.value)
    assert "secret cleanup failure" not in str(caught.value)
    assert list(tmp_path.glob("*报价执行报告*.xlsx")) == []
