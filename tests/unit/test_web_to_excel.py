from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import os

from openpyxl import load_workbook
from PIL import Image
import pytest

from quote_app.domain.models import QuoteMonth
from quote_app.evidence.models import MacCapturePolicy
from quote_app.services import web_to_excel
from quote_app.services.web_to_excel import (
    WebToExcelRequest,
    write_web_results_to_excel,
)
from tests.factories.web_run_factory import make_business_result, make_quote_row, make_tasks


def test_excel_thumbnail_preserves_high_resolution_for_readable_zoom() -> None:
    """Catches shrinking an Excel evidence image below the approved HD size."""
    source = BytesIO()
    Image.new("RGB", (2880, 1800), "white").save(source, format="PNG")

    payload = web_to_excel._excel_thumbnail(source.getvalue())

    with Image.open(BytesIO(payload)) as thumbnail:
        assert thumbnail.format == "JPEG"
        assert thumbnail.size == (1920, 1200)
    assert len(payload) <= 1536 * 1024


def test_default_excel_publication_rejects_mac_visual_review_evidence(tmp_path) -> None:
    """Catches publishing beta evidence from the ordinary batch entry point."""
    row = make_quote_row()
    task = make_tasks((row,))[0]
    result = make_business_result(tmp_path, task)
    assert result.evidence is not None
    beta_result = replace(
        result,
        evidence=replace(
            result.evidence,
            validation_code="CAPTURE_OK_MAC_VISUAL_REVIEW",
        ),
    )
    request = WebToExcelRequest(
        quote_month=QuoteMonth(2026, 8),
        rows=(row,),
        tasks=(task,),
        results=(beta_result,),
        output_dir=tmp_path,
        template_path="resources/templates/quote_template.xlsx",
    )

    with pytest.raises(ValueError, match="capture validation"):
        write_web_results_to_excel(request)


def test_explicit_mac_visual_review_publication_embeds_beta_evidence(
    tmp_path,
) -> None:
    """Catches beta evidence being lost after the explicit Darwin-only release gate."""
    row = make_quote_row()
    task = make_tasks((row,))[0]
    result = make_business_result(tmp_path, task)
    assert result.evidence is not None
    beta_result = replace(
        result,
        evidence=replace(
            result.evidence,
            validation_code="CAPTURE_OK_MAC_VISUAL_REVIEW",
        ),
    )
    request = WebToExcelRequest(
        quote_month=QuoteMonth(2026, 8),
        rows=(row,),
        tasks=(task,),
        results=(beta_result,),
        output_dir=tmp_path,
        template_path="resources/templates/quote_template.xlsx",
        capture_acceptance_policy=MacCapturePolicy.MAC_VISUAL_REVIEW_BETA,
    )

    output = write_web_results_to_excel(request)

    quote = load_workbook(output.quote_path)
    report = load_workbook(output.report_path)
    try:
        assert len(quote["5G手机"]._images) == 1
        assert report["处理明细"]["I2"].value == (
            "价格成功（4399）；截图成功（Mac 视觉复核）"
        )
    finally:
        quote.close()
        report.close()


def test_non_darwin_request_rejects_the_typed_mac_visual_review_policy(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path,
) -> None:
    """Catches a Windows batch run opting into the Mac-only release context."""
    from quote_app.services import web_to_excel

    monkeypatch.setattr(web_to_excel.host_platform, "system", lambda: "Windows")

    with pytest.raises(ValueError, match="only on Darwin"):
        WebToExcelRequest(
            quote_month=QuoteMonth(2026, 8),
            rows=(),
            tasks=(),
            results=(),
            output_dir=tmp_path,
            capture_acceptance_policy=MacCapturePolicy.MAC_VISUAL_REVIEW_BETA,
        )


def test_evidence_cache_reuses_unchanged_image_but_revalidates_changed_file(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path,
) -> None:
    """Break caught: every incremental snapshot rehashes all old screenshots."""
    from quote_app.services import web_to_excel

    row = make_quote_row()
    task = make_tasks((row,))[0]
    result = make_business_result(tmp_path, task)
    assert result.evidence is not None
    cache = web_to_excel.EvidenceValidationCache()
    reads = 0
    original = web_to_excel.read_validated_evidence

    def read_once(*args, **kwargs):
        nonlocal reads
        reads += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(web_to_excel, "read_validated_evidence", read_once)

    first = web_to_excel._prepare_evidence(
        (result,),
        capture_acceptance_policy=MacCapturePolicy.STRICT,
        evidence_cache=cache,
    )
    second = web_to_excel._prepare_evidence(
        (result,),
        capture_acceptance_policy=MacCapturePolicy.STRICT,
        evidence_cache=cache,
    )
    assert reads == 1
    assert first == second

    result.evidence.path.write_bytes(b"changed screenshot")
    _payloads, report_results = web_to_excel._prepare_evidence(
        (result,),
        capture_acceptance_policy=MacCapturePolicy.STRICT,
        evidence_cache=cache,
    )

    assert reads == 2
    assert report_results[0].state.value == "technical_failure"


def test_evidence_changed_after_validated_read_is_not_associated_with_cache(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path,
) -> None:
    """Break caught: a swapped file identity is cached with previously read bytes."""
    from quote_app.services import web_to_excel

    row = make_quote_row()
    task = make_tasks((row,))[0]
    result = make_business_result(tmp_path, task)
    assert result.evidence is not None
    cache = web_to_excel.EvidenceValidationCache()
    original = web_to_excel.read_validated_evidence

    def swap_after_read(*args, **kwargs):
        audit = original(*args, **kwargs)
        result.evidence.path.unlink()
        result.evidence.path.write_bytes(b"replacement")
        return audit

    monkeypatch.setattr(
        web_to_excel,
        "read_validated_evidence",
        swap_after_read,
    )

    payloads, report_results = web_to_excel._prepare_evidence(
        (result,),
        capture_acceptance_policy=MacCapturePolicy.STRICT,
        evidence_cache=cache,
    )

    assert payloads == {}
    assert report_results[0].state.value == "technical_failure"
    assert cache.total_payload_bytes == 0


def test_large_formal_screenshot_is_embedded_as_bounded_readable_thumbnail(
    tmp_path,
) -> None:
    """Break caught: 900 full-screen originals consume multi-gigabyte RAM/XLSX space."""
    from quote_app.services import web_to_excel

    row = make_quote_row()
    task = make_tasks((row,))[0]
    result = make_business_result(tmp_path, task)
    assert result.evidence is not None
    noisy = Image.frombytes("RGB", (1800, 1200), os.urandom(1800 * 1200 * 3))
    noisy.save(result.evidence.path, format="PNG")
    payload = result.evidence.path.read_bytes()
    result = replace(
        result,
        evidence=replace(
            result.evidence,
            sha256=sha256(payload).hexdigest(),
            pixel_width=1800,
            pixel_height=1200,
        ),
    )
    cache = web_to_excel.EvidenceValidationCache()

    payloads, report_results = web_to_excel._prepare_evidence(
        (result,),
        capture_acceptance_policy=MacCapturePolicy.STRICT,
        evidence_cache=cache,
    )

    thumbnail = payloads[result.task_id]
    assert report_results == (result,)
    assert len(thumbnail) <= 1536 * 1024
    assert cache.total_payload_bytes == len(thumbnail)
    with Image.open(BytesIO(thumbnail)) as image:
        assert image.format == "JPEG"
        assert image.width <= 1920
        assert image.height <= 1200
