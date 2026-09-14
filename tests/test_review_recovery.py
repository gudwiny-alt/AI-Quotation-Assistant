"""Regression coverage for editable evidence and independent batch progress."""

from dataclasses import replace
from zipfile import ZipFile
import pytest
from PIL import Image
from tests.test_review_support import session as base_session
from quote_app.services.review_workbook import digest
from quote_app.services.review_sources import workbook_channels
from quote_app.services.review_support import ReviewSession


@pytest.fixture
def session(tmp_path):
    return base_session.__wrapped__(tmp_path)


def test_repeated_save_keeps_workbook_bytes_and_one_backup(session):
    p = session.products[0]
    assert not session.is_written(p)
    session.save(p)
    first = session.quote_path.read_bytes()
    assert session.is_written(p)
    p.values["K"] = "2090.00"
    session.save(p)
    assert session.quote_path.read_bytes() == first
    assert len(list((session.quote_path.parent / "历史备份").glob("*.bak"))) == 1
    assert not list(session.quote_path.parent.glob("*.bak"))
    p.values["K"] = "2091"
    assert not session.is_written(p)


def test_owned_save_preserves_other_confirmations_but_not_external_changes(session):
    p = session.products[0]
    # Independent product fingerprint, not a second copy of the edited product.
    q = replace(p, id="other", values=dict(p.values), context=dict(p.context), channels=[])
    session.products.append(q)
    session._sources[q.id] = session._sources[p.id]
    session._source_known[q.id] = True
    session._confirmed[q.id] = session._fingerprint(q)
    session._reviews["other-review"] = {"product_id": q.id, "version": session._fingerprint(q)}
    session.save(p)
    assert session._confirmed[q.id] == session._fingerprint(q)
    assert session._reviews["other-review"]["version"] == session._fingerprint(q)
    session.quote_path.write_bytes(session.quote_path.read_bytes() + b"external")
    assert session._confirmed[q.id] != session._fingerprint(q)


def test_reviewer_defaults_survive_reload_and_history_is_immutable(session):
    p = session.products[0]
    session.set_reviewer("张三")
    assert session.reviewer_for(p) == "张三"
    session.set_reviewer("李四", p)
    session._history.append({"operator": "张三"})
    session.set_reviewer("王五")
    assert session.reviewer_for(p) == "李四"
    assert session._history[-1]["operator"] == "张三"
    restored = ReviewSession(session.model, session.month)
    assert restored.reviewer_for(restored.products[0]) == "李四"
    assert restored.batch_reviewer == "王五"


def test_replace_missing_and_existing_image_preserves_cells_and_other_channels(session, tmp_path):
    p = session.products[0]
    for channel, color in [("official", "red"), ("jd", "green"), ("official", "blue")]:
        image = tmp_path / f"{color}.png"
        Image.new("RGB", (640, 480), color).save(image)
        session.replace_evidence(p, channel, image, "商品页重新截图")
        channels = workbook_channels(session.quote_path, session._rows)
        record = next(c for c in channels if c.channel == channel and c.source_row_number == 2)
        assert digest(record.evidence_path) == digest(image)
        # Worksheet cells remain unchanged (only first drawing reference may be added).
        from quote_app.services.review_workbook import inspect

        assert inspect(session.quote_path, session.month)[2].get("K") is None
        if color == "blue":
            jd = next(c for c in channels if c.channel == "jd")
            assert digest(jd.evidence_path) == digest(tmp_path / "green.png")
            assert len([c for c in channels if c.channel == "official"]) == 1
    restored = ReviewSession.from_workbook(session.quote_path, session.month)
    t = next(t for t in restored.products[0].channels if t.channel == "official")
    assert digest(t.evidence_path) == digest(tmp_path / "blue.png")
    assert len([h for h in restored._history if h.get("action") == "replace_evidence"]) == 3
    assert (tmp_path / "red.png").is_file()


def test_evidence_transaction_rolls_back_on_sidecar_failure(session, tmp_path, monkeypatch):
    image = tmp_path / "test.png"
    Image.new("RGB", (400, 300), "white").save(image)
    before = session.quote_path.read_bytes()
    monkeypatch.setattr(session, "_persist", lambda: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        session.replace_evidence(session.products[0], "official", image, "重新截图")
    assert session.quote_path.read_bytes() == before
    assert not session._replacements


def test_invalid_image_and_stale_workbook_cannot_replace(session, tmp_path):
    image = tmp_path / "invalid.png"
    image.write_text("not a screenshot")
    before = session.quote_path.read_bytes()
    with pytest.raises(ValueError):
        session.replace_evidence(session.products[0], "official", image, "重新截图")
    assert session.quote_path.read_bytes() == before
    Image.new("RGB", (400, 300)).save(image)
    session.quote_path.write_bytes(before + b"changed")
    with pytest.raises(ValueError):
        session.replace_evidence(session.products[0], "official", image, "重新截图")


def test_confirm_two_real_rows_sequentially_and_noop_save_retains_both(session):
    from openpyxl import load_workbook
    from quote_app.desktop_state import DesktopState, TaskRow

    book = load_workbook(session.quote_path)
    sheet = book["5G手机"]
    for col in ("D", "E", "O"):
        sheet[f"{col}3"] = sheet[f"{col}2"].value
    sheet["C3"] = "002"
    book.save(session.quote_path)
    original = session.model.quote_rows[0]
    other = replace(original, source_row_number=3, material_code="002")
    obj = ReviewSession(
        DesktopState(quote_rows=(original, other), quote_path=session.quote_path), session.month
    )
    for p in obj.products:
        p.values.update(session.products[0].values)
        p.context.update(session.products[0].context, previous_price="3000")
        obj._sources[p.id].cells["J"] = 3000
        p.channels = [
            TaskRow(
                p.id,
                "Model",
                "official",
                price="3000",
                outcome="price_found",
                state="succeeded",
                source_row_number=p.output_row,
                material_code=p.material_code,
            )
        ]
        obj.confirm_product(p)
    assert all(obj._confirmed[p.id] == obj._fingerprint(p) for p in obj.products)
    backups = list((obj.quote_path.parent / "历史备份").glob("*.bak"))
    obj.confirm_product(obj.products[0])
    assert all(obj._confirmed[p.id] == obj._fingerprint(p) for p in obj.products)
    assert list((obj.quote_path.parent / "历史备份").glob("*.bak")) == backups
    # Editing either product changes only its valid confirmation.
    obj.products[1].values["K"] = "2080"
    assert obj._confirmed[obj.products[0].id] == obj._fingerprint(obj.products[0])
    assert obj._confirmed[obj.products[1].id] != obj._fingerprint(obj.products[1])


def test_existing_drawings_and_all_cells_are_unchanged_except_target_image(session, tmp_path):
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage

    image = tmp_path / "initial.png"
    Image.new("RGB", (500, 400), "red").save(image)
    book = load_workbook(session.quote_path)
    for anchor in ("AL2", "AM2", "AN2", "AN3"):
        book["5G手机"].add_image(XLImage(image), anchor)
    book.save(session.quote_path)
    obj = ReviewSession(session.model, session.month)
    replacement = tmp_path / "replacement.jpeg"
    Image.new("RGB", (550, 400), "blue").save(replacement)
    with ZipFile(obj.quote_path) as z:
        before = {n: z.read(n) for n in z.namelist()}
    obj.replace_evidence(obj.products[0], "official", replacement, "消除遮挡")
    with ZipFile(obj.quote_path) as z:
        changed = {n for n in before if z.read(n) != before[n]}
    assert changed <= {
        "xl/drawings/drawing1.xml",
        "xl/drawings/_rels/drawing1.xml.rels",
        "[Content_Types].xml",
    }
    book = load_workbook(obj.quote_path)
    photos = book["5G手机"]._images
    assert len(photos) == 4
    import hashlib

    for photo in photos:
        target = photo.anchor._from.col == 39 and photo.anchor._from.row == 1
        assert hashlib.sha256(photo._data()).hexdigest() == digest(replacement if target else image)


def test_noop_save_rollback_does_not_require_backup(session, monkeypatch):
    p = session.products[0]
    session.save(p)
    before = session.quote_path.read_bytes()
    monkeypatch.setattr(session, "_persist", lambda: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        session.save(p)
    assert session.quote_path.read_bytes() == before


def test_manual_audit_keeps_quotation_confirmation_but_failed_review_blocks_final(
    session, tmp_path
):
    from quote_app.desktop_state import TaskRow

    p = session.products[0]
    image = tmp_path / "tiny.png"
    Image.new("RGB", (10, 10)).save(image)
    p.channels.append(TaskRow("jd", channel="jd", evidence_path=image))
    c = next(c for c in session.evaluate(p) if c.code == "C01")
    session._confirmed[p.id] = session._fingerprint(p)
    session.review(c.id, "未通过", "图片不能看清", "张三")
    assert not session.unconfirmed_products()
    with pytest.raises(ValueError, match="只能导出草稿"):
        session.export_report(final=True)
    session.review(c.id, "通过", "查看原件，规格正确", "张三")
    assert not session.unconfirmed_products()


def test_failed_confirmation_record_is_not_marked_confirmed_in_memory(session, monkeypatch):
    from quote_app.desktop_state import TaskRow

    p = session.products[0]
    session._sources[p.id].cells["J"] = 2090
    p.channels = [
        TaskRow("jd", channel="jd", price="2200", state="succeeded", outcome="price_found")
    ]
    persist = session._persist
    calls = []

    def fail_second():
        calls.append(True)
        if len(calls) == 2:
            raise OSError("confirmation write failed")
        persist()

    monkeypatch.setattr(session, "_persist", fail_second)
    with pytest.raises(OSError):
        session.confirm_product(p)
    assert session.is_written(p)
    assert p in session.unconfirmed_products()
