from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from quote_app.domain.models import QuoteMonth
from quote_app.services.review_support import ReviewSession

MONTH = QuoteMonth(2026, 8)


def make_quote(tmp_path):
    book = load_workbook("resources/templates/quote_template.xlsx")
    sheet = book["5G手机"]
    for col, value in {"K": "2026年8月结算报价（元/台）"}.items():
        sheet[f"{col}1"] = value
    for col, value in {
        "A": "5G手机",
        "C": "001",
        "D": "HONOR_TEST_荣耀Magic8_16GB+512GB_天青釉_标准版",
        "E": "荣耀Magic8",
        "H": "2025-10-14",
        "I": 4459,
        "J": 4459,
        "K": 4400,
        "L": 4200,
        "M": 4400,
        "Q": 4999,
        "O": "4500到5000",
        "AI": 4999,
        "AJ": 5499,
        "AK": 5499,
    }.items():
        sheet[f"{col}2"] = value
    image = tmp_path / "evidence.png"
    Image.new("RGB", (40, 30), "blue").save(image)
    sheet.add_image(ExcelImage(image), "AL2")
    path = tmp_path / "2026年08月终端供货价报价表.xlsx"
    book.save(path)
    return path


def test_reopen_restores_channel_prices_and_embedded_evidence(tmp_path):
    path = make_quote(tmp_path)
    before = path.read_bytes()
    session = ReviewSession.from_workbook(path, MONTH)
    product = session.products[0]
    assert product.title == "荣耀Magic8"
    assert {c.channel: c.price for c in product.channels} == {
        "jd": "4999",
        "tmall": "5499",
        "official": "5499",
    }
    jd = next(c for c in product.channels if c.channel == "jd")
    assert jd.evidence_path.is_file()
    assert Image.open(jd.evidence_path).size == (40, 30)
    assert path.read_bytes() == before


def test_phone_rule_works_without_manual_category_confirmation(tmp_path):
    session = ReviewSession.from_workbook(make_quote(tmp_path), MONTH)
    product = session.products[0]
    check = next(c for c in session.evaluate(product) if c.code == "E02")
    assert check.status == "未通过"
    assert "4389" in check.comparison


def test_marketing_stock_and_entry_are_automatic_and_range_uses_upper_end(tmp_path):
    quote = make_quote(tmp_path)
    book = Workbook()
    sheet = book.active
    sheet.append(
        [
            "物料编码",
            "产品状态",
            "入库时间",
            "建议市场零售价",
            "市场通俗名称",
            "RAM大小(MB/GB)",
            "ROM大小(MB/GB)",
            "颜色",
        ]
    )
    sheet.append(
        ["001", "在库", "2025-10-14", "4500到5000", "荣耀Magic8", "16GB", "512GB", "天青釉"]
    )
    marketing = tmp_path / "营销.xlsx"
    book.save(marketing)
    session = ReviewSession.from_workbook(quote, MONTH, source_paths={"marketing": marketing})
    product = session.products[0]
    assert product.context["stock"] == "在库"
    assert product.context["entry_date"] == "2025-10-14"
    checks = {c.code: c for c in session.evaluate(product)}
    assert checks["E05"].status == "通过"
    assert "5000" in checks["E04"].comparison
    assert "4500到5000" in checks["E04"].reason
    assert "None" not in checks["E04"].comparison
    assert checks["E04"].status != "通过"


def test_warehouse_limit_uses_exact_upper_end():
    from quote_app.services.review_sources import warehouse_limit
    from decimal import Decimal

    assert warehouse_limit("4500到5000") == Decimal("5000")
    assert warehouse_limit("1500到2000") == Decimal("2000")
    assert warehouse_limit("1999.01至2099.005") == Decimal("2099.005")
    assert warehouse_limit("5000到4500") is None
    assert warehouse_limit("2000以上") is None


def test_history_reads_all_supplied_past_months(tmp_path):
    quote = make_quote(tmp_path)
    book = Workbook()
    sheet = book.active
    sheet.append(
        [
            "集团一级库物料编码",
            "2026年3月结算报价（元/台）",
            "2026年4月结算报价（元/台）",
            "2026年7月结算报价（元/台）",
            "2026年9月结算报价（元/台）",
        ]
    )
    sheet.append(["001", 4459, 4300, 4459, 100])
    path = tmp_path / "基础.xlsx"
    book.save(path)
    session = ReviewSession.from_workbook(quote, MONTH, source_paths={"base": path})
    check = next(c for c in session.evaluate(session.products[0]) if c.code == "E03")
    assert check.status == "未通过"
    assert "4300" in check.comparison
    assert "2026-04" in check.reason
    assert "2026-09" not in check.reason


def test_provenance_requires_exact_thumbnail_and_preserves_failure(tmp_path, monkeypatch):
    from quote_app.services import review_sources as sources
    from quote_app.services.web_to_excel import _excel_thumbnail
    from quote_app.desktop_state import TaskRow

    original = tmp_path / "original.png"
    Image.new("RGB", (900, 600), "#1255aa").save(original)
    thumb = tmp_path / "thumb.jpg"
    thumb.write_bytes(_excel_thumbnail(original.read_bytes()))
    database = tmp_path / "tasks.db"
    database.touch()
    saved = TaskRow(
        "real-task",
        channel="jd",
        material_code="001",
        price="4999",
        evidence_path=original,
        state="succeeded",
        outcome="price_found",
    )
    failure = TaskRow("failed", channel="tmall", material_code="001", state="technical_failure")
    monkeypatch.setattr(sources, "read_task_rows", lambda *a: ([saved, failure], ""))
    observed = TaskRow(
        "workbook:2:jd",
        channel="jd",
        material_code="001",
        price="4999",
        evidence_path=thumb,
        source_row_number=2,
    )
    restored = sources.recover_task_provenance([observed], database, run_id="run")
    assert [x.task_id for x in restored] == ["real-task", "failed"]
    assert restored[0].evidence_path == original
    observed.price = "5000"
    assert sources.recover_task_provenance([observed], database, run_id="run") == [observed]


def test_missing_file_cannot_gain_task_provenance(tmp_path, monkeypatch):
    from quote_app.services import review_sources as sources
    from quote_app.desktop_state import TaskRow

    database = tmp_path / "tasks.db"
    database.touch()
    task = TaskRow("no-image", channel="jd", material_code="001")
    monkeypatch.setattr(sources, "read_task_rows", lambda *a: ([task], ""))
    observed = TaskRow("workbook:2:jd", channel="jd", material_code="001", source_row_number=2)
    assert sources.recover_task_provenance([observed], database, run_id="run") == [observed]
