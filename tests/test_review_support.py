from pathlib import Path
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook
from quote_app.desktop_state import DesktopState, TaskRow
from quote_app.domain.models import QuoteMonth, QuoteRow, WebQuery
from quote_app.services.review_support import ReviewSession, can_confirm, CATEGORIES

MONTH = QuoteMonth(2026, 9)


@pytest.fixture
def session(tmp_path):
    template = Path("resources/templates/quote_template.xlsx")
    w = load_workbook(template)
    s = w["5G手机"]
    s["K1"] = "2026年9月结算报价（元/台）"
    s["C2"] = "001"
    s["D2"] = "Model"
    s["E2"] = "12 / 256 / 黑"
    s["O2"] = 3000
    s["Z2"] = "=K2/L2-1"
    path = tmp_path / "报价.xlsx"
    w.save(path)
    row = QuoteRow(
        2,
        "001",
        {"D": "Model", "E": "12 / 256 / 黑", "O": 3000},
        web_query=WebQuery(model_name="Model", ram="12", storage="256", color="黑"),
    )
    model = DesktopState(quote_rows=(row,), quote_path=path, run_id="test")
    obj = ReviewSession(model, MONTH)
    p = obj.products[0]
    p.values.update(K="2090", L="2000", M="2000", P="1", Q="3000")
    p.context.update(
        category="手机",
        stock="在库",
        entry_date="2025-01-01",
        first_quote_date="2026-08-01",
        youfu="否",
        qualification_note="产品经理提供库存凭据",
    )
    return obj


def check(session, code):
    return next(c for c in session.evaluate(session.products[0]) if c.code == code)


def test_precise_boundary_and_purchase_mean(session):
    assert check(session, "E02").status == "通过"
    assert check(session, "E01").status == "通过"
    session.products[0].values["K"] = "2090.0001"
    assert check(session, "E02").status == "未通过"
    assert check(session, "E01").status == "通过"
    for value in ("NaN", "Infinity", "-1", "0"):
        session.products[0].values["L"] = value
        assert check(session, "E02").status == "待补充"


def test_fourteen_months_requires_confirmation(session):
    p = session.products[0]
    p.context["youfu"] = "未确认"
    assert check(session, "E06").status == "待复核"
    p.context["youfu"] = "是"
    assert check(session, "E06").status == "不适用"
    assert not can_confirm(session.evaluate(p))
    p.context["youfu"] = "否"
    assert check(session, "E06").status == "通过"


def test_unknown_policy_never_passes_and_categories_conserve(session):
    checks = session.all_checks()
    assert len(checks) == len({c.id for c in checks})
    assert set(c.category for c in checks) == set(CATEGORIES)
    assert check(session, "E03").status != "通过"
    assert not can_confirm(checks)
    assert not can_confirm([])
    with pytest.raises(ValueError):
        session.export_report(final=True)


def test_writeback_preserves_unrelated_zip_parts_and_formulas(session):
    path = session.quote_path
    with ZipFile(path) as z:
        before = {n: z.read(n) for n in z.namelist()}
    session.products[0].values["AO"] = "=not a formula <说明>"
    assert session.save(session.products[0]) == path
    with ZipFile(path) as z:
        after = {n: z.read(n) for n in z.namelist()}
    changed = [n for n in before if before[n] != after[n]]
    assert changed == ["xl/worksheets/sheet1.xml"]
    w = load_workbook(path)
    s = w["5G手机"]
    assert s["K2"].value == 2090
    assert s["Z2"].value == "=K2/L2-1"
    assert s["AO2"].value == "=not a formula <说明>" and s["AO2"].data_type == "s"
    assert list(path.parent.glob("*.bak"))
    restored = ReviewSession(session.model, MONTH)
    assert restored.products[0].context["stock"] == "在库"
    assert restored.products[0].values["K"] == "2090"


def test_conflict_and_identity_cannot_overwrite(session):
    path = session.quote_path
    original = path.read_bytes()
    path.write_bytes(original + b"changed")
    with pytest.raises(ValueError, match="版本"):
        session.save(session.products[0])
    assert path.read_bytes() == original + b"changed"


def test_live_row_identity_is_checked(session):
    session.products[0].material_code = "other"
    before = session.quote_path.read_bytes()
    with pytest.raises(ValueError):
        session.save(session.products[0])
    assert before == session.quote_path.read_bytes()


def test_review_requires_evidence_and_invalidates_on_change(session, tmp_path):
    from PIL import Image

    image = tmp_path / "proof.png"
    Image.new("RGB", (10, 10)).save(image)
    p = session.products[0]
    p.channels.append(
        TaskRow(
            "t",
            model_name="Model",
            channel="jd",
            price="2200",
            outcome="price_found",
            state="succeeded",
            evidence_path=image,
            url="https://item.jd.com/1.html",
        )
    )
    c = check(session, "C01")
    assert c.status == "待复核"
    session.review(c.id, "通过", "核对页面正常", "张三")
    assert check(session, "C01").status == "通过"
    p.values["K"] = "2091"
    assert check(session, "C01").status == "待复核"
    with pytest.raises(ValueError):
        session.review(check(session, "E02").id, "通过", "忽略", "张三")


def test_same_model_different_spec_and_codes_are_distinct(session):
    r = QuoteRow(
        3,
        "002",
        {"D": "Model", "E": "16 / 512 / 白"},
        web_query=WebQuery(model_name="Model", ram="16", storage="512", color="白"),
    )
    session.model.quote_rows += (r,)
    s = ReviewSession(session.model, MONTH)
    assert len(s.products) == 2 and len({p.id for p in s.products}) == 2


def test_report_escapes_untrusted_values(session):
    session.products[0].title = "<script>alert(1)</script>"
    path = session.export_report()
    text = path.read_text()
    assert "<script>alert" not in text and "&lt;script&gt;" in text
    assert "草稿" in text and "未解决" in text


def test_from_workbook_reopens_manual_values(session):
    session.save(session.products[0])
    opened = ReviewSession.from_workbook(session.quote_path, MONTH)
    assert opened.products[0].values["K"] == "2090"
    assert opened.products[0].material_code == "001"
    assert not opened.products[0].channels
    assert check(opened, "C01").status == "未检查"


def test_safe_save_rejects_invalid_numeric_without_touching_file(session):
    before = session.quote_path.read_bytes()
    session.products[0].values["L"] = "NaN"
    with pytest.raises(ValueError):
        session.save(session.products[0])
    assert session.quote_path.read_bytes() == before


def test_replacement_failure_retains_original(session, monkeypatch):
    import quote_app.services.review_workbook as rw

    before = session.quote_path.read_bytes()

    def fail(*args):
        raise OSError("disk unavailable")

    monkeypatch.setattr(rw.os, "replace", fail)
    with pytest.raises(OSError):
        session.save(session.products[0])
    assert session.quote_path.read_bytes() == before


def test_same_code_wrong_spec_source_cannot_save(session):
    session.model.quote_rows[0].cells["E"] = "DIFFERENT SPEC"
    other = ReviewSession(session.model, MONTH)
    before = session.quote_path.read_bytes()
    with pytest.raises(ValueError):
        other.save(other.products[0])
    assert before == session.quote_path.read_bytes()


def test_attachment_replacement_invalidates_review(session, tmp_path):
    from PIL import Image

    image = tmp_path / "proof.png"
    Image.new("RGB", (10, 10)).save(image)
    p = session.products[0]
    p.channels.append(
        TaskRow(
            "t",
            channel="jd",
            price="2200",
            outcome="price_found",
            state="succeeded",
            evidence_path=image,
            url="https://item.jd.com/1.html",
        )
    )
    c = check(session, "C01")
    session.review(c.id, "通过", "页面可读", "张三")
    Image.new("RGB", (10, 10), color="red").save(image)
    assert check(session, "C01").status == "待复核"


def test_pending_first_quote_and_under_six_months(session):
    p = session.products[0]
    p.context["first_quote_date"] = ""
    assert check(session, "E07").status == "待补充"
    p.context["first_quote_date"] = "2026-08-01"
    assert check(session, "E07").status == "不适用"
    p.context["first_quote_date"] = "2025-01-01"
    assert check(session, "E07").status == "待复核"


def test_corrupt_png_does_not_pass_integrity(session, tmp_path):
    image = tmp_path / "proof.png"
    image.write_bytes(b"not a png")
    p = session.products[0]
    p.channels.append(TaskRow("t", evidence_path=image))
    assert check(session, "C01").status == "待补充"
    assert not check(session, "C01").human_reviewable


def test_extreme_nonfinite_input_does_not_crash_evaluation(session):
    p = session.products[0]
    for value in ("1e999999999", "1e-999999999", "sNaN"):
        p.values["L"] = value
        assert check(session, "E02").status == "待补充"


def test_sidecar_failure_rolls_back_workbook(session, monkeypatch):
    import quote_app.services.review_support as rs

    before = session.quote_path.read_bytes()

    def fail(*args):
        raise OSError("sidecar unavailable")

    monkeypatch.setattr(rs, "atomic_json", fail)
    with pytest.raises(OSError):
        session.save(session.products[0])
    assert session.quote_path.read_bytes() == before


def test_embedded_image_is_preserved(session, tmp_path):
    from PIL import Image
    from openpyxl.drawing.image import Image as ExcelImage

    image = tmp_path / "embedded.png"
    Image.new("RGB", (12, 12), color="blue").save(image)
    w = load_workbook(session.quote_path)
    w["5G手机"].add_image(ExcelImage(image), "AL2")
    w.save(session.quote_path)
    opened = ReviewSession(session.model, MONTH)
    opened.products[0].values.update(session.products[0].values)
    with ZipFile(opened.quote_path) as z:
        before = {n: z.read(n) for n in z.namelist() if "drawing" in n or "media" in n}
    opened.save(opened.products[0])
    with ZipFile(opened.quote_path) as z:
        after = {n: z.read(n) for n in z.namelist() if "drawing" in n or "media" in n}
    assert before and before == after
    assert len(load_workbook(opened.quote_path)["5G手机"]._images) == 1


def test_original_source_row_and_reopen_share_persistent_identity(session):
    row = session.model.quote_rows[0]
    row.source_row_number = 25
    session.model.rows = [
        TaskRow("precise", source_row_number=25, material_code="001", channel="jd", price="2300")
    ]
    original = ReviewSession(session.model, MONTH)
    p = original.products[0]
    p.values.update(session.products[0].values)
    p.context.update(session.products[0].context)
    original.save(p)
    reopened = ReviewSession.from_workbook(session.quote_path, MONTH)
    assert reopened.products[0].id == p.id
    assert reopened.products[0].context["qualification_note"] == p.context["qualification_note"]
    assert reopened.products[0].channels[0].task_id == "precise"


def test_workbook_gaps_keep_date_price_and_values_on_right_product(session):
    w = load_workbook(session.quote_path)
    s = w["5G手机"]
    s["C4"] = "002"
    s["D4"] = "Other"
    s["E4"] = "512白"
    s["H4"] = "2024-01-01"
    s["O4"] = 4567
    s["K4"] = 3456
    w.save(session.quote_path)
    reopened = ReviewSession.from_workbook(session.quote_path, MONTH)
    other = reopened.products[1]
    assert other.output_row == 4
    assert other.context["entry_date"] == "2024-01-01"
    assert other.context["warehouse_price"] == "4567"
    assert other.values["K"] == "3456"
    other.values["AO"] = "正确第4行"
    reopened.save(other)
    assert load_workbook(session.quote_path)["5G手机"]["AO4"].value == "正确第4行"


def test_known_previous_price_constrains_display_without_passing_history(session):
    from quote_app.services.review_support import price_ceiling

    session.products[0].context["previous_price"] = "1900"
    bound, explanation = price_ceiling(session.products[0])
    assert str(bound) == "1900" and "上期" in explanation
    assert check(session, "E03").status != "通过"


def test_noop_draft_save_does_not_invalidate_current_manual_review(session, tmp_path):
    from PIL import Image

    image = tmp_path / "proof.png"
    Image.new("RGB", (10, 10)).save(image)
    p = session.products[0]
    p.channels.append(
        TaskRow(
            "t",
            channel="jd",
            price="2200",
            outcome="price_found",
            state="succeeded",
            evidence_path=image,
            url="https://item.jd.com/1.html",
        )
    )
    session.save(p)
    session.review(check(session, "C01").id, "通过", "已看截图", "张三")
    session.save(p)
    assert check(session, "C01").status == "通过"


def test_reopen_retains_known_source_conflict(session):
    from quote_app.domain.models import Issue

    row = session.model.quote_rows[0]
    row.issues.append(Issue("MARKETING_CONFLICT", "营销表存在冲突候选", False))
    original = ReviewSession(session.model, MONTH)
    original.products[0].values.update(session.products[0].values)
    assert check(original, "A02").status == "未通过"
    original.save(original.products[0])
    reopened = ReviewSession.from_workbook(session.quote_path, MONTH)
    assert check(reopened, "A02").status == "未通过"
    assert not check(reopened, "A02").human_reviewable
    assert "冲突候选" in check(reopened, "A02").reason
    with pytest.raises(ValueError):
        reopened.review(check(reopened, "A02").id, "通过", "不能忽略冲突", "张三")


def test_without_source_metadata_association_cannot_be_reviewed_as_clean(session):
    reopened = ReviewSession.from_workbook(session.quote_path, MONTH)
    assert check(reopened, "A02").status == "待补充"
    assert not check(reopened, "A02").human_reviewable


def test_extreme_future_qualification_dates_do_not_crash(session):
    p = session.products[0]
    for key, code in [("entry_date", "E06"), ("first_quote_date", "E07")]:
        p.context[key] = "9999-12-31"
        assert check(session, code).status == "待补充"
        p.context[key] = "0001-01-01"
        assert check(session, code).status != "通过"
        p.context[key] = "2026-08-01"


def test_corrupt_workbook_is_technical_unchecked(session):
    session.quote_path.write_bytes(b"corrupt workbook")
    c = check(session, "F01")
    assert c.status == "未检查"
    assert "核验" in c.reason


def test_confirmed_wrong_row_remains_business_failure(session):
    session.products[0].material_code = "wrong-code"
    assert check(session, "F01").status == "未通过"


def test_actual_previous_quote_is_automatically_compared_with_scope(session):
    from quote_app.services.review_support import price_ceiling

    w = load_workbook(session.quote_path)
    w["5G手机"]["J2"] = 1800
    w.save(session.quote_path)
    opened = ReviewSession.from_workbook(session.quote_path, MONTH)
    product = opened.products[0]
    product.context["category"] = "手机"
    product.values.update(K="1700", L="2000", M="2500", P="2500", Q="2500")
    ceiling, explanation = price_ceiling(product)
    assert ceiling == 1800
    assert "上期" in explanation
    assert check(opened, "E03").status == "通过"
    product.values["K"] = "2000"
    assert check(opened, "E03").status == "未通过"


def test_reopen_previous_price_comes_from_j_not_sidecar_context(session):
    import json
    from quote_app.services.review_support import price_ceiling

    w = load_workbook(session.quote_path)
    w["5G手机"]["J2"] = 1800
    w.save(session.quote_path)
    opened = ReviewSession.from_workbook(session.quote_path, MONTH)
    p = opened.products[0]
    p.values.update(K="1700", L="2000", M="2500", P="2500", Q="2500")
    opened.save(p)
    data = json.loads(opened.store_path.read_text())
    data["products"][p.id]["context"]["previous_price"] = "9999"
    opened.store_path.write_text(json.dumps(data))
    reopened = ReviewSession.from_workbook(session.quote_path, MONTH)
    assert reopened.products[0].context["previous_price"] == "1800"
    assert price_ceiling(reopened.products[0])[0] == 1800
